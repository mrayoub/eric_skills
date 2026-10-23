#!/usr/bin/env python3
"""Render canonical my_skills deliverables using the ERIC_RC6 visual baseline."""

from __future__ import annotations

import argparse
import base64
import copy
import csv
import hashlib
import html
import io
import json
import textwrap
import re
import unicodedata
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
RENDERER_VERSION = "2.0.0"
DELIVERABLE_LABELS = {
    "briefing": "Entrevista Inicial de sete campos",
    "briefing_one_page": "Briefing",
    "briefing_improvement_questionnaire": "Perguntas para fechar lacunas",
    "briefing_validation": "Avaliação IQB",
    "innovation_potential_assessment": "Avaliação do potencial de inovação",
    "four_d_workshop_record": "Registro do workshop 4D",
    "four_d_report": "Relatório 4D",
    "three_e_workshop_record": "Registro do workshop 3E",
    "three_e_report": "Relatório 3E",
    "innovation_scorecard": "Avaliação das evidências de inovação",
    "experiment_plan": "Plano de experimentação",
    "test_card": "Cartão de teste",
    "learning_card": "Cartão de aprendizado",
    "experiment_report": "Relatório de experimentação",
    "evidence_recommendation": "Evidências e recomendação",
    "executive_presentation": "Apresentação executiva",
    "five_c_workshop_record": "Registro do workshop 5C",
    "five_c_report": "Relatório 5C",
    "cdt": "Cadastro de Demanda Tecnológica (CDT)",
}
COMPLETENESS_LABELS = {
    "complete": "Completo",
    "partial": "Parcial",
    "inconclusive": "Inconclusivo",
}
IQB_TABLE_WIDTHS = {
    "criterios": [14, 15, 10, 41, 10, 10],
    "dimensoes": [14, 7, 10, 12, 9, 10, 10, 28],
    "agrupamento": [35, 65],
}
IQB_SOURCE_LABELS = {
    "initial_interview": "Entrevista Inicial",
    "briefing": "Briefing",
    "transcription": "Transcrição",
    "initial_interview_model": "Modelo da Entrevista Inicial",
    "briefing_model": "Modelo da Entrevista Inicial (legado)",
    "absent": "Ausente",
}
CARD_LABELS = {
    "test_card": {
        "hypothesis": "1. Acreditamos que (hipótese)",
        "test": "2. Para verificar, vamos (teste)",
        "measure": "3. Vamos medir (medida)",
        "criteria": "4. Estaremos certos se (critério)",
    },
    "learning_card": {
        "hypothesis": "1. Acreditávamos que (hipótese testada)",
        "observations": "2. Observamos que (observações)",
        "learning": "3. Aprendemos que (inferência)",
        "action": "4. Portanto, vamos (ação)",
    },
}


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def canonical_bytes(value: Any) -> bytes:
    return (json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2) + "\n").encode("utf-8")


def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def safe_name(value: str) -> str:
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]*", value):
        raise ValueError(f"ID ou versão inadequado para nome de arquivo: {value!r}")
    return value


def identity_line(content: dict[str, Any]) -> str:
    meta = content["metadata"]
    state = COMPLETENESS_LABELS[meta.get("completeness") or "complete"]
    label = f'ID {meta["artifact_id"]} · versão {meta["content_version"]} · {state}'
    if isinstance(meta.get("date"), str) and meta["date"].strip():
        label += f' · data {meta["date"]}'
    return label


def footer_label(content: dict[str, Any]) -> str:
    label = identity_line(content)
    meta = content["metadata"]
    if content["deliverable_type"] == "learning_card" and meta.get("test_card_id"):
        label += f' · Test Card {meta["test_card_id"]} v{meta.get("test_card_version", "")}'
    return label


# Campos dos cards Strategyzer: uma frase objetiva cada.
CARD_FIELD_MAX_CHARS = 280


def header_label(content: dict[str, Any]) -> str:
    kind = content["deliverable_type"]
    profile_name = content.get("metadata", {}).get("content_profile")
    if profile_name:
        definition = read_json(ROOT / "content-profiles.json")["profiles"].get(profile_name)
        if definition and definition["deliverable_type"] == kind:
            return definition["label"]
    return DELIVERABLE_LABELS[kind]


def has_section_content(block: dict[str, Any]) -> bool:
    if block.get("type") in ("heading", "page_break"):
        return False
    data = block.get("data", {})

    def populated(value: Any) -> bool:
        if isinstance(value, str):
            return bool(value.strip())
        if isinstance(value, (list, dict)):
            values = value.values() if isinstance(value, dict) else value
            return any(populated(item) for item in values)
        return value is not None

    return any(populated(data.get(key)) for key in
               ("text", "description", "value", "content", "items", "rows", "cards"))


def normalized_title(value: str) -> str:
    decomposed = unicodedata.normalize("NFD", value.casefold())
    return " ".join("".join(char for char in decomposed
                            if unicodedata.category(char) != "Mn").split())


# Os entregáveis não usam travessão nem meia-risca, nem hífen com espaços no lugar deles.
DASH = re.compile(r"[\u2012\u2013\u2014\u2015]|(?<=\S) - (?=\S)")


def dash_exceptions() -> tuple[str, ...]:
    """Opções do Jira grafadas com hífen e espaços: a CDT precisa citá-las exatamente como na tela."""
    path = ROOT / "scripts/validators/cdt_options.json"
    if not path.is_file():
        return ()
    found: set[str] = set()

    def collect(value: Any) -> None:
        if isinstance(value, str) and DASH.search(value):
            found.add(value)
        elif isinstance(value, dict):
            for key, item in value.items():
                collect(key)
                collect(item)
        elif isinstance(value, list):
            for item in value:
                collect(item)

    collect(json.loads(path.read_text(encoding="utf-8")))
    return tuple(sorted(found, key=len, reverse=True))


DASH_EXCEPTIONS = dash_exceptions()


def dash_paths(value: Any, path: str = "") -> list[str]:
    """Caminhos dos textos do conteúdo que têm travessão."""
    if isinstance(value, str):
        for allowed in DASH_EXCEPTIONS:
            value = value.replace(allowed, "")
        return [path or "(raiz)"] if DASH.search(value) else []
    if isinstance(value, dict):
        return [found for key, item in value.items() for found in dash_paths(item, f"{path}.{key}" if path else key)]
    if isinstance(value, list):
        return [found for index, item in enumerate(value) for found in dash_paths(item, f"{path}[{index}]")]
    return []


def validate(content: dict[str, Any], profiles: dict[str, Any]) -> dict[str, Any]:
    try:
        from jsonschema import Draft202012Validator
        from referencing import Registry, Resource
    except ImportError as exc:
        raise RuntimeError("Validação requer jsonschema e referencing: pip install jsonschema") from exc

    kind = content.get("deliverable_type")
    profile = profiles["types"].get(kind)
    if profile is None:
        raise ValueError(f"deliverable_type sem perfil visual: {kind!r}")
    meta = content.get("metadata", {})
    if not isinstance(meta, dict):
        raise ValueError("metadata deve ser um objeto")
    for key in ("artifact_id", "content_version"):
        if not isinstance(meta.get(key), str):
            raise ValueError(f"metadata.{key} deve ser string")
        safe_name(meta[key])
    if meta.get("orientation") not in (None, "portrait" if profile["paper"] == "a4_portrait" else "landscape" if profile["paper"] == "a4_landscape" else "wide"):
        raise ValueError("metadata.orientation contradiz a variante visual do perfil")
    if meta.get("completeness") not in (None, "complete", "partial", "inconclusive"):
        raise ValueError("metadata.completeness deve ser complete, partial ou inconclusive")

    envelope = read_json(ROOT / "contracts" / "deliverable-content.schema.json")
    registry = Registry().with_resource(envelope["$id"], Resource.from_contents(envelope))
    effective_completeness = meta.get("completeness") or "complete"
    schema_file = profile.get("schema") if effective_completeness == "complete" else None
    schema = read_json(ROOT / "contracts" / schema_file) if schema_file else envelope
    errors = sorted(Draft202012Validator(schema, registry=registry).iter_errors(content), key=lambda e: str(e.path))
    if errors:
        raise ValueError("Conteúdo inválido: " + "; ".join(f"{list(e.path)}: {e.message}" for e in errors[:8]))

    # Perfil de conteúdo opcional: exige seções além do contrato do tipo nas versões completas.
    content_profile = meta.get("content_profile")
    if content_profile is not None:
        profiles_file = read_json(ROOT / "content-profiles.json")["profiles"]
        definition = profiles_file.get(content_profile)
        if definition is None:
            raise ValueError(f"metadata.content_profile desconhecido: {content_profile!r}; "
                             f"use um de {sorted(profiles_file)}")
        if definition["deliverable_type"] != kind:
            raise ValueError(f"{content_profile}: perfil de {definition['deliverable_type']}, não de {kind}")
        if effective_completeness == "complete":
            present = {block.get("data", {}).get("section_key") for block in content.get("blocks", [])
                       if has_section_content(block)}
            for key, label in definition["required_sections"].items():
                if key not in present:
                    raise ValueError(f"{content_profile}: falta a seção {key} ({label}); "
                                     "acrescente-a ou marque a versão como parcial")

    # O tipo do entregável já aparece no cabeçalho, acima do título; o título diz o assunto.
    title_key = normalized_title(str(content.get("title", "")))
    for name in (header_label(content), DELIVERABLE_LABELS[kind], kind.replace("_", " ")):
        normalized_name = normalized_title(name)
        if normalized_name and re.search(rf"(?<!\w){re.escape(normalized_name)}(?!\w)", title_key):
            raise ValueError(f"title: não repita o tipo do entregável ({name}), que já aparece no "
                             "cabeçalho; escreva o assunto ou a mensagem do documento")

    dashed = dash_paths(content)
    if dashed:
        raise ValueError("travessão no texto: " + ", ".join(dashed[:8])
                         + "; troque por dois-pontos, vírgula, parênteses ou ponto, e intervalos por \"de X a Y\"")

    if kind in ("test_card", "learning_card"):
        keys = ("hypothesis", "test", "measure", "criteria") if kind == "test_card" else ("hypothesis", "observations", "learning", "action")
        blocks = content["blocks"]
        if len(blocks) != 4:
            raise ValueError(f"{kind}: exige exatamente quatro campos")
        for block, key in zip(blocks, keys):
            data = block["data"]
            text = data.get("text")
            if data.get("section_key") != key:
                raise ValueError(f"{kind}: campo esperado na ordem do método: {key}")
            if block["type"] not in ("paragraph", "heading", "callout", "metric"):
                raise ValueError(f"{kind}.{key}: campo deve usar um bloco textual visível")
            if not isinstance(text, str) or not text.strip():
                raise ValueError(f"{kind}.{key}: exige texto ou indicação explícita de pendência")
            if len(text.strip()) > CARD_FIELD_MAX_CHARS:
                raise ValueError(f"{kind}.{key}: {len(text.strip())} caracteres; o campo do card aceita até "
                                 f"{CARD_FIELD_MAX_CHARS}. Escreva uma frase objetiva e mova detalhes ao plano ou caderno")

    ids = [block["id"] for block in content["blocks"]]
    if len(ids) != len(set(ids)):
        raise ValueError("IDs de blocos devem ser únicos")
    citation_ids = {citation["id"] for citation in content["citations"]}
    for block in content["blocks"]:
        unknown = set(block.get("citation_ids", [])) - citation_ids
        if block["type"] == "card_grid":
            cards = block["data"].get("cards", [])
            if not 1 <= len(cards) <= 12:
                raise ValueError(f"{block['id']}: card_grid exige 1 a 12 cards")
            if block["data"].get("columns", 2) not in (1, 2, 3, 4):
                raise ValueError(f"{block['id']}: card_grid aceita 1 a 4 colunas")
            for card in cards:
                if not isinstance(card, dict) or not card.get("title"):
                    raise ValueError(f"{block['id']}: cada card exige title")
                unknown |= set(card.get("citation_ids", [])) - citation_ids
        if unknown:
            raise ValueError(f"{block['id']}: citações inexistentes: {sorted(unknown)}")

    def nested_citations(value: Any) -> set[str]:
        found: set[str] = set()
        if isinstance(value, dict):
            for key, item in value.items():
                if key == "citation_id" and isinstance(item, str):
                    found.add(item)
                elif key == "citation_ids" and isinstance(item, list):
                    found.update(entry for entry in item if isinstance(entry, str))
                found.update(nested_citations(item))
        elif isinstance(value, list):
            for item in value:
                found.update(nested_citations(item))
        return found

    for block in content["blocks"]:
        unknown = nested_citations(block["data"]) - citation_ids
        if unknown:
            raise ValueError(f"{block['id']}: citações aninhadas inexistentes: {sorted(unknown)}")
    if effective_completeness == "complete":
        import importlib.util

        spec = importlib.util.spec_from_file_location("deliverable_validators", ROOT / "scripts/validators/__init__.py")
        validators = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(validators)
        semantic_errors = validators.run(content)
        if semantic_errors:
            raise ValueError("Inconsistência semântica: " + "; ".join(semantic_errors[:12]))
    return profile


# Rótulos em português para campos técnicos exibidos ao leitor; os demais viram "Nome do campo".
FIELD_LABELS = {
    "points": "Pontos", "responsibles": "Responsáveis", "sponsor": "Patrocinador",
    "owner": "Responsável", "name": "Nome", "score": "Nota", "overall_score": "Nota geral",
    "classification": "Classificação", "diagnosis": "Diagnóstico", "original": "Original",
    "analysis": "Análise", "suggestion": "Sugestão", "anti_clustering": "Verificação de notas agrupadas",
    "triggered": "Acionada", "original_scores": "Notas originais", "span": "Amplitude",
    "re_evaluation_notes": "Notas da reavaliação",
}


def field_label(key: str) -> str:
    return FIELD_LABELS.get(key, key.replace("_", " ").capitalize())


def block_text(block: dict[str, Any]) -> str:
    data = block.get("data", {})
    def value_text(value: Any) -> str:
        if isinstance(value, dict):
            return "; ".join(f"{field_label(key)}: {value_text(item)}" for key, item in value.items())
        if isinstance(value, list):
            return "; ".join(value_text(item) for item in value)
        return "" if value is None else str(value)

    parts = []
    for key in ("title", "label", "text", "description", "content", "value"):
        if key in data and data[key] not in (None, "", []):
            parts.append(value_text(data[key]))
    for key, value in data.items():
        if key in {"title", "label", "text", "description", "content", "value", "section_key", "layout", "level", "ordered", "headers", "columns", "rows", "tone"}:
            continue
        if value not in (None, "", []):
            parts.append(f"{field_label(key)}: {value_text(value)}")
    return "\n".join(parts) or json.dumps(data, ensure_ascii=False, sort_keys=True)


def table_rows(block: dict[str, Any]) -> tuple[list[str], list[list[str]]]:
    data = block.get("data", {})
    if isinstance(data.get("dimensions"), dict):
        _, _, headers, rows = table_views(block)[0]
        return headers, rows
    columns = data.get("headers") or data.get("columns") or []
    keys = [str(h.get("key", h.get("label", ""))) if isinstance(h, dict) else str(h) for h in columns]
    headers = [str(h.get("label", h.get("key", ""))) if isinstance(h, dict) else str(h) for h in columns]
    raw_rows = data.get("rows") or []
    if raw_rows and not headers and isinstance(raw_rows[0], dict):
        keys = headers = list(raw_rows[0])
    rows: list[list[str]] = []
    for row in raw_rows:
        if isinstance(row, list):
            rows.append([str(x) if x is not None else "" for x in row])
        elif isinstance(row, dict):
            values = [row.get(k) for k in keys] if keys else row.values()
            rows.append([str(value) if value is not None else "" for value in values])
        else:
            rows.append([str(row)])
    if not headers and not rows:
        rows = [[field_label(key), "; ".join(map(str, value)) if isinstance(value, list) else str(value)]
                for key, value in data.items() if key not in ("section_key", "layout")]
        headers = ["Campo", "Valor"]
    return headers, rows


def table_views(block: dict[str, Any]) -> list[tuple[str, str, list[str], list[list[str]]]]:
    """Retorna [(slug, legenda, cabeçalhos, linhas)]; três vistas para o IQB."""
    data = block.get("data", {})
    if not isinstance(data.get("dimensions"), dict):
        return [("", data.get("title", ""), *table_rows(block))]
    labels = {"clarity": "Clareza", "evidence": "Evidências", "outcomes": "Resultados", "actionability": "Acionabilidade"}
    criteria: list[list[str]] = []
    dimensions: list[list[str]] = []
    for dimension, detail in data["dimensions"].items():
        label = labels.get(dimension, dimension)
        for item in detail.get("items", []):
            criterion = str(item.get("criterion_id", ""))
            if criterion == detail.get("veto_item_id"):
                criterion += " (veto)"
            source = str(item.get("evidence_source", ""))
            criteria.append([label, criterion, "Sim" if item.get("answer") is True else "Não",
                             str(item.get("evidence", "")), IQB_SOURCE_LABELS.get(source, source),
                             str(item.get("citation_id") or "")])
        veto = f"Disparado: {detail.get('veto_detail') or 'sem detalhe'}" if detail.get("veto_triggered") else "Não disparado"
        dimensions.append([label, str(detail.get("yes_count", "")), str(detail.get("raw_score", "")),
                           veto, str(detail.get("veto_max_score", "")), str(detail.get("final_score", "")),
                           str(detail.get("label", "")), str(detail.get("justification", ""))])
    anti = data.get("anti_clustering", {})
    originals = anti.get("original_scores")
    audit = [["Disparada", "Sim" if anti.get("triggered") else "Não"],
             ["Amplitude", str(anti.get("span", ""))],
             ["Notas originais", ", ".join(map(str, originals)) if originals is not None else "não informadas"],
             ["Notas da reavaliação", str(anti.get("re_evaluation_notes") or "não informadas")]]
    return [
        ("criterios", "Critérios avaliados", ["Dimensão", "Critério", "Resposta", "Evidência", "Fonte", "Citação"], criteria),
        ("dimensoes", "Notas por dimensão", ["Dimensão", "Sim (de 9)", "Nota bruta", "Veto", "Teto do veto", "Nota final", "Rótulo", "Justificativa"], dimensions),
        ("agrupamento", "Auditoria de agrupamento", ["Campo", "Valor"], audit),
    ]


def card_text(card: dict[str, Any]) -> str:
    parts = []
    if card.get("value") is not None:
        parts.append(str(card["value"]))
    if card.get("text"):
        parts.append(str(card["text"]))
    parts.extend(str(x) for x in card.get("items", []))
    body = "; ".join(x for x in parts if x)
    title = str(card.get("title", ""))
    return f"{title}: {body}" if title and body else title or body


_PDF_MODULE: Any = None
PT_PER_MM = 72 / 25.4


def pdf_module() -> Any:
    """scripts/pdf_reportlab.py, que também define estilos e paginação usados pelo HTML."""
    global _PDF_MODULE
    if _PDF_MODULE is None:
        import importlib.util

        spec = importlib.util.spec_from_file_location("pdf_reportlab", ROOT / "scripts/pdf_reportlab.py")
        _PDF_MODULE = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(_PDF_MODULE)
    return _PDF_MODULE


def reportlab_styles(tokens: dict[str, Any], layout: str) -> dict[str, Any] | None:
    """Estilos do ReportLab para medir texto como o PDF; None sem a biblioteca instalada."""
    try:
        pdf_module()._register_fonts(tokens)
        return pdf_module().make_styles(tokens, layout)
    except ImportError:
        return None


def _number(value: float) -> str:
    return f"{round(value, 3):g}"


def style_rules(tokens: dict[str, Any], specs: dict[str, dict[str, Any]], scope: str = "") -> str:
    """Uma classe .st-<nome> por estilo do PDF: mesma família, peso, tamanho, entrelinha e cor."""
    rules = []
    for name, spec in specs.items():
        weight = 700 if spec["family"] == "font_brand" or spec["weight"] >= 600 else 500 if spec["weight"] >= 500 else 400
        selector = f"{scope}.st-{name}" + (f", {scope}.eyebrow" if name == "eyebrow" else "")
        rules.append(f'{selector} {{ font-family: {tokens["typography"][spec["family"]]}, Arial, sans-serif; '
                     f'font-weight: {weight}; font-size: {_number(spec["size"])}pt; '
                     f'line-height: {_number(spec["leading"])}pt; color: {tokens["color"][spec["color"]]}; }}')
    return "\n".join(rules) + "\n"


def base_css(tokens: dict[str, Any]) -> str:
    color = tokens["color"]
    font_faces = ""
    for family in dict.fromkeys((tokens["typography"]["font_brand"], tokens["typography"]["font_body"])):
        font_file = ROOT / "assets" / "fonts" / f"{family}-Variable.ttf"
        if not font_file.is_file():
            raise FileNotFoundError(f"Fonte {family} incorporada ausente em assets/fonts")
        data = base64.b64encode(font_file.read_bytes()).decode("ascii")
        font_faces += (f"@font-face {{ font-family: {family}; src: url(data:font/ttf;base64,{data}) "
                       "format('truetype'); font-weight: 100 900; }\n")
    return f"""{font_faces}* {{ box-sizing: border-box; }}
html {{ background: {color['surface_muted']}; }}
body {{ margin: 0; padding: 8mm 0; color: {color['text_body']}; font-family: {tokens['typography']['font_body']}, Arial, sans-serif; }}
p, h1, h2, h3, ul, ol {{ margin: 0; padding: 0; }}
sup {{ font-size: 7.5pt; line-height: 0; vertical-align: 3pt; }}
.gap {{ display: block; }}
.sheet {{ position: relative; margin: 0 auto 8mm; background: {color['surface']}; box-shadow: 0 3px 28px #10266126; }}
.eyebrow {{ text-transform: uppercase; }}
.masthead {{ display: grid; grid-template-columns: 1fr 40mm; align-items: center; padding-bottom: 5pt; border-bottom: 1.5pt solid {color['document_primary']}; }}
.masthead img {{ display: block; }}
.band {{ background: {color['document_deep']}; padding: 9pt 12pt; text-transform: uppercase; }}
.boxed-table {{ border-collapse: collapse; table-layout: fixed; }}
.boxed-table th {{ background: {color['document_primary']}; text-align: left; }}
.boxed-table th, .boxed-table td {{ border: .75pt solid {color['border']}; vertical-align: top; overflow-wrap: break-word; }}
.boxed-table tbody tr:nth-child(even) {{ background: {color['surface_muted']}; }}
.grid {{ display: grid; column-gap: 6pt; row-gap: 6pt; margin-bottom: 6pt; }}
.grid-card {{ padding: 7pt 9pt; border-left: 3pt solid {color['document_accent']}; background: {color['document_accent_soft']}; }}
.grid-card .st-card_title {{ text-transform: uppercase; }}
.grid-card.tone-neutral {{ background: {color['surface_muted']}; border-left-color: {color['border_strong']}; }}
.grid-card.tone-red {{ background: {color['danger_soft']}; border-left-color: {color['danger']}; }}
.grid-card.tone-green {{ background: #E8F4EE; border-left-color: #3F9270; }}
.grid-card.tone-yellow {{ background: #FFF4D9; border-left-color: #BD8A28; }}
.metric-box {{ background: {color['document_accent_soft']}; padding: 3pt 8pt; }}
.callout-box {{ background: {color['document_deep']}; border-left: 3pt solid {color['document_accent']}; padding: 8pt 9pt; }}
.missing-box {{ background: {color['danger_soft']}; border-left: 3pt solid {color['danger']}; padding: 3pt 8pt; }}
@media print {{ html, body {{ background: {color['surface']}; padding: 0; }} .sheet {{ margin: 0; box-shadow: none; }} }}
"""


def _gap(points: float) -> str:
    return f'<div class="gap" style="height: {_number(points)}pt"></div>'


# Negrito em Markdown (**texto**), usado na Descrição da CDT.
BOLD = re.compile(r"\*\*(.+?)\*\*")


def _text(value: Any) -> str:
    return BOLD.sub(r"<strong>\1</strong>", html.escape(str(value))).replace("\n", "<br>")


def add_rich_text(paragraph: Any, text: str) -> None:
    """Acrescenta o texto ao parágrafo do DOCX, com os trechos entre ** em negrito."""
    for index, part in enumerate(BOLD.split(text)):
        if part:
            paragraph.add_run(part).bold = bool(index % 2)


def _refs(ids: list[str]) -> str:
    return "".join(f"<sup>[{html.escape(str(ref))}]</sup>" for ref in ids)


def _box_width(inset: float) -> str:
    return f"width: calc(100% - {_number(inset)}pt);" if inset else "width: 100%;"


def _logo(tokens: dict[str, Any]) -> tuple[str, float, float]:
    """URI do logo e tamanho em mm, com a mesma escala do PDF (cabe em 38 x 11 mm)."""
    data = (ROOT / "assets" / "logo_ccee.png").read_bytes()
    width, height = int.from_bytes(data[16:20], "big"), int.from_bytes(data[20:24], "big")
    scale = min(38 / width, 11 / height)
    return "data:image/png;base64," + base64.b64encode(data).decode("ascii"), width * scale, height * scale


def html_masthead(content: dict[str, Any], tokens: dict[str, Any]) -> str:
    uri, width, height = _logo(tokens)
    return (f'<header class="masthead"><span class="eyebrow">{html.escape(header_label(content))}</span>'
            f'<img src="{uri}" alt="CCEE" style="width: {_number(width)}mm; height: {_number(height)}mm;"></header>')


def html_flowables(block: dict[str, Any], inset: float = 0) -> str:
    """Espelha block_flowables do PDF."""
    kind, data = block["type"], block.get("data", {})
    refs = _refs(block.get("citation_ids", []))
    if kind in ("heading", "page_break"):
        return ""
    if kind == "table":
        headings, rows = table_rows(block)
        ncols = max(1, len(headings), *(len(row) for row in rows)) if rows else max(1, len(headings))
        cell = 'style="padding: 5pt 6pt;"'
        head = "".join(f'<th class="st-table_header" {cell}>{_text(headings[i] if i < len(headings) else "")}</th>'
                       for i in range(ncols))
        body = "".join("<tr>" + "".join(f'<td class="st-table_body" {cell}>{_text(row[i] if i < len(row) else "")}</td>'
                                        for i in range(ncols)) + "</tr>" for row in rows)
        table = (f'<table class="boxed-table" style="{_box_width(inset)}"><thead><tr>{head}</tr></thead>'
                 f'<tbody>{body}</tbody></table>')
        return table + (f'<p class="st-body">{refs}</p>' if refs else "")
    if kind == "list":
        return "".join(f'<p class="st-body">• {_text(item)}{refs}</p>' for item in data.get("items", []))
    if kind == "metric":
        return f'<div class="metric-box" style="{_box_width(inset)}"><p class="st-metric">{_text(block_text(block))}{refs}</p></div>'
    if kind == "callout":
        return (f'<div class="callout-box" style="{_box_width(inset)}">'
                f'<p class="st-body_inverse">{_text(block_text(block))}{refs}</p></div>')
    return f'<p class="st-body">{_text(block_text(block))}{refs}</p>'


def html_card_grid(block: dict[str, Any], inset: float = 0, max_columns: int = 4) -> str:
    """Espelha _report_card_grid do PDF."""
    data = block.get("data", {})
    columns = max(1, min(max_columns, int(data.get("columns", 2))))
    result = f'<p class="st-card_title" style="margin-bottom: 5pt;">{_text(data["title"])}</p>' if data.get("title") else ""
    cards = []
    for card in data.get("cards", []):
        tone = card.get("tone", "blue")
        tone = tone if tone in ("blue", "neutral", "red", "green", "yellow") else "blue"
        lines = [f'<p class="st-card_title">{_text(card.get("title", ""))}</p>']
        if card.get("value") is not None:
            lines.append(f'<p class="st-metric">{_text(card["value"])}</p>')
        if card.get("text"):
            lines.append(f'<p class="st-body">{_text(card["text"])}</p>')
        lines.extend(f'<p class="st-body">• {_text(item)}</p>' for item in card.get("items", []))
        if card.get("citation_ids"):
            lines.append(f'<p class="st-source">{_refs(card["citation_ids"])}</p>')
        cards.append(f'<article class="grid-card tone-{tone}">{"".join(lines)}</article>')
    result += (f'<div class="grid" style="grid-template-columns: repeat({columns}, minmax(0, 1fr)); {_box_width(inset)}">'
               + "".join(cards) + "</div>")
    if block.get("citation_ids"):
        result += f'<p class="st-source">{_refs(block["citation_ids"])}</p>'
    return result


def html_report_table(block: dict[str, Any], inset: float, width_pt: float, measure: dict[str, Any] | None,
                      tokens: dict[str, Any]) -> str:
    """Espelha _report_table do PDF, com as mesmas larguras de coluna quando o ReportLab está instalado."""
    data = block.get("data", {})
    views = table_views(block) if isinstance(data.get("dimensions"), dict) else [("", data.get("title", ""), *table_rows(block))]
    result = ""
    for slug, caption, headings, rows in views:
        if slug == "dimensoes":
            result += html_score_panel({**block, "citation_ids": []}, tokens, measure, width_pt) + _gap(7)
            continue
        if caption:
            result += f'<p class="st-card_title" style="margin-bottom: 4pt;">{_text(caption)}</p>'
        ncols = max(1, len(headings), *(len(row) for row in rows)) if rows else max(1, len(headings))
        proportions = IQB_TABLE_WIDTHS.get(slug)
        padding = 4 if proportions else 6
        columns = ""
        if measure is not None:
            widths = pdf_module()._content_widths([headings, *rows], ncols, width_pt, padding, measure,
                                                  proportions if proportions and len(proportions) == ncols else None)
            columns = "<colgroup>" + "".join(f'<col style="width: {_number(100 * w / sum(widths))}%">' for w in widths) + "</colgroup>"
        cell = f'style="padding: {padding}pt;"'
        head = "".join(f'<th class="st-table_header" {cell}>{_text(headings[i] if i < len(headings) else "")}</th>'
                       for i in range(ncols))
        body = "".join("<tr>" + "".join(f'<td class="st-table_body" {cell}>{_text(row[i] if i < len(row) else "")}</td>'
                                        for i in range(ncols)) + "</tr>" for row in rows)
        result += (f'<table class="boxed-table" style="{_box_width(inset)}">{columns}<thead><tr>{head}</tr></thead>'
                   f'<tbody>{body}</tbody></table>' + _gap(7))
    if block.get("citation_ids"):
        result += f'<p class="st-source">{_refs(block["citation_ids"])}</p>'
    return result


def html_report_block(block: dict[str, Any], inset: float, width_pt: float, measure: dict[str, Any] | None,
                      tokens: dict[str, Any]) -> str:
    """Espelha _report_block do PDF."""
    if block["type"] == "table":
        return html_report_table(block, inset, width_pt, measure, tokens)
    if block["type"] == "card_grid":
        return html_card_grid(block, inset)
    items = html_flowables(block, inset)
    if block.get("data", {}).get("status") in ("missing", "pending", "inconclusive"):
        return f'<div class="missing-box" style="{_box_width(inset)}">{items}</div>'
    return items


def score_summary(block: dict[str, Any]) -> str:
    """Total e classe do painel de notas por extenso, para os formatos sem o painel visual."""
    if not pdf_module().is_score_panel(block):
        return ""
    panel = pdf_module().score_panel_data(block)
    return f"Total ponderado: {panel['total']} · Classificação: {panel['badge']}"


def html_score_panel(block: dict[str, Any], tokens: dict[str, Any], measure: dict[str, Any] | None, width_pt: float) -> str:
    """Espelha _score_panel do PDF, com as mesmas larguras de coluna quando o ReportLab está instalado."""
    panel = pdf_module().score_panel_data(block)
    widths = pdf_module().panel_widths(panel, measure, width_pt) if measure is not None else list(panel["shares"])
    background, foreground = panel["tone"]
    color = tokens["color"]
    badge = (f'<span class="score-badge st-badge" style="background: {color[background]}; color: {color[foreground]};">'
             f'{html.escape(panel["badge"])}</span>')
    columns = "".join(f'<col style="width: {_number(100 * width / sum(widths))}%">' for width in widths)
    heads = "".join(f'<th class="st-score_head">{_text(head)}</th>' for head in panel["heads"])
    rows = ""
    for row in panel["rows"]:
        cells = ""
        for index, value in enumerate(row["cells"]):
            if index == panel["bar"]:
                bar = f'<div style="width: {_number(row["bar"])}%"></div>' if row["bar"] is not None else ""
                cells += f'<td><div class="score-bar">{bar}</div><p class="st-score_value">{_text(value)}</p></td>'
            else:
                style = "st-card_title" if index in panel["strong"] else "st-table_body"
                cells += f'<td class="{style}">{_text(value)}</td>'
        rows += f"<tr>{cells}</tr>"
    total = [f'<td class="st-card_title" colspan="{panel["value"]}">{_text(panel["total_label"])}</td>']
    for index in range(panel["value"], len(panel["heads"])):
        if index == panel["value"]:
            total.append(f'<td class="st-card_title">{_text(panel["total"])}</td>')
        elif index == panel["badge_column"]:
            total.append(f'<td class="st-score_value">{_text(panel["badge"])}</td>')
        else:
            total.append("<td></td>")
    rows += f'<tr class="score-total">{"".join(total)}</tr>'
    result = (f'<div class="score-panel"><div class="score-header"><span class="st-section">{_text(panel["title"].upper())}</span>'
              f'<span class="st-section">{_text(panel["total_prefix"])}: {_text(panel["total"])}</span>{badge}</div>'
              f'<table class="score-table"><colgroup>{columns}</colgroup><thead><tr>{heads}</tr></thead><tbody>{rows}</tbody></table></div>')
    if block.get("citation_ids"):
        result += _gap(3) + f'<p class="st-source">{_refs(block["citation_ids"])}</p>'
    return result


def html_flow_row(row: list[dict[str, Any]], tokens: dict[str, Any], width_pt: float, measure: dict[str, Any] | None) -> str:
    """Espelha _flow_row do PDF: card compacto, caixa escura sem título ou painel de notas."""
    pdf = pdf_module()
    half_pt = (width_pt - 8) / 2
    if len(row) == 1 and row[0]["mode"] == "tabular_list":
        cell = row[0]
        cell_pt = half_pt if cell["half"] else width_pt
        inner = "".join(html_score_panel(block, tokens, measure, cell_pt) if pdf.is_score_panel(block)
                        else html_report_block(block, 0, cell_pt, measure, tokens) for block in cell["blocks"])
        return f'<div style="width: {"calc(50% - 4pt)" if cell["half"] else "100%"};">{inner}</div>'
    cells = []
    for cell in row:
        dark = cell["mode"] == "content_only"
        cell_pt = (half_pt if cell["half"] else width_pt) - 18
        items = []
        label = ""
        if dark and cell["heading"] is not None:
            label = f'<p class="flow-label st-callout_label">{_text(block_text(cell["heading"]))}</p>' + _gap(3)
        for block in cell["blocks"]:
            if dark and block["type"] not in ("table", "card_grid"):
                items.append(f'<p class="st-body_inverse">{_text(block_text(block))}{_refs(block.get("citation_ids", []))}</p>')
            else:
                items.append(html_report_block(block, 0, cell_pt, measure, tokens))
        body = label + _gap(4).join(items)
        has_band = cell["mode"] == "title_and_content" and cell["heading"] is not None
        band = f'<h2 class="flow-band st-section">{_text(block_text(cell["heading"]))}</h2>' if has_band else ""
        kind = "flow-dark" if dark else "flow-card"
        cells.append(f'<section class="{kind}">{band}<div class="flow-body">{body}</div></section>')
    columns = "minmax(0, 1fr) minmax(0, 1fr)" if row[0]["half"] else "minmax(0, 1fr)"
    return f'<div class="flow-row" style="grid-template-columns: {columns};">{"".join(cells)}</div>'


def flow_css(tokens: dict[str, Any]) -> str:
    color = tokens["color"]
    return f""".flow-row {{ display: grid; column-gap: 8pt; break-inside: avoid; }}
.flow-card {{ display: flex; flex-direction: column; border: .75pt solid {color['border']}; background: {color['surface_muted']}; }}
.flow-band {{ background: {color['document_deep']}; padding: 5pt 9pt; text-transform: uppercase; }}
.flow-body {{ padding: 8pt 9pt; }}
.flow-label {{ text-transform: uppercase; }}
.flow-dark {{ background: {color['document_deep']}; border-left: 3pt solid {color['document_accent']}; }}
.score-panel {{ break-inside: avoid; }}
.score-header {{ display: grid; grid-template-columns: 1fr auto 82pt; align-items: center; column-gap: 8pt; background: {color['document_deep']}; padding: 6pt 9pt; }}
.score-badge {{ display: block; text-align: center; border-radius: 8pt; padding: 3pt 0; }}
.score-table {{ width: 100%; border-collapse: collapse; table-layout: fixed; border: .75pt solid {color['border']}; }}
.score-table th {{ background: {color['surface_muted']}; text-align: left; }}
.score-table th, .score-table td {{ padding: 5pt 7pt; vertical-align: middle; border-bottom: .75pt solid {color['border']}; overflow-wrap: break-word; }}
.score-table tr.score-total td {{ background: {color['surface_muted']}; border-top: 1.5pt solid {color['border_strong']}; border-bottom: 0; }}
.score-bar {{ height: 4pt; margin: 3pt 0 0; border-radius: 2pt; background: {color['border']}; }}
.score-bar > div {{ height: 100%; border-radius: 2pt; background: {color['document_accent']}; }}
"""


def html_sources(content: dict[str, Any]) -> str:
    lines = []
    for citation in content["citations"]:
        line = f"<b>{html.escape(citation['id'])}</b>: {html.escape(citation['source_id'])}"
        if citation.get("locator"):
            line += f", {html.escape(citation['locator'])}"
        lines.append(f'<p class="st-source">{line}</p>')
    return "".join(lines)


def page_css(tokens: dict[str, Any], paper: str, footer: str | None) -> str:
    page = tokens["page"][paper]
    margins = (f"{page['margin_top']}mm {page['margin_right']}mm "
               f"{page['margin_bottom'] + 10}mm {page['margin_left']}mm")
    if footer is None:
        return (f"@page {{ size: {page['width']}mm {page['height']}mm; margin: 0; }}\n"
                f".sheet {{ width: {page['width']}mm; height: {page['height']}mm; padding: {margins}; overflow: hidden; }}\n")
    spec = pdf_module().style_specs(tokens, "report_portrait")["footer"]
    label = json.dumps(footer + " · página ", ensure_ascii=False).replace("<", "\\003c ")
    return (f"@page {{ size: {page['width']}mm {page['height']}mm; margin: {margins}; "
            f"@bottom-right {{ content: {label} counter(page) \" de \" counter(pages); "
            f"font-family: {tokens['typography'][spec['family']]}; font-size: {_number(spec['size'])}pt; "
            f"color: {tokens['color'][spec['color']]}; }} }}\n"
            f".sheet {{ width: {page['width']}mm; min-height: {page['height']}mm; padding: {margins}; }}\n"
            ".sheet > footer.meta { position: absolute; right: "
            f"{page['margin_right']}mm; bottom: {page['margin_bottom'] / 2}mm; }}\n"
            "@media print { .sheet { width: auto; min-height: 0; padding: 0; } .sheet > footer.meta { display: none; } }\n")


def html_document(content: dict[str, Any], css_text: str, body: str) -> str:
    return ("<!doctype html>\n<html lang=\"pt-BR\"><head><meta charset=\"utf-8\">"
            f"<title>{html.escape(content['title'])}</title><style>{css_text}</style></head>"
            f"<body>{body}</body></html>\n")


def render_report_html(content: dict[str, Any], tokens: dict[str, Any], profile: dict[str, Any]) -> str:
    """Relatórios em retrato e paisagem com a mesma sequência de blocos e medidas de _render_report."""
    content = pdf_module().with_panel_totals(content)
    layout, page = profile["layout"], tokens["page"][profile["paper"]]
    usable_pt = (page["width"] - page["margin_left"] - page["margin_right"]) * PT_PER_MM
    column_pt = (usable_pt - 10) / 2
    measure = reportlab_styles(tokens, layout)
    parts = [html_masthead(content, tokens), _gap(8), f'<h1 class="st-report_title">{_text(content["title"])}</h1>']
    if content.get("summary"):
        parts += [_gap(6), f'<p class="st-summary">{_text(content["summary"])}</p>']
    parts.append(_gap(12))

    def band(title: str) -> str:
        return f'<h2 class="band st-report_section">{_text(title)}</h2>'

    remaining = content["blocks"]
    slots = (load_presentation_layout(content["deliverable_type"])
             if content["deliverable_type"] in pdf_module().FLOW_LAYOUT_TYPES else None)
    if slots:
        rows, remaining = pdf_module().plan_flow(content, slots)
        for row in rows:
            parts += [html_flow_row(row, tokens, usable_pt, measure), _gap(8)]
        parts.append(_gap(4))
    sections = pdf_module().group_sections(remaining, block_text)
    position = 0
    while position < len(sections):
        item = sections[position]
        if item[0] == "page_break":
            parts.append('<div class="page-break" style="break-before: page;"></div>')
            position += 1
            continue
        if (not item[3] and position + 1 < len(sections)
                and content["deliverable_type"] not in pdf_module().SINGLE_COLUMN_TYPES
                and sections[position + 1][0] == "section" and not sections[position + 1][3]):
            cells = []
            for _, title, blocks, _ in sections[position:position + 2]:
                inner = "".join(html_report_block(block, 22, column_pt - 22, measure, tokens) + _gap(7) for block in blocks)
                cells.append(f'<section class="report-section">{band(title)}{_gap(7)}{inner}</section>')
            parts.append('<div class="section-pair" style="display: grid; grid-template-columns: minmax(0, 1fr) 10pt minmax(0, 1fr); '
                         f'break-inside: avoid;">{cells[0]}<div></div>{cells[1]}</div>' + _gap(10))
            position += 2
            continue
        _, title, blocks, _ = item
        inner = "".join(html_report_block(block, 0, usable_pt - 22, measure, tokens) + _gap(7) for block in blocks)
        parts.append(f'<section class="report-section">{band(title)}{_gap(7)}'
                     f'<div style="padding: 0 11pt;">{inner}</div></section>' + _gap(10))
        position += 1
    if content["citations"]:
        parts += ['<h2 class="st-field">Fontes</h2>', _gap(4), html_sources(content)]
    label = footer_label(content)
    parts.append(f'<footer class="meta st-footer">{html.escape(label)}</footer>')
    css_text = (base_css(tokens) + style_rules(tokens, pdf_module().style_specs(tokens, layout))
                + flow_css(tokens) + page_css(tokens, profile["paper"], label))
    return html_document(content, css_text, f'<article class="sheet document">{"".join(parts)}</article>')


def render_card_html(content: dict[str, Any], tokens: dict[str, Any], profile: dict[str, Any]) -> str:
    """Test Card e Learning Card numa página, como o ramo card_portrait de render_pdf."""
    page = tokens["page"][profile["paper"]]
    fields = CARD_LABELS[content["deliverable_type"]]
    parts = [html_masthead(content, tokens), _gap(9), f'<h1 class="st-title">{_text(content["title"])}</h1>']
    if content.get("summary"):
        parts += [_gap(6), f'<p class="st-summary">{_text(content["summary"])}</p>']
    parts.append(_gap(9))
    for block in content["blocks"]:
        key = block.get("data", {}).get("section_key")
        if key in fields:
            field_block = {**block, "type": "paragraph"} if block["type"] == "heading" else block
            parts.append(f'<section class="card-field"><h2 class="st-field">{_text(fields[key])}</h2>{_gap(5)}'
                         f'{html_flowables(field_block)}</section>' + _gap(10))
        else:
            parts.append(html_flowables(block))
    if content["citations"]:
        parts += [_gap(7), '<h2 class="st-field">Fontes</h2>', html_sources(content)]
    footer_css = (f".sheet > footer.meta {{ position: absolute; right: {page['margin_right']}mm; "
                  f"bottom: {page['margin_bottom'] / 2}mm; font-weight: 400; }}\n"
                  f".card-field {{ border: .75pt solid {tokens['color']['border']}; "
                  f"border-left: 3pt solid {tokens['color']['document_accent']}; padding: 10pt 12pt; }}\n")
    parts.append(f'<footer class="meta st-footer">{html.escape(footer_label(content))}</footer>')
    css_text = (base_css(tokens) + style_rules(tokens, pdf_module().style_specs(tokens, "card_portrait"))
                + page_css(tokens, profile["paper"], None) + footer_css)
    return html_document(content, css_text, f'<article class="sheet document">{"".join(parts)}</article>')


def render_presentation_html(content: dict[str, Any], tokens: dict[str, Any]) -> str:
    """Slides 16:9 com os slots do layout.json e a mesma sequência de _render_presentation."""
    slots = load_presentation_layout(content["deliverable_type"]) or {}
    for name in ("logo", "titulo_slide", "corpo_slide", "rodape_slide"):
        if name not in slots or any(key not in slots[name] for key in ("x", "y", "width", "height")):
            raise ValueError(f"layout.json da apresentação sem o slot {name}")
    page = tokens["page"]["presentation_16_9"]
    width_mm, height_mm = float(page["width_in"]) * 25.4, float(page["height_in"]) * 25.4

    def rect(name: str) -> tuple[float, float, float, float]:
        return tuple(float(slots[name][key]) for key in ("x", "y", "width", "height"))  # type: ignore[return-value]

    title_x, title_y, title_w, title_h = rect("titulo_slide")
    body_x, body_y, body_w, body_h = rect("corpo_slide")
    logo_x, logo_y, logo_w, logo_h = rect("logo")
    measure = reportlab_styles(tokens, "presentation")
    if measure is not None:
        from reportlab.platypus import Paragraph

        slide_source = pdf_module()._paragraph_style(tokens, "slide_source", pdf_module().slide_style_specs(tokens)["source"])
        body_size = (body_w * PT_PER_MM, body_h * PT_PER_MM)

        def sources_fit(lines: list[str]) -> bool:
            flow = [Paragraph(pdf_module()._markup(line), slide_source) for line in lines]
            return sum(item.wrap(*body_size)[1] + 12 for item in flow) <= body_size[1]
    else:
        def sources_fit(lines: list[str]) -> bool:
            return len(lines) <= 8
    slides = pdf_module().plan_slides(content, block_text, table_rows, table_views,
                                      body_w * PT_PER_MM, body_h * PT_PER_MM, sources_fit)
    uri, image_w, image_h = _logo(tokens)
    scale = min(min(logo_w, 38) / image_w, min(logo_h, 11) / image_h)
    logo = (f'<img class="slide-logo" src="{uri}" alt="CCEE" style="left: {_number(logo_x + logo_w - image_w * scale)}mm; '
            f'top: {_number(logo_y + (logo_h - image_h * scale) / 2)}mm; width: {_number(image_w * scale)}mm; '
            f'height: {_number(image_h * scale)}mm;">')

    def frame(left: float, top: float, width: float, height: float, inner: str) -> str:
        return (f'<div class="slide-frame" style="left: {_number(left)}mm; top: {_number(top)}mm; '
                f'width: {_number(width)}mm; height: {_number(height)}mm;">{inner}</div>')

    identity = html.escape(identity_line(content))
    footer_x, footer_y, footer_w, footer_h = rect("rodape_slide")
    rendered = []
    for number, (kind, payload) in enumerate(slides, 1):
        parts = [logo]
        if kind == "cover":
            parts.append('<div class="cover-bar"></div>')
            parts.append(frame(title_x, title_y, title_w, title_h + 34,
                               f'<div class="cover-label"><span class="eyebrow">{html.escape(header_label(content))}</span></div>{_gap(16)}'
                               f'<h1 class="st-cover">{_text(content["title"])}</h1>'))
            parts.append(frame(body_x, body_y + 20, body_w, body_h - 40,
                               f'<p class="st-body">{_text(content.get("summary", ""))}</p>'))
        elif kind == "section":
            title, blocks = payload
            parts.append(frame(title_x, title_y, title_w, title_h, f'<h2 class="st-slide_title">{_text(title)}</h2>'))
            items = []
            for block in blocks:
                data = block.get("data", {})
                if block["type"] == "metric":
                    value = data.get("value") or data.get("text") or block_text(block)
                    items.append(f'<p class="st-metric">{_text(value)}{_refs(block.get("citation_ids", []))}</p>')
                    support = data.get("description") or data.get("title")
                    if support:
                        items.append(_gap(8) + f'<p class="st-body">{_text(support)}</p>')
                elif block["type"] == "card_grid":
                    items.append(html_card_grid(block, max_columns=3))
                else:
                    items.append(html_flowables(block))
                items.append(_gap(12))
            parts.append(frame(body_x, body_y, body_w, body_h, "".join(items)))
        else:
            earlier = any(item[0] == "sources" for item in slides[:number - 1])
            title = "Fontes (cont.)" if earlier else "Fontes"
            parts.append(frame(title_x, title_y, title_w, title_h, f'<h2 class="st-slide_title">{title}</h2>'))
            parts.append(frame(body_x, body_y, body_w, body_h,
                               "".join(f'<p class="st-source">{_text(line)}</p>{_gap(12)}' for line in payload)))
        parts.append(f'<div class="slide-footer st-slide_footer" style="left: {_number(footer_x)}mm; top: {_number(footer_y)}mm; '
                     f'width: {_number(footer_w)}mm; height: {_number(footer_h)}mm;">'
                     f'<footer class="meta">{identity}</footer><span>{number} / {len(slides)}</span></div>')
        rendered.append(f'<section class="sheet slide">{"".join(parts)}</section>')
    slide_specs = pdf_module().slide_style_specs(tokens)
    footer_spec = slide_specs.pop("footer")
    color = tokens["color"]
    css_text = (base_css(tokens) + style_rules(tokens, pdf_module().style_specs(tokens, "presentation"))
                + style_rules(tokens, slide_specs, ".slide ") + style_rules(tokens, {"slide_footer": footer_spec})
                + f"@page {{ size: {_number(width_mm)}mm {_number(height_mm)}mm; margin: 0; }}\n"
                f".slide {{ width: {_number(width_mm)}mm; height: {_number(height_mm)}mm; overflow: hidden; break-after: page; }}\n"
                ".slide:last-child { break-after: auto; }\n"
                ".slide-logo, .slide-frame, .slide-footer { position: absolute; }\n"
                f".cover-bar {{ position: absolute; left: 0; top: 0; width: 6mm; height: 100%; background: {color['document_accent']}; }}\n"
                ".slide-footer { display: flex; align-items: center; justify-content: space-between; }\n"
                ".cover-label .eyebrow { display: block; }\n"
                ".slide-footer footer, .slide-footer span { font: inherit; color: inherit; }\n")
    return html_document(content, css_text, "".join(rendered))


def load_presentation_layout(deliverable_type: str) -> dict[str, dict[str, Any]] | None:
    """Lê _DESIGN_SYSTEM/layouts/<tipo>.layout.json e devolve os slots por nome,
    em milímetros (CSS aceita mm nativamente, sem conversão de unidade).
    None se o arquivo não existir."""
    path = ROOT / "layouts" / f"{deliverable_type}.layout.json"
    if not path.is_file():
        return None
    data = json.loads(path.read_text(encoding="utf-8"))
    return {s["name"]: s for s in data.get("slots", []) if s.get("name")}


ONE_PAGE_SLOTS = ("titulo_documento", "identificacao_versao")


def render_one_page_html(content: dict[str, Any], tokens: dict[str, Any], profile: dict[str, Any],
                         logo_uri: str, logo_size: tuple[int, int]) -> str:
    """Resumo de uma página nos slots do layout.json, espelhando _render_one_page_layout do PDF.

    Mesmas posições, recuos, cores, fontes e regras de conteúdo; blocos sem slot são recusados.
    """
    slots = load_presentation_layout(content["deliverable_type"]) or {}
    missing = [name for name in ONE_PAGE_SLOTS if name not in slots]
    if missing:
        raise ValueError("layout.json do resumo sem slots obrigatórios: " + ", ".join(missing))
    mapped_keys = {slot.get("section_key") or name for name, slot in slots.items() if name not in ONE_PAGE_SLOTS}
    unmapped = [block["id"] for block in content["blocks"]
                if block["type"] != "page_break" and block.get("data", {}).get("section_key") not in mapped_keys]
    if unmapped:
        raise ValueError("executive_one_page: blocos sem slot no layout: " + ", ".join(unmapped))
    color = tokens["color"]
    page = tokens["page"][profile["paper"]]
    gap = 1.5

    def cited(value: Any, block: dict[str, Any]) -> str:
        return html.escape(str(value)) + "".join(f"<sup>[{html.escape(x)}]</sup>" for x in block.get("citation_ids", []))

    def box(name: str, inner: str, css_class: str) -> str:
        slot = slots[name]
        x, y, width, height = (float(slot[key]) for key in ("x", "y", "width", "height"))
        if min(x, y, width, height) < 0 or x + width > page["width"] + .1 or y + height > page["height"] + .1:
            raise ValueError(f"layout.json do resumo: slot {name} fora da página")
        if css_class in ("op-hmw", "op-field", "op-people"):
            x, y, width, height = x + gap, y + gap, width - 2 * gap, height - 2 * gap
        return (f'<div class="op-slot {css_class}" data-slot="{html.escape(name)}" style="left: {x}mm; top: {y}mm; '
                f'width: {width}mm; height: {height}mm;">{inner}</div>')

    # Cabeçalho na faixa acima do título, com as mesmas medidas do PDF.
    title_slot = slots["titulo_documento"]
    tx, ty, tw = (float(title_slot[key]) for key in ("x", "y", "width"))
    band_top = 3.0
    available = ty - band_top - 2
    if available < 5:
        raise ValueError("layout.json do resumo: titulo_documento precisa de ao menos 10 mm livres acima")
    image_width, image_height = logo_size
    scale = min(38 / image_width, min(11, available) / image_height)
    logo_width, logo_height = image_width * scale, image_height * scale
    parts = [
        f'<img class="op-logo" src="{logo_uri}" alt="CCEE" style="left: {tx + tw - logo_width}mm; top: {band_top}mm; '
        f'width: {logo_width}mm; height: {logo_height}mm;">',
        f'<span class="eyebrow op-eyebrow" style="left: {tx}mm; top: {band_top}mm; height: {logo_height}mm;">'
        f'{html.escape(header_label(content))}</span>',
        f'<div class="op-rule" style="left: {tx}mm; top: {band_top + logo_height + 1.5}mm; width: {tw}mm;"></div>',
        box("titulo_documento", f'<h1 class="st-title">{html.escape(content["title"])}</h1>', "op-title"),
    ]
    by_key: dict[str, list[dict[str, Any]]] = {}
    for block in content["blocks"]:
        by_key.setdefault(block.get("data", {}).get("section_key"), []).append(block)
    for name, slot in slots.items():
        key = slot.get("section_key") or name
        if name in ONE_PAGE_SLOTS or key in ONE_PAGE_SLOTS or key not in by_key:
            continue
        blocks = by_key[key]
        if key == "hmw":
            parts.append(box(name, "".join(f"<p>{cited(block_text(block), block)}</p>" for block in blocks), "op-hmw"))
            continue
        items = []
        for block in blocks:
            data = block.get("data", {})
            if data.get("title"):
                items.append(f'<h2>{html.escape(data["title"])}</h2>')
            if block["type"] == "table" and data.get("rows"):
                headers, rows = table_rows(block)
                lines = ([" | ".join(headers)] if headers else []) + [" | ".join(row) for row in rows]
                items.extend(f"<p>{html.escape(line)}</p>" for line in lines)
            elif data.get("description") or data.get("text"):
                items.append(f'<p>{html.escape(data.get("description") or data.get("text"))}</p>')
            elif data.get("value") is not None or data.get("content"):
                items.append(f'<p>{html.escape(str(data["value"] if data.get("value") is not None else data["content"]))}</p>')
            elif block["type"] == "table" or data.get("cards") or not (data.get("points") or data.get("items") or data.get("title")):
                items.extend(f"<p>{html.escape(line)}</p>" for line in block_text(block).splitlines() if line.strip())
            for point in data.get("points", []) or data.get("items", []):
                items.append(f"<p>• {html.escape(str(point))}</p>")
            if block.get("citation_ids"):
                items.append(f"<p>{cited('', block)}</p>")
        parts.append(box(name, "".join(items), "op-people" if key == "people" else "op-field"))
    footer = []
    if content["citations"]:
        sources = "; ".join(f"{citation['id']}: {citation['source_id']}"
                            + (f", {citation['locator']}" if citation.get("locator") else "")
                            for citation in content["citations"])
        footer.append(f"<p>{html.escape('Fontes: ' + sources)}</p>")
    footer.append(f"<p>{html.escape(identity_line(content))}</p>")
    parts.append(box("identificacao_versao", "".join(footer), "op-footer"))

    styles = tokens["typography"]["styles"]
    body = tokens["typography"]["font_body"]
    one_page_css = f"""
@page {{ size: {page['width']}mm {page['height']}mm; margin: 0; }}
.one-page {{ width: {page['width']}mm; height: {page['height']}mm; overflow: hidden; }}
.op-slot, .op-logo, .op-eyebrow, .op-rule {{ position: absolute; }}
.op-slot {{ display: flex; flex-direction: column; }}
.op-slot p, .op-slot h1, .op-slot h2 {{ margin: 0; padding: 0; border: 0; }}
.op-eyebrow {{ display: flex; align-items: center; }}
.op-rule {{ border-top: 1.5pt solid {color['document_primary']}; }}
.op-title {{ justify-content: center; }}
.op-hmw {{ justify-content: center; background: {color['document_deep']}; border-left: 3pt solid {color['document_accent']}; padding: 5mm; }}
.op-hmw p {{ font-family: {body}, Arial, sans-serif; font-size: 15pt; line-height: 19pt; font-weight: 700; color: {color['inverse']}; }}
.op-field, .op-people {{ padding: 4mm; }}
.op-field {{ background: {color['document_accent_soft']}; border-left: 3pt solid {color['document_accent']}; }}
.op-people {{ background: {color['surface_muted']}; border-left: 3pt solid {color['border_strong']}; }}
.op-field h2, .op-people h2 {{ font-family: {body}, Arial, sans-serif; font-size: 11pt; line-height: 14pt; font-weight: 700; color: {color['document_primary']}; margin-bottom: 5pt; }}
.op-field p, .op-people p {{ font-size: 10pt; line-height: 13.5pt; color: {color['text_body']}; margin-bottom: 4pt; }}
.op-footer p {{ font-size: {styles['footer']['size']}pt; line-height: {styles['footer']['line_height']}; font-weight: 500; color: {color['text_muted']}; }}
.op-slot sup {{ font-size: 7.5pt; }}
"""
    title = html.escape(content["title"])
    return ("<!doctype html>\n<html lang=\"pt-BR\"><head><meta charset=\"utf-8\">"
            f"<title>{title}</title><style>{base_css(tokens)}{style_rules(tokens, pdf_module().style_specs(tokens, profile['layout']))}"
            f"{one_page_css}</style></head>"
            f'<body><article class="sheet document one-page">{"".join(parts)}</article></body></html>\n')


def render_html(content: dict[str, Any], tokens: dict[str, Any], profile: dict[str, Any]) -> str:
    """HTML com a mesma estrutura, medidas e estilos do PDF de cada layout."""
    logo = ROOT / "assets" / "logo_ccee.png"
    if not logo.is_file():
        raise FileNotFoundError("Logo CCEE ausente; não é possível assegurar timbre idempotente")
    layout = profile["layout"]
    if layout == "executive_one_page":
        logo_bytes = logo.read_bytes()
        uri = "data:image/png;base64," + base64.b64encode(logo_bytes).decode("ascii")
        logo_size = (int.from_bytes(logo_bytes[16:20], "big"), int.from_bytes(logo_bytes[20:24], "big"))
        return render_one_page_html(content, tokens, profile, uri, logo_size)
    if layout == "presentation":
        return render_presentation_html(content, tokens)
    if layout == "card_portrait":
        return render_card_html(content, tokens, profile)
    return render_report_html(content, tokens, profile)


def render_md(content: dict[str, Any]) -> str:
    def grid_cell(text: str) -> str:
        return text.replace("|", "\\|").replace("\r\n", "\n").replace("\r", "\n").replace("\n", "<br>")

    lines = [f"# {content['title']}", "", identity_line(content), ""]
    meta = content.get("metadata", {})
    if content["deliverable_type"] == "learning_card" and meta.get("test_card_id"):
        lines += [f'Test Card: {meta["test_card_id"]} v{meta.get("test_card_version", "")}', ""]
    if content.get("summary"):
        lines += [content["summary"], ""]
    for block in content["blocks"]:
        kind = block["type"]
        card_key = block.get("data", {}).get("section_key")
        card_label = CARD_LABELS.get(content["deliverable_type"], {}).get(card_key)
        if kind == "page_break":
            continue
        if card_label:
            lines += ["## " + card_label, "", block_text(block), ""]
        elif kind == "heading":
            lines += ["## " + block_text(block), ""]
        elif kind == "card_grid":
            if block["data"].get("title"):
                lines += ["### " + str(block["data"]["title"]), ""]
            lines += ["| Card | Conteúdo |", "| --- | --- |"]
            for card in block["data"]["cards"]:
                text = card_text(card)
                if card.get("citation_ids"):
                    text += "\nFontes: " + ", ".join(card["citation_ids"])
                lines.append("| " + grid_cell(str(card["title"])) + " | " + grid_cell(text) + " |")
            lines.append("")
        elif kind == "table":
            for _, legend, headers, rows in table_views(block):
                if legend:
                    lines += ["### " + str(legend), ""]
                if headers:
                    lines += ["| " + " | ".join(value.replace("|", "\\|") for value in headers) + " |",
                              "| " + " | ".join("---" for _ in headers) + " |"]
                lines += ["| " + " | ".join(value.replace("|", "\\|") for value in row) + " |" for row in rows] + [""]
            if score_summary(block):
                lines += [score_summary(block), ""]
        elif kind == "list":
            lines += [f"- {item}" for item in block.get("data", {}).get("items", [])] + [""]
        else:
            lines += [block_text(block), ""]
        if block.get("citation_ids"):
            lines += ["Fontes: " + ", ".join(block["citation_ids"]), ""]
    if content["citations"]:
        lines += ["## Fontes", ""]
        lines += [f"- {c['id']}: {c['source_id']} {c.get('locator', '')}" for c in content["citations"]]
    return "\n".join(lines).rstrip() + "\n"


def render_docx(content: dict[str, Any], tokens: dict[str, Any], profile: dict[str, Any], path: Path) -> None:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Cm, Pt, RGBColor
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError as exc:
        raise RuntimeError("DOCX requer python-docx: pip install python-docx") from exc
    def shade(cell: Any, fill: str) -> None:
        element = OxmlElement("w:shd")
        element.set(qn("w:fill"), fill.lstrip("#"))
        cell._tc.get_or_add_tcPr().append(element)

    def band(paragraph: Any) -> None:
        element = OxmlElement("w:shd")
        element.set(qn("w:fill"), tokens["color"]["document_deep"].lstrip("#"))
        paragraph._p.get_or_add_pPr().append(element)
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor(255, 255, 255)

    doc = Document()
    section = doc.sections[0]
    page = tokens["page"][profile["paper"]]
    section.page_width, section.page_height = Cm(page["width"] / 10), Cm(page["height"] / 10)
    section.top_margin, section.right_margin = Cm(page["margin_top"] / 10), Cm(page["margin_right"] / 10)
    section.bottom_margin, section.left_margin = Cm(page["margin_bottom"] / 10), Cm(page["margin_left"] / 10)
    style = doc.styles["Normal"]
    style.font.name = tokens["typography"]["font_body"]
    style.font.size = Pt(tokens["typography"]["styles"]["body"]["size"])
    for key in ("Title", "Heading 1", "Heading 2"):
        doc.styles[key].font.name = tokens["typography"]["font_brand"]
        doc.styles[key].font.color.rgb = RGBColor.from_string(tokens["color"]["document_primary"].lstrip("#"))
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.add_run(header_label(content).upper() + "    ")
    header.add_run().add_picture(str(ROOT / "assets" / "logo_ccee.png"), width=Cm(3.2))
    doc.add_heading(content["title"], 0)
    identity = doc.add_paragraph(identity_line(content))
    for run in identity.runs:
        run.font.size = Pt(tokens["typography"]["styles"]["footer"]["size"])
        run.font.color.rgb = RGBColor.from_string(tokens["color"]["text_muted"].lstrip("#"))
    meta = content.get("metadata", {})
    if content["deliverable_type"] == "learning_card" and meta.get("test_card_id"):
        doc.add_paragraph(f'Test Card: {meta["test_card_id"]} v{meta.get("test_card_version", "")}')
    if content.get("summary"):
        doc.add_paragraph(content["summary"])
    for block in content["blocks"]:
        kind = block["type"]
        card_key = block.get("data", {}).get("section_key")
        card_label = CARD_LABELS.get(content["deliverable_type"], {}).get(card_key)
        if card_label:
            doc.add_heading(card_label, 1)
        if kind == "heading":
            if card_label:
                doc.add_paragraph(block_text(block))
            else:
                band(doc.add_heading(block_text(block), 1))
        elif kind == "card_grid":
            cards = block["data"]["cards"]
            columns = block["data"].get("columns", 2)
            for offset in range(0, len(cards), columns):
                row_cards = cards[offset:offset + columns]
                table = doc.add_table(rows=1, cols=len(row_cards))
                table.style = "Table Grid"
                for cell, card in zip(table.rows[0].cells, row_cards):
                    cell.text = card_text(card)
                    tones = {"blue": tokens["color"]["document_accent_soft"], "neutral": tokens["color"]["surface_muted"], "red": tokens["color"]["danger_soft"], "green": "#E8F4EE", "yellow": "#FFF4D9"}
                    shade(cell, tones.get(card.get("tone", "blue"), tones["blue"]))
                    if card.get("citation_ids"):
                        cell.add_paragraph("Fontes: " + ", ".join(card["citation_ids"]))
        elif kind == "table":
            for slug, legend, heads, rows in table_views(block):
                if legend:
                    doc.add_paragraph(str(legend))
                width = max(1, len(heads), *(len(row) for row in rows))
                table = doc.add_table(rows=1 if heads else 0, cols=width)
                table.style = "Table Grid"
                if heads:
                    for cell, value in zip(table.rows[0].cells, heads):
                        cell.text = value
                for row in rows:
                    cells = table.add_row().cells
                    for cell, value in zip(cells, row):
                        cell.text = value
                if isinstance(block["data"].get("dimensions"), dict):
                    table.autofit = False
                    available_width = section.page_width - section.left_margin - section.right_margin
                    for column, percentage in zip(table.columns, IQB_TABLE_WIDTHS[slug]):
                        column.width = int(available_width * percentage / 100)
                    for row in table.rows:
                        for cell, column in zip(row.cells, table.columns):
                            cell.width = column.width
                            for paragraph in cell.paragraphs:
                                for run in paragraph.runs:
                                    run.font.size = Pt(8)
            if score_summary(block):
                doc.add_paragraph(score_summary(block))
        elif kind == "list":
            for item in block.get("data", {}).get("items", []):
                doc.add_paragraph(str(item), style="List Bullet")
        elif kind == "page_break":
            doc.add_page_break()
        else:
            if content["deliverable_type"] in CARD_LABELS:
                doc.add_paragraph(block_text(block))
            else:
                table = doc.add_table(rows=1, cols=1)
                table.style = "Table Grid"
                cell = table.rows[0].cells[0]
                add_rich_text(cell.paragraphs[0], block_text(block))
                shade(cell, tokens["color"]["surface_muted"])
        if block.get("citation_ids"):
            doc.add_paragraph("Fontes: " + ", ".join(block["citation_ids"]))
    if content["citations"]:
        doc.add_heading("Fontes", 1)
        for citation in content["citations"]:
            doc.add_paragraph(f"{citation['id']}: {citation['source_id']} {citation.get('locator', '')}")
    doc.core_properties.title = content["title"]
    doc.core_properties.author = "my_skills"
    doc.save(path)


def load_pptx_layout(deliverable_type: str) -> dict[str, Any] | None:
    """Lê _DESIGN_SYSTEM/layouts/<tipo>.layout.json, se existir, e converte mm -> polegadas.
    Retorna None quando o arquivo não existe (render_pptx usa as posições hardcoded nesse caso)."""
    path = ROOT / "layouts" / f"{deliverable_type}.layout.json"
    if not path.is_file():
        return None
    data = json.loads(path.read_text(encoding="utf-8"))
    mm_to_in = 1 / 25.4
    canvas = data.get("canvas", {})
    slots = {}
    for slot in data.get("slots", []):
        name = slot.get("name")
        if not name:
            continue
        slots[name] = {k: slot[k] * mm_to_in for k in ("x", "y", "width", "height") if k in slot}
    return {
        "width_in": canvas.get("width", 0) * mm_to_in,
        "height_in": canvas.get("height", 0) * mm_to_in,
        "slots": slots,
    }


def render_pptx(content: dict[str, Any], tokens: dict[str, Any], path: Path) -> None:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Mm, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
    except ImportError as exc:
        raise RuntimeError("PPTX requer python-pptx: pip install python-pptx") from exc
    pdf = pdf_module()
    layout = load_pptx_layout(content["deliverable_type"])

    def box(name: str, fx: float, fy: float, fw: float, fh: float) -> tuple[Any, Any, Any, Any]:
        slot = layout["slots"].get(name) if layout else None
        if slot and all(k in slot for k in ("x", "y", "width", "height")):
            return Inches(slot["x"]), Inches(slot["y"]), Inches(slot["width"]), Inches(slot["height"])
        return Inches(fx), Inches(fy), Inches(fw), Inches(fh)

    prs = Presentation()
    # O modelo padrão usa meia-risca como marcador dos níveis 2 e 4; os entregáveis não usam travessão.
    for master in prs.slide_masters:
        for bullet in master.element.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar"):
            if DASH.search(bullet.get("char", "")):
                bullet.set("char", "•")
    if layout:
        prs.slide_width, prs.slide_height = Inches(layout["width_in"]), Inches(layout["height_in"])
    else:
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    primary = RGBColor.from_string(tokens["color"]["document_primary"].lstrip("#"))
    logo = ROOT / "assets" / "logo_ccee.png"

    def color(name: str) -> Any:
        return RGBColor.from_string(tokens["color"][name].lstrip("#"))

    table_size = Pt(max(12, min(14, tokens["typography"]["styles"]["slide_body"]["size"])))

    def add_table(s: Any, table: tuple[str, list[str], list[list[str]], list[float]], x: Any, y: Any, width: Any, height: Any) -> None:
        caption, heads, rows, heights = table
        top = y
        if caption:
            caption_box = s.shapes.add_textbox(x, y, width, Inches(.45))
            caption_box.text_frame.word_wrap = True
            run = caption_box.text_frame.paragraphs[0].add_run()
            run.text = caption
            run.font.name, run.font.size, run.font.bold = tokens["typography"]["font_brand"], table_size, True
            run.font.color.rgb = primary
            top = y + Inches(.5)
        columns = max(1, len(heads), *(len(row) for row in rows)) if rows else max(1, len(heads))
        frame = s.shapes.add_table(len(rows) + 1, columns, x, top, width, Pt(sum(heights)))
        grid = frame.table
        for index, row_height in enumerate(heights):
            grid.rows[index].height = Pt(row_height)
        for index in range(columns):
            grid.columns[index].width = int(width / columns)
        for row_index, values in enumerate([heads, *rows]):
            for column in range(columns):
                cell = grid.cell(row_index, column)
                cell.text = str(values[column]) if column < len(values) else ""
                cell.fill.solid()
                cell.fill.fore_color.rgb = color("document_primary") if row_index == 0 else (
                    color("surface") if row_index % 2 else color("surface_muted"))
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.font.name = tokens["typography"]["font_body"]
                    paragraph.font.size = table_size
                    paragraph.font.bold = row_index == 0
                    paragraph.font.color.rgb = color("inverse") if row_index == 0 else color("text_body")

    def slide(title: str, body: str = "", table: tuple[str, list[str], list[list[str]], list[float]] | None = None) -> None:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        # O slot "logo" é a faixa do cabeçalho: como no HTML e no PDF, o logo cabe no
        # slot e em 38 x 11 mm, alinhado à direita e centrado na altura.
        logo_x, logo_y, logo_w, logo_h = box("logo", 11.1, 0.35, 1.7, 0)
        picture = s.shapes.add_picture(str(logo), logo_x, logo_y)
        max_w = min(logo_w, Mm(38))
        max_h = min(logo_h, Mm(11)) if logo_h else Mm(11)
        scale = min(max_w / picture.width, max_h / picture.height)
        picture.width, picture.height = int(picture.width * scale), int(picture.height * scale)
        picture.left = logo_x + logo_w - picture.width
        picture.top = logo_y + (logo_h - picture.height) // 2 if logo_h else logo_y
        title_x, title_y, title_w, title_h = box("titulo_slide", .5, .5, 10.3, .8)
        title_box = s.shapes.add_textbox(title_x, title_y, title_w, title_h)
        run = title_box.text_frame.paragraphs[0].add_run()
        run.text, run.font.name, run.font.size, run.font.color.rgb = title, tokens["typography"]["font_brand"], Pt(tokens["typography"]["styles"]["slide_title"]["size"]), primary
        body_x, body_y, body_w, body_h = box("corpo_slide", .6, 1.55, 12.1, 5.35)
        if table is not None:
            add_table(s, table, body_x, body_y, body_w, body_h)
        else:
            body_box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, body_x, body_y, body_w, body_h)
            body_box.fill.solid()
            body_box.fill.fore_color.rgb = color("document_accent_soft")
            body_box.line.color.rgb = color("border")
            body_box.text_frame.margin_left = Inches(.2)
            body_box.text_frame.margin_top = Inches(.2)
            body_box.text_frame.word_wrap = True
            for line in body.splitlines() or [""]:
                para = body_box.text_frame.paragraphs[0] if not body_box.text_frame.text else body_box.text_frame.add_paragraph()
                para.text = line
                para.font.name, para.font.size = tokens["typography"]["font_body"], Pt(tokens["typography"]["styles"]["slide_body"]["size"])
                # Sem cor explícita, o texto de uma forma herda o branco do tema e some no fundo claro.
                para.font.color.rgb = color("text_body")
        footer_x, footer_y, footer_w, footer_h = box("rodape_slide", .6, 7.02, 12.1, .33)
        footer = s.shapes.add_textbox(footer_x, footer_y, footer_w, footer_h)
        footer.text_frame.word_wrap = True
        paragraph = footer.text_frame.paragraphs[0]
        paragraph.text = identity_line(content)
        paragraph.font.name = tokens["typography"]["font_body"]
        paragraph.font.size = Pt(tokens["typography"]["styles"]["slide_footer"]["size"])
        paragraph.font.color.rgb = RGBColor.from_string(tokens["color"]["text_muted"].lstrip("#"))
        eyebrow = s.shapes.add_textbox(logo_x, logo_y, logo_w - Mm(42), logo_h or Mm(11))
        eyebrow_run = eyebrow.text_frame.paragraphs[0].add_run()
        eyebrow_run.text = header_label(content).upper()
        eyebrow_run.font.name = tokens["typography"]["font_body"]
        eyebrow_run.font.size = Pt(tokens["typography"]["styles"]["cover_eyebrow"]["size"])
        eyebrow_run.font.color.rgb = color("document_accent")

    slide(content["title"], content.get("summary", ""))
    title = "Conteúdo"
    lines: list[str] = []
    for block in content["blocks"]:
        if block["type"] == "heading" or block["type"] == "page_break":
            if lines:
                slide(title, "\n".join(lines))
                lines = []
            if block["type"] == "heading":
                title = block_text(block)
            continue
        if block["type"] == "table":
            # Tabela nativa: o texto acumulado sai antes; até 8 linhas por lâmina, cabeçalho repetido.
            if lines:
                slide(title, "\n".join(lines))
                lines = []
            refs = " [" + ", ".join(block["citation_ids"]) + "]" if block.get("citation_ids") else ""
            for _, legend, heads, rows in table_views(block):
                body_x, body_y, body_w, body_h = box("corpo_slide", .6, 1.55, 12.1, 5.35)
                chunks = pdf.paginate_table_rows(heads, rows, body_w / 12700, body_h / 12700)
                for chunk, heights in chunks:
                    slide(title, table=((str(legend) if legend else "") + refs, heads, chunk, heights))
                    if not title.endswith("(continuação)"):
                        title += " (continuação)"
            continue
        line = block_text(block)
        if block.get("citation_ids"):
            line += " [" + ", ".join(block["citation_ids"]) + "]"
        # Divide em limites de palavra: nenhuma palavra fica partida entre lâminas.
        chunks = [piece for paragraph in line.split("\n")
                  for piece in (textwrap.wrap(paragraph, 650, break_long_words=False, break_on_hyphens=False) or [""])]
        for chunk in chunks:
            if lines and (sum(len(x) for x in lines) + len(chunk) > 650 or len(lines) >= 8):
                slide(title, "\n".join(lines))
                if not title.endswith("(continuação)"):
                    title += " (continuação)"
                lines = []
            lines.append(chunk)
    if lines:
        slide(title, "\n".join(lines))
    if content["citations"]:
        slide("Fontes", "\n".join(f"{c['id']}: {c['source_id']}" + (f", {c['locator']}" if c.get('locator') else '')
                                     for c in content["citations"]))
    prs.save(path)


def write_csv_tables(content: dict[str, Any], out: Path, stem: str) -> list[Path]:
    written = []
    for block in content["blocks"]:
        export = block.get("export", {})
        if block["type"] != "table" or not export.get("csv"):
            continue
        variant = safe_name(export["variant_key"])
        for index, (slug, _, headers, rows) in enumerate(table_views(block)):
            suffix = f"-{slug}" if index else ""
            path = out / f"{stem}-{variant}{suffix}.csv"
            with path.open("w", encoding="utf-8", newline="") as handle:
                writer = csv.writer(handle)
                if headers:
                    writer.writerow(headers)
                writer.writerows(rows)
            written.append(path)
    if not written:
        raise ValueError("CSV solicitado sem bloco table marcado export.csv=true")
    return written


def check_pdf_quality(path: Path, tokens: dict[str, Any]) -> list[str]:
    import importlib.util

    spec = importlib.util.spec_from_file_location("pdf_quality", ROOT / "scripts/pdf_quality.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module.check_pdf(path, min_size=float(tokens["density"]["minimum_font_size"]))


def render_pdf_reportlab(content: dict[str, Any], tokens: dict[str, Any],
                         profile: dict[str, Any], path: Path) -> None:
    pdf_module().render_pdf(content, tokens, profile, path,
                      deliverable_label=header_label(content),
                      footer_label=footer_label(content),
                      field_labels=CARD_LABELS.get(content["deliverable_type"]),
                      block_text=block_text, table_rows=table_rows,
                      table_views=table_views, iqb_widths=IQB_TABLE_WIDTHS,
                      load_presentation_layout=load_presentation_layout,
                      identity_line=identity_line)


# Na CDT, o que vai ao Jira não mostra fonte (o Jira não tem como referenciá-la); as fontes aparecem só nas seções
# que ficam fora do Jira. O JSON canônico mantém as citações, que o checador usa.
CDT_SECTIONS_WITH_SOURCES = ("readiness", "publication_checklist", "quality_review", "data_to_obtain", "traceability")


def presentation_content(content: dict[str, Any]) -> dict[str, Any]:
    if content.get("deliverable_type") != "cdt":
        return content
    shown = copy.deepcopy(content)
    for block in shown.get("blocks", []):
        if block.get("data", {}).get("section_key") not in CDT_SECTIONS_WITH_SOURCES:
            block.pop("citation_ids", None)
    return shown


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("--out-dir", required=True, type=Path)
    parser.add_argument("--formats", help="json,pdf,docx,md,txt; pptx só para apresentação; csv só para tabelas exportáveis")
    parser.add_argument("--pdf-engine", choices=("reportlab",), default="reportlab",
                        help="motor do PDF (só reportlab; opção mantida por compatibilidade)")
    args = parser.parse_args()
    content = read_json(args.input)
    tokens = read_json(ROOT / "tokens.json")
    profiles = read_json(ROOT / "profiles.json")
    profile = validate(content, profiles)
    allowed = profile.get("formats", profiles["default_formats"]) + ["html", "csv"]
    requested = args.formats.split(",") if args.formats else profile.get("formats", profiles["default_formats"])
    requested = list(dict.fromkeys(["json", *requested]))
    invalid = set(requested) - set(allowed)
    if invalid:
        raise ValueError(f"Formatos incompatíveis com {content['deliverable_type']}: {sorted(invalid)}")
    out = args.out_dir
    out.mkdir(parents=True, exist_ok=True)
    stem = safe_name(content["metadata"]["artifact_id"]) + "-v" + safe_name(content["metadata"]["content_version"])
    files: list[Path] = []
    canonical = canonical_bytes(content)
    path = out / f"{stem}.json"
    path.write_bytes(canonical)
    files.append(path)
    content = presentation_content(content)
    source_html: str | None = None
    if "html" in requested:
        source_html = render_html(content, tokens, profile)
    if "html" in requested:
        path = out / f"{stem}.html"
        path.write_text(source_html or "", encoding="utf-8")
        files.append(path)
    unavailable: dict[str, str] = {}
    if "pdf" in requested:
        path = out / f"{stem}.pdf"
        try:
            render_pdf_reportlab(content, tokens, profile, path)
        except ValueError:
            # Conteúdo que não cabe (PdfOverflow) continua fatal: corrija o conteúdo.
            path.unlink(missing_ok=True)
            raise
        except Exception as exc:
            import sys

            # Dependência ausente ou falha inesperada do motor: entrega os demais formatos.
            path.unlink(missing_ok=True)
            unavailable["pdf"] = f"{type(exc).__name__}: {exc}"[:300]
            print(f"AVISO: PDF não gerado ({unavailable['pdf']}). Instale: pip install -r requirements.txt.",
                  file=sys.stderr)
        else:
            if profile["layout"] in ("executive_one_page", "card_portrait"):
                try:
                    from pypdf import PdfReader
                except ImportError as exc:
                    raise RuntimeError("Verificação de página única requer pypdf: pip install pypdf") from exc
                count = len(PdfReader(str(path)).pages)
                if count != 1:
                    path.unlink(missing_ok=True)
                    raise ValueError(f"{profile['layout']} excedeu uma página ({count}); revise ou divida o conteúdo sem reduzir a fonte")
            try:
                problems = check_pdf_quality(path, tokens)
            except ImportError as exc:
                path.unlink(missing_ok=True)
                raise RuntimeError("Verificação de qualidade do PDF requer pdfminer.six: pip install pdfminer.six") from exc
            if problems:
                path.unlink(missing_ok=True)
                raise ValueError("PDF reprovado na verificação de qualidade: " + "; ".join(problems[:5])
                                 + ". Divida o conteúdo ou mude a variante; não reduza a fonte")
            files.append(path)
    if "md" in requested or "txt" in requested:
        markdown = render_md(content)
        if "md" in requested:
            path = out / f"{stem}.md"
            path.write_text(markdown, encoding="utf-8")
            files.append(path)
        if "txt" in requested:
            path = out / f"{stem}.txt"
            path.write_text(re.sub(r"^#{1,6} ", "", markdown, flags=re.MULTILINE), encoding="utf-8")
            files.append(path)
    if "docx" in requested:
        path = out / f"{stem}.docx"
        render_docx(content, tokens, profile, path)
        files.append(path)
    if "pptx" in requested:
        path = out / f"{stem}.pptx"
        render_pptx(content, tokens, path)
        files.append(path)
    if "csv" in requested:
        files.extend(write_csv_tables(content, out, stem))
    manifest = {
        "artifact_id": content["metadata"]["artifact_id"],
        "content_version": content["metadata"]["content_version"],
        "deliverable_type": content["deliverable_type"],
        "layout": profile["layout"],
        "paper": profile["paper"],
        "renderer_version": RENDERER_VERSION,
        "profile_version": profiles["profile_version"],
        "branding_version": tokens["profile_version"],
        "content_sha256": sha(canonical),
        "profile_sha256": sha(canonical_bytes(profiles)),
        "tokens_sha256": sha(canonical_bytes(tokens)),
        "files": {path.name: sha(path.read_bytes()) for path in files},
        "unavailable": unavailable,
    }
    if any(file.suffix == ".pdf" for file in files):
        manifest["pdf_engine"] = args.pdf_engine
    (out / f"{stem}.manifest.json").write_bytes(canonical_bytes(manifest))
    print(json.dumps({"manifest": f"{stem}.manifest.json", "files": list(manifest["files"]),
                      "unavailable": unavailable}, ensure_ascii=False))


if __name__ == "__main__":
    main()
