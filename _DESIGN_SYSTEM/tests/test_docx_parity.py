"""Proofs for editable DOCX visual parity against the 19 sample PDFs."""
from __future__ import annotations

import hashlib
import importlib.util
import json
import copy
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from docx import Document
from docx.oxml.ns import qn
from PIL import Image
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[2]
SAMPLES = ROOT / "_SAMPLES_ENTREGAVEIS"
COMPARATOR = ROOT / "_TOOLS/check_docx_pdf_parity.py"
RENDER = ROOT / "_DESIGN_SYSTEM/scripts/render.py"


def load_module(path: Path, name: str):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def normalized(value: str) -> str:
    return " ".join(value.replace("**", "").split()).casefold()


def docx_text(path: Path) -> str:
    document = Document(path)
    return " ".join(node.text or "" for node in document.element.body.xpath(".//w:t"))


def text_atoms(data: object):
    if isinstance(data, dict):
        for key, value in data.items():
            if key in {"text", "description", "title", "label", "content", "justification", "evidence"}:
                if isinstance(value, str) and len(value.strip()) >= 8:
                    for line in value.splitlines():
                        if len(line.strip()) >= 8:
                            yield line
            elif isinstance(value, (dict, list)):
                yield from text_atoms(value)
    elif isinstance(data, list):
        for value in data:
            yield from text_atoms(value)


@unittest.skipUnless(shutil.which("soffice") and shutil.which("pdftoppm"),
                     "comparação visual requer LibreOffice e pdftoppm")
class DocxParityTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        digest = hashlib.sha256()
        for path in [RENDER, COMPARATOR, *sorted(SAMPLES.glob("*/input.json")),
                     *sorted(SAMPLES.glob("*/*.pdf"))]:
            digest.update(path.read_bytes())
        cls.cache = Path(tempfile.gettempdir()) / "eric-docx-parity-tests" / digest.hexdigest()[:16]
        report_path = cls.cache / "report.json"
        if not report_path.is_file():
            result = subprocess.run([sys.executable, str(COMPARATOR), "--out-dir", str(cls.cache)],
                                    capture_output=True, text=True)
            if result.returncode:
                raise AssertionError(result.stdout + "\n" + result.stderr)
        cls.report = json.loads(report_path.read_text(encoding="utf-8"))
        cls.rows = {item["type"]: item for item in cls.report["results"]}
        cls.renderer = load_module(RENDER, "render_for_docx_parity")
        cls.tokens = json.loads((ROOT / "_DESIGN_SYSTEM/tokens.json").read_text(encoding="utf-8"))
        cls.profiles = json.loads((ROOT / "_DESIGN_SYSTEM/profiles.json").read_text(encoding="utf-8"))

    def test_all_sample_page_counts(self):
        self.assertEqual(self.report["sample_count"], 19)
        for kind, row in self.rows.items():
            with self.subTest(kind=kind):
                self.assertEqual(row["reference_pages"], row["docx_pages"])
                self.assertEqual(row["failures"], [])
                sizes = row["docx_sizes_pt"]
                paper = self.profiles["types"][kind]["paper"]
                self.assertTrue(all((width > height) == (paper == "a4_landscape")
                                    for width, height in sizes))

    def test_one_page_briefing_geometry(self):
        row = self.rows["briefing_one_page"]
        self.assertEqual(row["docx_pages"], 1)
        self.assertGreaterEqual(row["pages"][0]["blue_region_iou"], .65)
        path = next((self.cache / "generated/briefing_one_page").glob("*.docx"))
        document = Document(path)
        table = document.tables[0]
        self.assertEqual(len(table.rows), 5)
        self.assertEqual(len(table.columns), 5)
        self.assertEqual(round(table.rows[2].height.cm, 1), 3.2)
        card_row = table.rows[3]
        for position, label in ((0, "Problema ou oportunidade"), (2, "Relevância"),
                                (4, "Resultados esperados")):
            self.assertIn(label, card_row.cells[position].text)
        self.assertIn("Responsáveis", table.rows[4].cells[0].text)
        self.assertIn("ID sample-briefing-one-page", "\n".join(p.text for p in document.paragraphs))
        sample = json.loads((SAMPLES / "briefing_one_page/input.json").read_text(encoding="utf-8"))
        profile = self.renderer.validate(sample, self.profiles)
        changed = copy.deepcopy(self.renderer.load_presentation_layout("briefing_one_page"))
        changed["hmw"]["height"] = 35
        with tempfile.TemporaryDirectory() as temporary, mock.patch.object(
                self.renderer, "load_presentation_layout", return_value=changed):
            adjusted = Path(temporary) / "changed-layout.docx"
            self.renderer.render_docx(sample, self.tokens, profile, adjusted)
            self.assertEqual(round(Document(adjusted).tables[0].rows[2].height.cm, 1), 3.5)

    def test_report_structure_and_safe_area(self):
        card_types = {"test_card", "learning_card", "briefing_one_page"}
        for kind, row in self.rows.items():
            if kind in card_types:
                continue
            with self.subTest(kind=kind):
                self.assertEqual(row["clipped_characters"], 0)
                self.assertGreaterEqual(row["pages"][0]["blue_region_iou"], .25)
                for page_number, page in enumerate(row["pages"], 1):
                    self.assertLessEqual(page["mean_rgb_difference"], 42.0,
                                         f"{kind} página {page_number}: composição distante do PDF")
                path = next((self.cache / "generated" / kind).glob("*.docx"))
                document = Document(path)
                body_xml = document.element.body.xml
                self.assertIn(self.tokens["color"]["document_deep"].lstrip("#"), body_xml)
                self.assertIn(self.tokens["color"]["document_primary"].lstrip("#"), body_xml)
        reference = PdfReader(next((SAMPLES / "four_d_report").glob("sample-*.pdf")))
        converted = PdfReader(next((self.cache / "converted").glob("sample-four-d-report*.pdf")))
        for section in ("JORNADA DO PROCESSO", "ALTERNATIVAS", "PRIORIZAÇÃO", "CONCLUSÃO E PRÓXIMOS PASSOS"):
            def pages_with_heading(reader):
                return [number for number, page in enumerate(reader.pages, 1)
                        if section in (page.extract_text() or "").splitlines()]
            self.assertEqual(pages_with_heading(reference), pages_with_heading(converted), section)

    def test_cards_one_page_and_order(self):
        for kind in ("test_card", "learning_card"):
            with self.subTest(kind=kind):
                row = self.rows[kind]
                self.assertEqual(row["docx_pages"], 1)
                self.assertEqual(row["clipped_characters"], 0)
                path = next((self.cache / "generated" / kind).glob("*.docx"))
                document = Document(path)
                labels = list(self.renderer.CARD_LABELS[kind].values())
                text = "\n".join(p.text for p in document.paragraphs)
                positions = [text.index(label) for label in labels]
                self.assertEqual(positions, sorted(positions))
                self.assertGreaterEqual(document.element.body.xml.count('w:val="single"'), 12)
                self.assertGreaterEqual(document.element.body.xml.count('w:color="4C73D1"'), 4)

    def test_all_samples_editable(self):
        for kind in self.rows:
            with self.subTest(kind=kind):
                path = next((self.cache / "generated" / kind).glob("*.docx"))
                document = Document(path)
                text_nodes = document.element.body.xpath(".//w:t")
                images = document.element.body.xpath(".//a:blip")
                self.assertGreater(len(text_nodes), 8)
                self.assertLessEqual(len(images), 1, "só o logo pode ser imagem")
                self.assertIn("EXEMPLO FICTÍCIO", docx_text(path).upper())

    def test_all_sample_logos_uncropped(self):
        source = Image.open(ROOT / "_DESIGN_SYSTEM/assets/logo_ccee.png").convert("RGB")
        def logo_pixels(image, x_start=0, y_end=None):
            y_end = y_end or image.height
            return [(x, y) for y in range(y_end) for x in range(x_start, image.width)
                    if (lambda color: color[0] < 30 and color[1] < 30 and color[2] > 80)(image.getpixel((x, y)))]
        points = logo_pixels(source)
        expected_ratio = (max(x for x, _ in points) - min(x for x, _ in points) + 1) / (
            max(y for _, y in points) - min(y for _, y in points) + 1)
        for kind in self.rows:
            with self.subTest(kind=kind):
                image = Image.open(self.cache / "pictures" / f"{kind}-1-docx.png").convert("RGB")
                points = logo_pixels(image, int(image.width * .58), int(image.height * .15))
                self.assertGreater(len(points), 50, "logo ausente do cabeçalho")
                width = max(x for x, _ in points) - min(x for x, _ in points) + 1
                height = max(y for _, y in points) - min(y for _, y in points) + 1
                self.assertLessEqual(abs(width / height / expected_ratio - 1), .12,
                                     "proporção do logo indica corte no DOCX convertido")

    def test_all_section_bands_match_pdf_padding(self):
        for kind in self.rows:
            if kind in {"test_card", "learning_card", "briefing_one_page"}:
                continue
            with self.subTest(kind=kind):
                path = next((self.cache / "generated" / kind).glob("*.docx"))
                document = Document(path)
                bands = [paragraph for paragraph in document.element.body.xpath(".//w:p")
                         if any(shade.get(qn("w:fill")) == self.tokens["color"]["document_deep"].lstrip("#")
                                for shade in paragraph.xpath("./w:pPr/w:shd"))]
                self.assertTrue(bands, "faixa de título ausente")
                inset = "180" if kind == "innovation_potential_assessment" else "240"
                vertical = "40" if kind == "innovation_potential_assessment" else "72"
                for paragraph in bands:
                    self.assertTrue(any(tab.get(qn("w:pos")) == inset
                                        for tab in paragraph.xpath("./w:pPr/w:tabs/w:tab")),
                                    "faixa sem recuo horizontal do PDF")
                    for edge in ("top", "bottom"):
                        borders = paragraph.xpath(f"./w:pPr/w:pBdr/w:{edge}")
                        self.assertEqual(len(borders), 1, f"padding vertical {edge} ausente")
                        self.assertEqual(borders[0].get(qn("w:sz")), vertical)
                    self.assertTrue(paragraph.xpath("./w:r[1]/w:tab"), "texto da faixa não começa no recuo")

        def first_band(image_path):
            image = Image.open(image_path).convert("RGB")
            width, height = image.size
            rows = []
            for y in range(int(height * .15), int(height * .45)):
                dark = [x for x in range(width) if (
                    lambda color: color[0] < 40 and color[1] < 60 and 60 < color[2] < 140
                )(image.getpixel((x, y)))]
                if len(dark) > width * .6:
                    rows.append((y, min(dark), max(dark)))
            self.assertTrue(rows, "faixa azul de título não encontrada")
            top = rows[0][0]
            band_rows = [row for row in rows if top <= row[0] <= top + 80]
            bottom = band_rows[-1][0]
            left = min(row[1] for row in band_rows)
            right = max(row[2] for row in band_rows)
            white = [(x, y) for y in range(top, bottom + 1) for x in range(left, right + 1)
                     if all(channel > 180 for channel in image.getpixel((x, y)))]
            self.assertTrue(white, "rótulo branco não encontrado")
            text_top = min(y for _, y in white)
            text_bottom = max(y for _, y in white)
            return (min(x for x, _ in white) - left, bottom - top + 1,
                    (text_top + text_bottom - top - bottom) / 2)

        for kind in ("cdt", "briefing_validation"):
            with self.subTest(kind=kind):
                reference = first_band(self.cache / "pictures" / f"{kind}-1-reference.png")
                editable = first_band(self.cache / "pictures" / f"{kind}-1-docx.png")
                self.assertGreaterEqual(editable[0], 12, "título encostado na borda azul")
                self.assertLessEqual(abs(reference[0] - editable[0]), 4, "recuo do título difere do PDF")
                self.assertLessEqual(abs(reference[1] - editable[1]), 5, "altura da faixa difere do PDF")
                self.assertLessEqual(abs(editable[2]), 3, "título não está centralizado verticalmente")

    def test_all_samples_semantic_parity(self):
        for kind in self.rows:
            with self.subTest(kind=kind):
                source = json.loads((SAMPLES / kind / "input.json").read_text(encoding="utf-8"))
                path = next((self.cache / "generated" / kind).glob("*.docx"))
                pdf = next((SAMPLES / kind).glob("sample-*.pdf"))
                texts = (normalized(docx_text(path)), normalized(" ".join(p.extract_text() or "" for p in PdfReader(pdf).pages)))
                expected_state = self.renderer.COMPLETENESS_LABELS[source["metadata"].get("completeness") or "complete"]
                for target in texts:
                    self.assertIn(normalized(source["title"]), target)
                    self.assertIn(normalized(source["metadata"]["artifact_id"]), target)
                    self.assertIn(normalized(source["metadata"]["content_version"]), target)
                    self.assertIn(normalized(expected_state), target)
                    if kind != "cdt":
                        for citation in source["citations"]:
                            self.assertIn(normalized(citation["source_id"]), target)
                    for atom in text_atoms(source["blocks"]):
                        self.assertIn(normalized(atom), target, f"{kind}: {atom[:60]}")
                for atom in text_atoms(source["blocks"]):
                    if len(normalized(atom)) >= 15:
                        self.assertEqual(texts[0].count(normalized(atom)), texts[1].count(normalized(atom)),
                                         f"{kind}: duplicação ou omissão de {atom[:60]}")

    def test_missing_font_fails_explicitly(self):
        sample = json.loads((SAMPLES / "test_card/input.json").read_text(encoding="utf-8"))
        profile = self.renderer.validate(sample, self.profiles)
        with tempfile.TemporaryDirectory() as temporary, mock.patch.object(self.renderer, "ROOT", Path(temporary)):
            with self.assertRaisesRegex(FileNotFoundError, "DOCX requer a fonte"):
                self.renderer.render_docx(sample, self.tokens, profile, Path(temporary) / "out.docx")
        with mock.patch.dict(sys.modules, {"docx": None}):
            with self.assertRaisesRegex(RuntimeError, "DOCX requer python-docx"):
                self.renderer.render_docx(sample, self.tokens, profile, Path(tempfile.gettempdir()) / "unused.docx")

    def test_comparator_artifacts_and_failures(self):
        comparator = load_module(COMPARATOR, "docx_parity_comparator")
        self.assertEqual(len(self.rows), 19)
        for kind, row in self.rows.items():
            for page in range(1, row["docx_pages"] + 1):
                self.assertTrue((self.cache / "pictures" / f"{kind}-{page}-side-by-side.png").is_file())
                self.assertTrue((self.cache / "pictures" / f"{kind}-{page}-side-by-side-diff.png").is_file())
        self.assertIn("page_count", comparator.comparison_failures([[595, 842]], [[595, 842], [595, 842]], 0))
        self.assertNotIn("page_size", comparator.comparison_failures([[595.28, 841.89]], [[595.48, 841.89]], 0))
        self.assertIn("page_size", comparator.comparison_failures([[595.28, 841.89]], [[595.49, 841.89]], 0))
        self.assertIn("clipped_text", comparator.comparison_failures([[595, 842]], [[595, 842]], 1))


if __name__ == "__main__":
    unittest.main()
