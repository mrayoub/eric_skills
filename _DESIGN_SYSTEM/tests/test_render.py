from __future__ import annotations

import json
import copy
import importlib.util
import os
import re
import subprocess
import sys
import tempfile
import unittest
import unittest.mock
from pathlib import Path

from docx import Document
from pypdf import PdfReader
from pptx import Presentation
from jsonschema import Draft202012Validator
from referencing import Registry, Resource

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "render.py"


def load_render_module():
    spec = importlib.util.spec_from_file_location("render_module", SCRIPT)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def pdf_available() -> bool:
    try:
        import weasyprint
        return True
    except (ImportError, OSError):
        return False


def reportlab_available() -> bool:
    try:
        import reportlab
        return True
    except ImportError:
        return False


def content(kind: str, artifact: str) -> dict:
    blocks = [
        {"id": "h-1", "type": "heading", "data": {"text": "Seção principal", "level": 2}},
        {"id": "p-1", "type": "paragraph", "data": {"text": "Evidência resumida e verificável."}, "citation_ids": ["F01"]},
        {"id": "t-1", "type": "table", "data": {"headers": ["Medida", "Valor"], "rows": [["Casos", "12"]]}, "export": {"csv": True, "variant_key": "medidas"}},
    ]
    if kind in ("test_card", "learning_card"):
        keys = ("hypothesis", "test", "measure", "criteria") if kind == "test_card" else ("hypothesis", "observations", "learning", "action")
        blocks = [{"id": key, "type": "paragraph", "data": {
            "section_key": key, "text": "Evidência resumida e verificável."}} for key in keys]
    return {
        "schema_version": "1.0.0", "deliverable_type": kind, "locale": "pt-BR",
        "title": "Avaliação de exemplo", "summary": "Resumo curto.", "blocks": blocks,
        "citations": [{"id": "F01", "source_id": "documento-1", "locator": "p. 2"}],
        "metadata": {"artifact_id": artifact, "content_version": "1.0.0", "completeness": "partial"},
    }


class TableRowsTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        spec = importlib.util.spec_from_file_location("render", SCRIPT)
        cls.render = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(cls.render)

    def test_column_key_differs_from_label(self):
        self.assertEqual(self.render.table_rows({"data": {
            "columns": [{"key": "count", "label": "Quantidade"}],
            "rows": [{"count": 12}],
        }}), (["Quantidade"], [["12"]]))

    def test_reordered_columns(self):
        self.assertEqual(self.render.table_rows({"data": {
            "columns": [{"key": "b", "label": "B"}, {"key": "a", "label": "A"}],
            "rows": [{"a": 1, "b": 2}],
        }}), (["B", "A"], [["2", "1"]]))

    def test_zero_and_false_are_preserved(self):
        for value, expected in [(0, "0"), (False, "False")]:
            with self.subTest(value=value):
                self.assertEqual(self.render.table_rows({"data": {
                    "columns": [{"key": "count", "label": "Quantidade"}],
                    "rows": [{"count": value}],
                }}), (["Quantidade"], [[expected]]))

    def test_missing_field_is_empty(self):
        self.assertEqual(self.render.table_rows({"data": {
            "columns": [{"key": "count", "label": "Quantidade"}], "rows": [{}],
        }}), (["Quantidade"], [[""]]))

    def test_simple_headers_and_list_rows(self):
        self.assertEqual(self.render.table_rows({"data": {
            "headers": ["Medida", "Valor"], "rows": [["Casos", "12"]],
        }}), (["Medida", "Valor"], [["Casos", "12"]]))

    def test_string_headers_with_object_rows(self):
        self.assertEqual(self.render.table_rows({"data": {
            "headers": ["x"], "rows": [{"x": "v"}],
        }}), (["x"], [["v"]]))

    def test_column_without_label(self):
        self.assertEqual(self.render.table_rows({"data": {
            "columns": [{"key": "k"}], "rows": [{"k": "v"}],
        }}), (["k"], [["v"]]))


REQUIRED_SECTION_TYPES = ("four_d_workshop_record", "three_e_workshop_record", "three_e_report",
                          "five_c_workshop_record", "five_c_report", "innovation_scorecard")


class RenderTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        spec = importlib.util.spec_from_file_location("required_sections_render", SCRIPT)
        cls.renderer = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(cls.renderer)
        cls.profiles = json.loads((ROOT / "profiles.json").read_text())
        cls.sections = json.loads((ROOT / "response-sections.json").read_text())

    def complete_section_document(self, kind):
        item = content(kind, kind)
        item["metadata"]["completeness"] = "complete"
        item["blocks"] = [{"id": key, "type": "paragraph", "data": {
            "section_key": key, "text": "Registro fictício " + key}, "citation_ids": ["F01"]}
            for key in self.sections[kind]]
        return item

    def render(self, item: dict, formats: str, folder: Path):
        source = folder / "input.json"
        source.write_text(json.dumps(item, ensure_ascii=False))
        result = subprocess.run([sys.executable, str(SCRIPT), str(source), "--out-dir", str(folder), "--formats", formats], capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        output = json.loads(result.stdout)
        self.assertEqual(output["unavailable"], {})
        return output

    def test_one_page_contract_only_accepts_landscape(self):
        item = json.loads((ROOT.parent / "_SAMPLES_ENTREGAVEIS/briefing_one_page/input.json").read_text())
        envelope = json.loads((ROOT / "contracts/deliverable-content.schema.json").read_text())
        registry = Registry().with_resource(envelope["$id"], Resource.from_contents(envelope))
        for path in (ROOT / "contracts/briefing-one-page.schema.json",
                     ROOT.parent / "briefing-inovacao/references/briefing-one-page.schema.json"):
            with self.subTest(schema=path):
                schema = json.loads(path.read_text())
                validator = Draft202012Validator(schema, registry=registry)
                self.assertEqual(list(validator.iter_errors(item)), [])
                portrait = copy.deepcopy(item)
                portrait["metadata"]["orientation"] = "portrait"
                self.assertTrue(list(validator.iter_errors(portrait)))

    def test_briefing_modes_keep_supported_orientation(self):
        for kind, orientation in (("briefing", "portrait"),
                                  ("briefing_improvement_questionnaire", "portrait"),
                                  ("briefing_one_page", "landscape")):
            with self.subTest(kind=kind):
                item = json.loads((ROOT.parent / "_SAMPLES_ENTREGAVEIS" / kind / "input.json").read_text())
                self.assertEqual(item["metadata"]["orientation"], orientation)
                self.renderer.validate(item, self.profiles)

    def test_null_completeness_enforces_complete_validation(self):
        learning = self.card_document("learning_card", "partial")
        report = content("three_e_report", "invalid-report")
        potential = json.loads((ROOT.parent / "_SAMPLES_ENTREGAVEIS/innovation_potential_assessment/input.json").read_text())
        next(block["data"] for block in potential["blocks"] if block["data"].get("section_key") == "scores")["overall_score"] = 80
        for original in (learning, report, potential):
            for state in ("complete", "missing", None):
                with self.subTest(kind=original["deliverable_type"], state=state):
                    item = copy.deepcopy(original)
                    if state == "missing":
                        item["metadata"].pop("completeness", None)
                    else:
                        item["metadata"]["completeness"] = state
                    before = copy.deepcopy(item)
                    with self.assertRaises(ValueError):
                        self.renderer.validate(item, self.profiles)
                    self.assertEqual(item, before)

    def test_editable_formats_show_identity_and_state(self):
        for kind in ("three_e_report", "learning_card", "executive_presentation"):
            item = self.card_document(kind, "partial") if kind == "learning_card" else content(kind, "identity-test")
            identity = f'ID {item["metadata"]["artifact_id"]} · versão 1.0.0 · Parcial'
            with tempfile.TemporaryDirectory() as temp:
                folder = Path(temp)
                editable = "pptx" if kind == "executive_presentation" else "docx"
                self.render(item, "md,txt," + editable, folder)
                stem = item["metadata"]["artifact_id"] + "-v1.0.0"
                for fmt in ("md", "txt", editable):
                    with self.subTest(kind=kind, format=fmt):
                        path = folder / f"{stem}.{fmt}"
                        if fmt == "docx":
                            text = "\n".join(p.text for p in Document(path).paragraphs)
                            self.assertIn(identity, text)
                        elif fmt == "pptx":
                            for slide in Presentation(path).slides:
                                text = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
                                self.assertIn(identity, text)
                        else:
                            self.assertIn(identity, path.read_text())

    def test_identity_for_all_states_and_supported_profiles(self):
        for kind in ("briefing_one_page", "three_e_report", "learning_card", "executive_presentation"):
            source = ROOT.parent / "_SAMPLES_ENTREGAVEIS" / kind / "input.json"
            original = json.loads(source.read_text())
            for state in ("partial", "inconclusive", "complete", "missing", None):
                with self.subTest(kind=kind, state=state), tempfile.TemporaryDirectory() as temp:
                    item = copy.deepcopy(original)
                    if state == "missing":
                        item["metadata"].pop("completeness", None)
                    else:
                        item["metadata"]["completeness"] = state
                    label = {"partial": "Parcial", "inconclusive": "Inconclusivo"}.get(state, "Completo")
                    identity = f'ID {item["metadata"]["artifact_id"]} · versão {item["metadata"]["content_version"]} · {label}'
                    before = copy.deepcopy(item)
                    self.renderer.validate(item, self.profiles)
                    folder = Path(temp)
                    # O PDF legado da apresentação tem rodapé fora da página;
                    # o portão o reprova até o motor ReportLab cobrir esse tipo.
                    has_pdf = pdf_available() and kind != "executive_presentation"
                    editable = "pptx" if kind == "executive_presentation" else "docx"
                    self.render(item, "md,txt,html," + editable + (",pdf" if has_pdf else ""), folder)
                    stem = item["metadata"]["artifact_id"] + "-v" + item["metadata"]["content_version"]
                    for fmt in ("md", "txt", "html"):
                        self.assertIn(identity, (folder / f"{stem}.{fmt}").read_text())
                    if editable == "docx":
                        self.assertIn(identity, "\n".join(p.text for p in Document(folder / f"{stem}.docx").paragraphs))
                    else:
                        presentation = Presentation(folder / f"{stem}.pptx")
                        for slide in presentation.slides:
                            footers = [shape for shape in slide.shapes if shape.has_text_frame and identity in shape.text]
                            self.assertEqual(len(footers), 1)
                            footer = footers[0]
                            self.assertLessEqual(footer.top + footer.height, presentation.slide_height)
                            self.assertLessEqual(footer.left + footer.width, presentation.slide_width)
                            other_text = [shape for shape in slide.shapes if shape.has_text_frame and shape.shape_id != footer.shape_id]
                            self.assertGreaterEqual(footer.top, max(shape.top + shape.height for shape in other_text))
                            self.assertEqual(footer.text_frame.paragraphs[0].font.name, "Roboto")
                            self.assertGreaterEqual(footer.text_frame.paragraphs[0].font.size.pt, 7.5)
                        html = (folder / f"{stem}.html").read_text()
                        self.assertEqual(html.count(f'<footer class="meta">{identity}</footer>'), len(presentation.slides))
                        self.assertTrue(any("Fontes" in shape.text for shape in presentation.slides[-1].shapes if shape.has_text_frame))
                    if has_pdf:
                        pdf = PdfReader(folder / f"{stem}.pdf")
                        for page in pdf.pages:
                            self.assertIn(identity, " ".join(page.extract_text().split()))
                        if kind == "learning_card":
                            self.assertEqual(len(pdf.pages), 1)
                            self.assertIn("Test Card sample-test-card v1.0.0", pdf.pages[0].extract_text())
                    canonical = json.loads((folder / f"{stem}.json").read_text())
                    self.assertEqual(canonical, before)
                    self.assertEqual(item, before)

    def test_identity_uses_only_provided_date(self):
        item = content("three_e_report", "dated-identity")
        before = copy.deepcopy(item)
        self.assertEqual(self.renderer.identity_line(item), "ID dated-identity · versão 1.0.0 · Parcial")
        self.assertEqual(item, before)
        item["metadata"]["date"] = "2026-08-21"
        self.assertEqual(self.renderer.identity_line(item), "ID dated-identity · versão 1.0.0 · Parcial · data 2026-08-21")
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            self.render(item, "md,txt,html,docx", folder)
            for fmt in ("md", "txt", "html"):
                self.assertIn("data 2026-08-21", (folder / f"dated-identity-v1.0.0.{fmt}").read_text())
            self.assertIn("data 2026-08-21", "\n".join(p.text for p in Document(folder / "dated-identity-v1.0.0.docx").paragraphs))

    def test_reader_labels_are_portuguese_and_canonical_keys_are_preserved(self):
        for kind, label in [
            ("experiment_plan", "Plano de experimentação"),
            ("test_card", "Cartão de teste"),
            ("executive_presentation", "Apresentação executiva"),
        ]:
            with self.subTest(kind=kind), tempfile.TemporaryDirectory() as folder:
                path = Path(folder)
                item = content(kind, "linguagem")
                result = self.render(item, "html", path)
                html_path = next(path / name for name in result["files"] if name.endswith(".html"))
                rendered = html_path.read_text()
                self.assertIn(f'<span class="eyebrow">{label}</span>', rendered)
                self.assertIn("Evidência resumida e verificável.", rendered)
                self.assertNotIn(f'>{kind.replace("_", " ")}</span>', rendered)
                if kind != "executive_presentation":
                    self.assertIn("Parcial", rendered)
                    self.assertNotIn("perfil report_portrait", rendered)
                json_path = next(path / name for name in result["files"] if name.endswith(".json"))
                self.assertEqual(json.loads(json_path.read_text()), item)

    def test_iqb_accepts_zero_to_three_strengths_and_questions(self):
        original = json.loads((ROOT.parent / "_SAMPLES_ENTREGAVEIS/briefing_validation/input.json").read_text())
        for key in ("strengths", "provocations"):
            for count in range(4):
                with self.subTest(section=key, count=count):
                    item = copy.deepcopy(original)
                    next(b["data"] for b in item["blocks"] if b["data"].get("section_key") == key)["items"] = [f"Achado fictício {index}" for index in range(count)]
                    if count == 0:
                        item["summary"] += " Nenhum item sustentado registrado nesta seção do exemplo fictício."
                    before = copy.deepcopy(item)
                    self.renderer.validate(item, self.profiles)
                    self.assertEqual(item, before)
        item = copy.deepcopy(original)
        note = "Não há pontos fortes sustentados nem perguntas relevantes neste exemplo fictício."
        item["summary"] = note
        for block in item["blocks"]:
            if block["data"].get("section_key") in ("strengths", "provocations"):
                block["data"]["items"] = []
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            self.render(item, "md,html,docx", folder)
            stem = item["metadata"]["artifact_id"] + "-v" + item["metadata"]["content_version"]
            for fmt in ("md", "html"):
                self.assertIn(note, (folder / f"{stem}.{fmt}").read_text())
            self.assertIn(note, "\n".join(p.text for p in Document(folder / f"{stem}.docx").paragraphs))
            canonical = json.loads((folder / f"{stem}.json").read_text())
            self.assertEqual(canonical, item)

    def test_iqb_rejects_four_strengths_or_questions(self):
        original = json.loads((ROOT.parent / "_SAMPLES_ENTREGAVEIS/briefing_validation/input.json").read_text())
        for key in ("strengths", "provocations"):
            for case in ("four", "missing"):
                with self.subTest(section=key, case=case):
                    item = copy.deepcopy(original)
                    if case == "missing":
                        item["blocks"] = [b for b in item["blocks"] if b["data"].get("section_key") != key]
                    else:
                        next(b["data"] for b in item["blocks"] if b["data"].get("section_key") == key)["items"] = ["Achado fictício"] * 4
                    with self.assertRaisesRegex(ValueError, "Conteúdo inválido"):
                        self.renderer.validate(item, self.profiles)

    def test_iqb_derivatives_keep_scores_vetoes_and_audit(self):
        path = ROOT.parent / "_SAMPLES_ENTREGAVEIS/briefing_validation/input.json"
        if not path.exists():
            self.skipTest("Amostra IQB ausente")
        item = json.loads(path.read_text(encoding="utf-8"))
        scores = next(block["data"] for block in item["blocks"] if block["id"] == "scores")
        scores["dimensions"]["clarity"]["items"][0]["citation_id"] = "FIC-01"
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            self.render(item, "md,docx,html,txt", folder)
            stem = f"{item['metadata']['artifact_id']}-v{item['metadata']['content_version']}"
            doc = Document(folder / f"{stem}.docx")
            doc_text = "\n".join([p.text for p in doc.paragraphs] +
                                 [" | ".join(cell.text for cell in row.cells)
                                  for table in doc.tables for row in table.rows])
            for format in ("md", "docx", "html", "txt"):
                with self.subTest(format=format):
                    text = doc_text if format == "docx" else (folder / f"{stem}.{format}").read_text(encoding="utf-8")
                    if format == "html":
                        text = re.sub(r"<style>.*?</style>", "", text, flags=re.DOTALL)
                        text = re.sub(r"<[^>]*>", " ", text.replace("</tr>", "</tr>\n"))
                    for prefix in ("CLA", "EVI", "OUT", "ACT"):
                        for index in range(1, 11):
                            self.assertIn(f"{prefix}-{index:02d}", text)
                    self.assertTrue("Nota bruta" in text, f"{format}: Nota bruta ausente")
                    self.assertTrue("Nota final" in text, f"{format}: Nota final ausente")
                    self.assertIn(scores["dimensions"]["clarity"]["justification"], text)
                    self.assertIn("Auditoria de agrupamento", text)
                    self.assertRegex(text, rf"Amplitude[\s:|]+{scores['anti_clustering']['span']}\b")
                    self.assertTrue(any("CLA-01" in line and "FIC-01" in line for line in text.splitlines()))

    def test_card_grid_citations_survive_md_txt(self):
        for second_refs in (["F02"], ["F01", "F02"]):
            with self.subTest(second_refs=second_refs), tempfile.TemporaryDirectory() as temp:
                item = content("four_d_report", "grid-citations")
                cards = [{"title": "Primeiro card", "text": "Primeiro achado", "citation_ids": ["F01"]},
                         {"title": "Segundo | card\r\ncontinuação", "text": "Segundo | achado\nOutra linha", "citation_ids": second_refs},
                         {"title": "Terceiro card", "text": "Achado sem fonte"}]
                item["blocks"] = [{"id": "grid", "type": "card_grid", "data": {"cards": cards}, "citation_ids": ["FB"]}]
                item["citations"] = [{"id": key, "source_id": "fonte-ficticia-" + key} for key in ("F01", "F02", "FB")]
                before = copy.deepcopy(item)
                folder = Path(temp)
                self.render(item, "md,txt,html,docx", folder)
                for fmt in ("md", "txt"):
                    rendered = (folder / f"grid-citations-v1.0.0.{fmt}").read_text()
                    rows = [line for line in rendered.splitlines() if line.startswith(("| Primeiro card |", "| Segundo", "| Terceiro card |"))]
                    self.assertEqual(len(rows), 3)
                    self.assertIn("Fontes: F01", rows[0])
                    self.assertNotIn("F02", rows[0])
                    self.assertIn("Fontes: " + ", ".join(second_refs), rows[1])
                    self.assertNotIn("Fontes:", rows[2])
                    self.assertTrue(all("FB" not in row for row in rows))
                    self.assertIn("Fontes: FB", rendered.splitlines())
                    self.assertIn("- F01: fonte-ficticia-F01", rendered)
                    self.assertIn("- F02: fonte-ficticia-F02", rendered)
                    self.assertIn(r"Segundo \| card<br>continuação", rows[1])
                    self.assertIn(r"Segundo \| achado<br>Outra linha", rows[1])
                    self.assertNotIn("\r", rows[1])
                html = (folder / "grid-citations-v1.0.0.html").read_text()
                articles = re.findall(r'<article class="grid-card[^"]*">(.*?)</article>', html, re.DOTALL)
                self.assertEqual([re.findall(r"<sup>\[(.*?)\]</sup>", article) for article in articles],
                                 [["F01"], second_refs, []])
                doc = Document(folder / "grid-citations-v1.0.0.docx")
                cells = [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
                self.assertEqual(len(cells), 3)
                self.assertIn("Fontes: F01", cells[0])
                self.assertIn("Fontes: " + ", ".join(second_refs), cells[1])
                self.assertNotIn("Fontes:", cells[2])
                self.assertTrue(any(p.text == "Fontes: FB" for p in doc.paragraphs))
                self.assertEqual(item, before)
                item["blocks"][0]["data"]["cards"][0]["citation_ids"] = ["UNKNOWN"]
                with self.assertRaisesRegex(ValueError, "citações inexistentes"):
                    self.renderer.validate(item, self.profiles)

    def test_md_escapes_pipes(self):
        item = content("innovation_potential_assessment", "pipe-test")
        item["blocks"][2]["data"]["rows"] = [["Casos", "a|b"]]
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            self.render(item, "md", folder)
            self.assertIn(r"a\|b", (folder / "pipe-test-v1.0.0.md").read_text(encoding="utf-8"))

    @unittest.skipUnless(pdf_available(), "WeasyPrint/pango indisponível")
    def test_iqb_pdf_keeps_wide_table_columns(self):
        path = ROOT.parent / "_SAMPLES_ENTREGAVEIS/briefing_validation/input.json"
        if not path.exists():
            self.skipTest("Amostra IQB ausente")
        item = json.loads(path.read_text(encoding="utf-8"))
        scores = next(block["data"] for block in item["blocks"] if block["id"] == "scores")
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            self.render(item, "pdf", folder)
            stem = f"{item['metadata']['artifact_id']}-v{item['metadata']['content_version']}"
            reader = PdfReader(folder / f"{stem}.pdf")
            text = " ".join(" ".join(page.extract_text() for page in reader.pages).split())
            for label in ("Citação", "Nota bruta", "Nota final", "Rótulo", "Justificativa", "Auditoria de agrupamento"):
                self.assertTrue(label in text, f"PDF: coluna/seção {label!r} ausente")
            for detail in scores["dimensions"].values():
                self.assertTrue(detail["justification"] in text, "PDF: justificativa cortada")

    @unittest.skipUnless(pdf_available(), "WeasyPrint/pango indisponível")
    def test_portrait_and_landscape_and_semantic_identity(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            result = self.render(content("innovation_potential_assessment", "potencial-01"), "html,pdf,docx,md,txt,csv", folder)
            self.assertEqual(len(PdfReader(str(folder / "potencial-01-v1.0.0.pdf")).pages), 1)
            page = PdfReader(str(folder / "potencial-01-v1.0.0.pdf")).pages[0]
            self.assertLess(float(page.mediabox.width), float(page.mediabox.height))
            doc = Document(folder / "potencial-01-v1.0.0.docx")
            self.assertIn("Seção principal", "\n".join(p.text for p in doc.paragraphs))
            self.assertIn("Fontes: F01", "\n".join(p.text for p in doc.paragraphs))
            self.assertIn("Seção principal", (folder / "potencial-01-v1.0.0.md").read_text())
            self.assertIn("Seção principal", (folder / "potencial-01-v1.0.0.html").read_text())
            self.assertIn("medidas", " ".join(result["files"]))
            first = (folder / "potencial-01-v1.0.0.html").read_bytes()
            self.render(content("innovation_potential_assessment", "potencial-01"), "html,pdf,docx,md,txt,csv", folder)
            self.assertEqual(first, (folder / "potencial-01-v1.0.0.html").read_bytes())

            self.render(content("four_d_workshop_record", "workshop-4d-01"), "pdf", folder)
            page = PdfReader(str(folder / "workshop-4d-01-v1.0.0.pdf")).pages[0]
            self.assertGreater(float(page.mediabox.width), float(page.mediabox.height))

    @unittest.skipUnless(pdf_available(), "WeasyPrint/pango indisponível")
    def test_card_is_one_page_and_presentation_is_editable(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            card = content("test_card", "TC-01")
            card["blocks"] = [
                {"id": f"f-{i}", "type": "paragraph", "data": {"section_key": key, "text": label}}
                for i, (key, label) in enumerate(zip(("hypothesis", "test", "measure", "criteria"),
                                                   ("Acreditamos que", "Para verificar, vamos", "Vamos medir", "Estaremos certos se")), 1)
            ]
            self.render(card, "pdf,docx,md", folder)
            self.assertEqual(len(PdfReader(str(folder / "TC-01-v1.0.0.pdf")).pages), 1)
            self.assertIn("Acreditamos que", (folder / "TC-01-v1.0.0.md").read_text())
            presentation = content("executive_presentation", "apresentacao-01")
            self.render(presentation, "pptx", folder)
            prs = Presentation(folder / "apresentacao-01-v1.0.0.pptx")
            self.assertGreaterEqual(len(prs.slides), 2)
            self.assertTrue(any("Seção principal" in shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame))
            self.assertTrue(any("Fontes" in shape.text for shape in prs.slides[-1].shapes if shape.has_text_frame))

    def test_missing_pdf_backend_still_renders_other_formats(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            backend = folder / "backend"
            backend.mkdir()
            (backend / "weasyprint.py").write_text('raise OSError("pango ausente (simulado)")\n', encoding="utf-8")
            source = folder / "input.json"
            source.write_text(json.dumps(content("learning_card", "LC-fallback")), encoding="utf-8")
            out = folder / "output"
            env = {**os.environ, "PYTHONPATH": str(backend)}
            result = subprocess.run([sys.executable, str(SCRIPT), str(source), "--out-dir", str(out),
                                     "--formats", "pdf,md,docx"], env=env, capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            output = json.loads(result.stdout)
            self.assertEqual(output["unavailable"], {"pdf": "OSError: pango ausente (simulado)"})
            self.assertIn("AVISO: PDF não gerado", result.stderr)
            self.assertIn("brew install pango", result.stderr)
            self.assertTrue((out / "LC-fallback-v1.0.0.md").is_file())
            self.assertTrue((out / "LC-fallback-v1.0.0.docx").is_file())
            manifest = json.loads((out / output["manifest"]).read_text())
            self.assertEqual(manifest["unavailable"], output["unavailable"])
            self.assertNotIn("LC-fallback-v1.0.0.pdf", output["files"])
            self.assertFalse((out / "LC-fallback-v1.0.0.pdf").exists())

    def test_pdf_write_failure_removes_partial_pdf(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            backend = folder / "backend"
            backend.mkdir()
            (backend / "weasyprint.py").write_text(
                "from pathlib import Path\n"
                "class HTML:\n"
                "    def __init__(self, **kwargs): pass\n"
                "    def write_pdf(self, path):\n"
                "        Path(path).write_bytes(b'PDF incompleto')\n"
                "        raise OSError('x' * 400)\n", encoding="utf-8")
            source = folder / "input.json"
            source.write_text(json.dumps(content("learning_card", "LC-write-failure")), encoding="utf-8")
            out = folder / "output"
            result = subprocess.run([sys.executable, str(SCRIPT), str(source), "--out-dir", str(out),
                                     "--formats", "pdf,md,docx"],
                                    env={**os.environ, "PYTHONPATH": str(backend)}, capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            output = json.loads(result.stdout)
            self.assertEqual(len(output["unavailable"]["pdf"]), 300)
            self.assertFalse((out / "LC-write-failure-v1.0.0.pdf").exists())
            self.assertTrue((out / "LC-write-failure-v1.0.0.md").exists())
            self.assertTrue((out / "LC-write-failure-v1.0.0.docx").exists())

    def test_card_field_over_limit_is_rejected_in_any_state(self):
        for completeness in ("complete", "partial"):
            item = content("test_card", "TC-long")
            item["metadata"]["completeness"] = completeness
            item["blocks"][3]["data"]["text"] = "x" * 281
            with tempfile.TemporaryDirectory() as temp:
                folder = Path(temp)
                source = folder / "input.json"
                source.write_text(json.dumps(item), encoding="utf-8")
                result = subprocess.run([sys.executable, str(SCRIPT), str(source), "--out-dir", str(folder),
                                         "--formats", "md"], capture_output=True, text=True)
                self.assertNotEqual(result.returncode, 0)
                self.assertIn("test_card.criteria: 281 caracteres", result.stderr)
                self.assertFalse((folder / "TC-long-v1.0.0.md").exists())

    def test_card_field_at_limit_is_accepted(self):
        item = content("test_card", "TC-limit")
        item["blocks"][3]["data"]["text"] = "x" * 280
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            source = folder / "input.json"
            source.write_text(json.dumps(item), encoding="utf-8")
            result = subprocess.run([sys.executable, str(SCRIPT), str(source), "--out-dir", str(folder),
                                     "--formats", "md"], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)

    @unittest.skipUnless(pdf_available(), "WeasyPrint/pango indisponível")
    def test_oversized_card_remains_fatal(self):
        # Os campos já têm limite de caracteres; o título longo exercita a checagem de página única.
        item = content("test_card", "TC-oversized")
        item["title"] = "Título longo para verificar o limite do card. " * 120
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            source = folder / "input.json"
            source.write_text(json.dumps(item), encoding="utf-8")
            result = subprocess.run([sys.executable, str(SCRIPT), str(source), "--out-dir", str(folder),
                                     "--formats", "pdf,md"], capture_output=True, text=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("excedeu uma página", result.stderr)
            self.assertNotIn("AVISO: PDF não gerado", result.stderr)
            self.assertFalse((folder / "TC-oversized-v1.0.0.pdf").exists())
            self.assertFalse((folder / "TC-oversized-v1.0.0.manifest.json").exists())

    def test_rejects_wrong_orientation(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            item = content("innovation_scorecard", "score-01")
            item["metadata"]["orientation"] = "portrait"
            source = folder / "input.json"
            source.write_text(json.dumps(item))
            result = subprocess.run([sys.executable, str(SCRIPT), str(source), "--out-dir", str(folder)], capture_output=True, text=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("orientation", result.stderr)

    def card_document(self, kind, state):
        keys = ("hypothesis", "test", "measure", "criteria") if kind == "test_card" else ("hypothesis", "observations", "learning", "action")
        item = content(kind, "TC-state" if kind == "test_card" else "LC-state")
        item["metadata"]["completeness"] = state
        item["blocks"] = [{"id": key, "type": "paragraph", "data": {
            "section_key": key, "text": "Pendente: registrar " + key}} for key in keys]
        item["citations"] = []
        if state == "complete":
            item["citations"] = [{"id": "F01", "source_id": "registro-ficticio"}]
            if kind == "learning_card":
                item["metadata"].update(test_card_id="TC-executed", test_card_version="1")
                item["blocks"][1]["citation_ids"] = ["F01"]
        return item

    def test_cards_keep_four_fields_for_all_states(self):
        for kind in ("test_card", "learning_card"):
            for state in ("complete", "partial", "inconclusive"):
                original = self.card_document(kind, state)
                with self.subTest(kind=kind, state=state, case="valid"):
                    self.renderer.validate(original, self.profiles)
                    self.assertEqual(original, self.card_document(kind, state))
                for case in ("only_hypothesis", "missing", "duplicate", "inverted", "extra", "blank", "page_break", "list", "table", "card_grid"):
                    with self.subTest(kind=kind, state=state, case=case):
                        item = copy.deepcopy(original)
                        if case == "only_hypothesis":
                            item["blocks"] = item["blocks"][:1]
                        elif case == "missing":
                            item["blocks"].pop()
                        elif case == "duplicate":
                            item["blocks"][2]["data"]["section_key"] = item["blocks"][0]["data"]["section_key"]
                        elif case == "inverted":
                            item["blocks"][2:4] = reversed(item["blocks"][2:4])
                        elif case == "extra":
                            item["blocks"].append({"id": "extra", "type": "paragraph", "data": {"text": "Extra"}})
                        elif case == "blank":
                            item["blocks"][2]["data"]["text"] = " \n\t "
                        else:
                            item["blocks"][2]["type"] = case
                            data = item["blocks"][2]["data"]
                            if case == "list":
                                data["items"] = ["Item diferente"]
                            elif case == "table":
                                data.update(headers=["Coluna"], rows=[["Linha diferente"]])
                            elif case == "card_grid":
                                data["cards"] = [{"title": "Card próprio", "text": "Texto diferente"}]
                        before = copy.deepcopy(item)
                        with self.assertRaises(ValueError):
                            self.renderer.validate(item, self.profiles)
                        self.assertEqual(item, before)
            with self.subTest(kind=kind, case="CLI"), tempfile.TemporaryDirectory() as temp:
                item = self.card_document(kind, "partial")
                item["blocks"] = item["blocks"][:1]
                folder = Path(temp)
                source = folder / "input.json"
                source.write_text(json.dumps(item))
                result = subprocess.run([sys.executable, str(SCRIPT), str(source),
                                         "--out-dir", str(folder / "output"), "--formats", "md"],
                                        capture_output=True, text=True)
                self.assertNotEqual(result.returncode, 0)
                self.assertFalse((folder / "output").exists())
            with self.subTest(kind=kind, case="default_complete"):
                item = self.card_document(kind, "complete")
                del item["metadata"]["completeness"]
                self.renderer.validate(item, self.profiles)
                if kind == "learning_card":
                    del item["metadata"]["test_card_id"]
                    with self.assertRaisesRegex(ValueError, "Conteúdo inválido"):
                        self.renderer.validate(item, self.profiles)

    def test_partial_cards_preserve_pending_fields_in_exports(self):
        for kind in ("test_card", "learning_card"):
            for state in ("partial", "inconclusive"):
                with self.subTest(kind=kind, state=state), tempfile.TemporaryDirectory() as temp:
                    item = self.card_document(kind, state)
                    texts = (["Hipótese pendente de definição.", "Pendente: definir teste antes da coleta.",
                              "Pendente: definir medida antes da coleta.", "Pendente: definir critério antes da coleta."]
                             if kind == "test_card" else
                             ["Hipótese do teste ainda não informada.", "Não observado: coleta ainda não realizada.",
                              "Aprendizado pendente de observação.", "Ação pendente de análise."])
                    for block, text, block_type in zip(item["blocks"], texts, ("paragraph", "heading", "callout", "metric")):
                        block["data"]["text"] = text
                        block["type"] = block_type
                    before = copy.deepcopy(item)
                    folder = Path(temp)
                    has_pdf = pdf_available()
                    self.render(item, "md,html,docx" + (",pdf" if has_pdf else ""), folder)
                    self.assertEqual(item, before)
                    stem = item["metadata"]["artifact_id"] + "-v1.0.0"
                    exports = {fmt: (folder / f"{stem}.{fmt}").read_text() for fmt in ("md", "html")}
                    exports["docx"] = "\n".join(p.text for p in Document(folder / f"{stem}.docx").paragraphs)
                    if has_pdf:
                        exports["pdf"] = "\n".join(p.extract_text() for p in PdfReader(folder / f"{stem}.pdf").pages)
                    for fmt, rendered in exports.items():
                        for text in texts:
                            self.assertIn(text.casefold(), " ".join(rendered.split()).casefold(), (kind, state, fmt, text))
                    item["blocks"][0]["citation_ids"] = ["INEXISTENTE"]
                    with self.assertRaisesRegex(ValueError, "citações inexistentes"):
                        self.renderer.validate(item, self.profiles)

    def test_complete_cards_require_four_fields_and_learning_evidence(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            card = content("test_card", "TC-02")
            card["metadata"]["completeness"] = "complete"
            card["blocks"] = [
                {"id": key, "type": "paragraph", "data": {"section_key": key, "text": key}}
                for key in ("hypothesis", "test", "measure", "criteria")
            ]
            self.render(card, "md", folder)
            card["blocks"].pop()
            source = folder / "input.json"
            source.write_text(json.dumps(card))
            result = subprocess.run([sys.executable, str(SCRIPT), str(source), "--out-dir", str(folder), "--formats", "md"], capture_output=True, text=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("Conteúdo inválido", result.stderr)

            learning = content("learning_card", "LC-02")
            learning["metadata"].update(completeness="complete", test_card_id="TC-02", test_card_version="1.0.0")
            learning["blocks"] = [
                {"id": key, "type": "paragraph", "data": {"section_key": key, "text": key}}
                for key in ("hypothesis", "observations", "learning", "action")
            ]
            source.write_text(json.dumps(learning))
            result = subprocess.run([sys.executable, str(SCRIPT), str(source), "--out-dir", str(folder), "--formats", "md"], capture_output=True, text=True)
            self.assertNotEqual(result.returncode, 0)
            learning["blocks"][1]["citation_ids"] = ["F01"]
            self.render(learning, "md", folder)

    def test_learning_card_link_required_only_for_complete_content(self):
        sample = json.loads((ROOT.parent / "_SAMPLES_ENTREGAVEIS/learning_card/input.json").read_text())
        for field, value in [("test_card_id", None), ("test_card_version", None),
                             ("test_card_id", ""), ("test_card_id", "invalid id"),
                             ("test_card_version", ""), ("test_card_version", "invalid version")]:
            with self.subTest(field=field, value=value), tempfile.TemporaryDirectory() as temp:
                item = json.loads(json.dumps(sample))
                if value is None:
                    del item["metadata"][field]
                else:
                    item["metadata"][field] = value
                folder = Path(temp)
                source = folder / "input.json"
                source.write_text(json.dumps(item))
                result = subprocess.run([sys.executable, str(SCRIPT), str(source),
                                         "--out-dir", str(folder / "output"), "--formats", "md"],
                                        capture_output=True, text=True)
                self.assertNotEqual(result.returncode, 0)
                self.assertIn("Conteúdo inválido", result.stderr)
                self.assertFalse((folder / "output").exists())
        with tempfile.TemporaryDirectory() as temp:
            sample["metadata"]["completeness"] = "partial"
            del sample["metadata"]["test_card_id"]
            del sample["metadata"]["test_card_version"]
            self.render(sample, "md", Path(temp))

    def test_learning_card_link_versions_and_visible_exports(self):
        sample = json.loads((ROOT.parent / "_SAMPLES_ENTREGAVEIS/learning_card/input.json").read_text())
        for version in ("1", "1.0.0"):
            with self.subTest(version=version), tempfile.TemporaryDirectory() as temp:
                sample["metadata"]["test_card_version"] = version
                folder = Path(temp)
                formats = "html,docx,md,txt" + (",pdf" if pdf_available() else "")
                self.render(sample, formats, folder)
                stem = "sample-learning-card-v1.0.0"
                link = "Test Card: sample-test-card v" + version
                for extension in ("md", "txt"):
                    self.assertIn(link, (folder / f"{stem}.{extension}").read_text())
                self.assertIn("Test Card sample-test-card v" + version,
                              (folder / f"{stem}.html").read_text())
                self.assertIn(link, "\n".join(p.text for p in Document(folder / f"{stem}.docx").paragraphs))
                if pdf_available():
                    text = "\n".join(p.extract_text() for p in PdfReader(folder / f"{stem}.pdf").pages)
                    self.assertIn("Test Card sample-test-card v" + version, text)

    def test_required_sections_reject_hidden_block_types(self):
        for kind in REQUIRED_SECTION_TYPES:
            original = self.complete_section_document(kind)
            for index, section_block in enumerate(original["blocks"]):
                for block_type in ("page_break", "list", "table", "card_grid"):
                    with self.subTest(kind=kind, section=section_block["id"], block_type=block_type):
                        item = copy.deepcopy(original)
                        block = item["blocks"][index]
                        block["type"] = block_type
                        if block_type == "list":
                            block["data"]["items"] = ["Item próprio da lista"]
                        elif block_type == "table":
                            block["data"].update(headers=["Coluna"], rows=[["Linha própria da tabela"]])
                        elif block_type == "card_grid":
                            block["data"]["cards"] = [{"title": "Card próprio", "text": "Texto próprio da grade"}]
                        with self.assertRaisesRegex(ValueError, "Conteúdo inválido"):
                            self.renderer.validate(item, self.profiles)
            with self.subTest(kind=kind, boundary="CLI"), tempfile.TemporaryDirectory() as temp:
                item = copy.deepcopy(original)
                item["blocks"][0]["type"] = "page_break"
                folder = Path(temp)
                source = folder / "input.json"
                source.write_text(json.dumps(item))
                result = subprocess.run([sys.executable, str(SCRIPT), str(source),
                                         "--out-dir", str(folder / "output"), "--formats", "md"],
                                        capture_output=True, text=True)
                self.assertNotEqual(result.returncode, 0)
                self.assertIn("Conteúdo inválido", result.stderr)
                self.assertFalse((folder / "output").exists())

    def test_required_sections_reject_missing_blank_and_uncited_summaries(self):
        for kind in REQUIRED_SECTION_TYPES:
            original = self.complete_section_document(kind)
            for index, section_block in enumerate(original["blocks"]):
                cases = ["missing_section", "blank_text", "missing_type", "missing_data"]
                if kind.startswith("five_c"):
                    cases += ["missing_citations", "empty_citations"]
                for case in cases:
                    with self.subTest(kind=kind, section=section_block["id"], case=case):
                        item = copy.deepcopy(original)
                        block = item["blocks"][index]
                        if case == "missing_section":
                            item["blocks"].pop(index)
                        elif case == "blank_text":
                            block["data"]["text"] = " \n\t "
                        elif case == "missing_type":
                            del block["type"]
                        elif case == "missing_data":
                            del block["data"]
                        elif case == "missing_citations":
                            del block["citation_ids"]
                        else:
                            block["citation_ids"] = []
                        with self.assertRaisesRegex(ValueError, "Conteúdo inválido"):
                            self.renderer.validate(item, self.profiles)

    def test_required_section_text_survives_exports(self):
        allowed = ("paragraph", "heading", "callout", "metric")
        for kind_index, kind in enumerate(REQUIRED_SECTION_TYPES):
            with self.subTest(kind=kind), tempfile.TemporaryDirectory() as temp:
                item = self.complete_section_document(kind)
                item["blocks"] = []
                texts = []
                for section_index, key in enumerate(self.sections[kind]):
                    for variant, block_type in enumerate(allowed):
                        text = f"Registro fictício {kind_index} seção {section_index} variante {variant}."
                        texts.append(text)
                        item["blocks"].append({"id": f"{key}-{block_type}", "type": block_type,
                                               "data": {"section_key": key, "text": text}, "citation_ids": ["F01"]})
                item["blocks"].insert(2, {"id": "extra-break", "type": "page_break", "data": {}})
                item["blocks"] += [
                    {"id": "extra-list", "type": "list", "data": {"items": ["Lista complementar"]}},
                    {"id": "extra-table", "type": "table", "data": {"headers": ["Coluna"], "rows": [["Tabela complementar"]]}},
                    {"id": "extra-grid", "type": "card_grid", "data": {"cards": [{"title": "Card complementar", "text": "Texto complementar"}]}},
                ]
                folder = Path(temp)
                has_pdf = pdf_available()
                formats = "html,md,txt,docx" + (",pdf" if has_pdf else "")
                self.render(item, formats, folder)
                stem = f"{kind}-v1.0.0"
                exports = {fmt: (folder / f"{stem}.{fmt}").read_text() for fmt in ("html", "md", "txt")}
                doc = Document(folder / f"{stem}.docx")
                exports["docx"] = "\n".join([p.text for p in doc.paragraphs] +
                                            [cell.text for table in doc.tables for row in table.rows for cell in row.cells])
                if has_pdf:
                    exports["pdf"] = "\n".join(page.extract_text() for page in PdfReader(folder / f"{stem}.pdf").pages)
                for fmt, rendered in exports.items():
                    normalized = " ".join(rendered.split())
                    for text in texts:
                        self.assertIn(text.casefold(), normalized.casefold(), (kind, fmt, text))

    def test_every_type_has_schema_and_complete_sections_are_enforced(self):
        profiles = json.loads((ROOT / "profiles.json").read_text())
        sections = json.loads((ROOT / "response-sections.json").read_text())
        for kind, profile in profiles["types"].items():
            self.assertTrue(profile.get("schema"), kind)
            self.assertTrue((ROOT / "contracts" / profile["schema"]).is_file(), kind)
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            for kind, keys in sections.items():
                item = content(kind, kind)
                item["metadata"]["completeness"] = "complete"
                item["blocks"] = [
                    {"id": key, "type": "paragraph", "data": {"section_key": key, "text": "Síntese registrada."}, "citation_ids": ["F01"]}
                    for key in keys
                ]
                if kind == "innovation_potential_assessment":
                    item["blocks"].insert(0, {
                        "id": "scores", "type": "table", "data": {
                            "section_key": "scores", "columns": ["criterion", "score", "weight", "justification"],
                            "rows": [{"criterion": key, "score": 80, "weight": weight, "justification": "Evidência de teste."}
                                     for key, weight in [("impact", 0.30), ("viability", 0.25),
                                                         ("strategic_alignment", 0.20), ("urgency", 0.15), ("innovation", 0.10)]],
                            "overall_score": 80.0, "classification": "priorizar",
                        },
                    })
                self.render(item, "md", folder)
                item["blocks"].pop()
                source = folder / "input.json"
                source.write_text(json.dumps(item))
                result = subprocess.run([sys.executable, str(SCRIPT), str(source), "--out-dir", str(folder), "--formats", "md"], capture_output=True, text=True)
                self.assertNotEqual(result.returncode, 0, kind)


class VisualTokenContractTests(unittest.TestCase):
    """Garante que DOCX/PPTX herdam fonte/cor de tokens.json, não de literais soltos."""

    def setUp(self):
        self.tokens = json.loads((ROOT / "tokens.json").read_text(encoding="utf-8"))

    def test_docx_normal_and_title_follow_tokens(self):
        from docx import Document
        path = ROOT.parent / "_SAMPLES_ENTREGAVEIS/briefing_validation/input.json"
        if not path.exists():
            self.skipTest("Amostra ausente")
        with tempfile.TemporaryDirectory() as temp:
            out = Path(temp)
            subprocess.run([sys.executable, str(SCRIPT), str(path), "--out-dir", str(out), "--formats", "docx"],
                           check=True, capture_output=True, text=True)
            docx_file = next(out.glob("*.docx"))
            doc = Document(str(docx_file))
            self.assertEqual(doc.styles["Normal"].font.name, self.tokens["typography"]["font_body"])
            self.assertEqual(doc.styles["Title"].font.name, self.tokens["typography"]["font_brand"])
            expected_rgb = self.tokens["color"]["document_primary"].lstrip("#").upper()
            self.assertEqual(str(doc.styles["Title"].font.color.rgb), expected_rgb)

    def test_pptx_title_and_body_follow_tokens(self):
        from pptx import Presentation
        path = ROOT.parent / "_SAMPLES_ENTREGAVEIS/executive_presentation/input.json"
        if not path.exists():
            self.skipTest("Amostra ausente")
        with tempfile.TemporaryDirectory() as temp:
            out = Path(temp)
            subprocess.run([sys.executable, str(SCRIPT), str(path), "--out-dir", str(out), "--formats", "pptx"],
                           check=True, capture_output=True, text=True)
            pptx_file = next(out.glob("*.pptx"))
            prs = Presentation(str(pptx_file))
            slide = prs.slides[0]
            title_run = slide.shapes[1].text_frame.paragraphs[0].runs[0]
            self.assertEqual(title_run.font.name, self.tokens["typography"]["font_brand"])
            body_paragraph = slide.shapes[2].text_frame.paragraphs[0]
            self.assertEqual(body_paragraph.font.name, self.tokens["typography"]["font_body"])
            footer_paragraph = slide.shapes[3].text_frame.paragraphs[0]
            self.assertEqual(footer_paragraph.font.name, self.tokens["typography"]["font_body"])


class PptxLayoutJsonTests(unittest.TestCase):
    """Garante que render_pptx usa _DESIGN_SYSTEM/layouts/<tipo>.layout.json quando existe,
    e cai para as posições hardcoded quando o arquivo não existe."""

    def test_existing_layout_matches_hardcoded_positions(self):
        from pptx import Presentation
        path = ROOT.parent / "_SAMPLES_ENTREGAVEIS/executive_presentation/input.json"
        if not path.exists():
            self.skipTest("Amostra ausente")
        with tempfile.TemporaryDirectory() as temp:
            out = Path(temp)
            subprocess.run([sys.executable, str(SCRIPT), str(path), "--out-dir", str(out), "--formats", "pptx"],
                           check=True, capture_output=True, text=True)
            prs = Presentation(str(next(out.glob("*.pptx"))))
            self.assertAlmostEqual(prs.slide_width.inches, 13.333, places=2)
            self.assertAlmostEqual(prs.slide_height.inches, 7.5, places=2)
            title_shape = prs.slides[0].shapes[1]
            # executive_presentation.layout.json: titulo_slide em x=16 mm, y=32 mm.
            self.assertAlmostEqual(title_shape.left.inches, 16 / 25.4, places=2)
            self.assertAlmostEqual(title_shape.top.inches, 32 / 25.4, places=2)
            # O logo cabe no slot "logo" (16, 16, 307 x 16 mm) e em 38 x 11 mm, alinhado à direita.
            logo_shape = prs.slides[0].shapes[0]
            self.assertLessEqual(logo_shape.width.mm, 38.01)
            self.assertLessEqual(logo_shape.height.mm, 11.01)
            self.assertAlmostEqual(logo_shape.left.mm + logo_shape.width.mm, 16 + 307, places=1)
            self.assertGreaterEqual(logo_shape.top.mm, 16 - 0.01)
            self.assertLessEqual(logo_shape.top.mm + logo_shape.height.mm, 16 + 16 + 0.01)
            self.assertLessEqual(logo_shape.top.mm + logo_shape.height.mm, title_shape.top.mm + 0.01)
            # O texto do corpo usa text_body; sem cor explícita ficaria branco sobre fundo claro.
            body_shape = prs.slides[0].shapes[2]
            colors = {str(p.font.color.rgb) for p in body_shape.text_frame.paragraphs if p.text}
            self.assertEqual(colors, {"30343B"})

    def test_custom_layout_file_moves_the_title_box(self):
        # Desvio documentado em relação ao plano 029: um deliverable_type inventado
        # (__plan029_fake__) falha antes de chegar em render_pptx, em validate()
        # (render.py:106 — "deliverable_type sem perfil visual"), porque profiles.json
        # não tem entrada para tipos fictícios. Confirmado empiricamente rodando
        # scripts/render.py via subprocess com deliverable_type="__plan029_fake__":
        # ValueError: deliverable_type sem perfil visual: '__plan029_fake__'.
        # Por isso, como o próprio plano antecipa, usamos um deliverable_type real
        # ("executive_presentation") e fazemos mock.patch.object em
        # render.load_pptx_layout para devolver o layout fake (já convertido para
        # polegadas, no formato que load_pptx_layout devolveria), chamando
        # render.render_pptx diretamente em processo (não via subprocess, já que o
        # patch só vale no processo atual).
        from unittest.mock import patch
        from pptx import Presentation
        sample_path = ROOT.parent / "_SAMPLES_ENTREGAVEIS/executive_presentation/input.json"
        if not sample_path.exists():
            self.skipTest("Amostra ausente")
        content = json.loads(sample_path.read_text(encoding="utf-8"))
        tokens = json.loads((ROOT / "tokens.json").read_text(encoding="utf-8"))
        mm_to_in = 1 / 25.4
        fake_layout = {
            "width_in": 338.66 * mm_to_in,
            "height_in": 190.5 * mm_to_in,
            "slots": {
                "logo": {"x": 281.94 * mm_to_in, "y": 8.89 * mm_to_in,
                         "width": 43.18 * mm_to_in, "height": 16.19 * mm_to_in},
                "titulo_slide": {"x": 50.0 * mm_to_in, "y": 60.0 * mm_to_in,
                                 "width": 200.0 * mm_to_in, "height": 20.32 * mm_to_in},
                "corpo_slide": {"x": 15.24 * mm_to_in, "y": 39.37 * mm_to_in,
                                "width": 307.34 * mm_to_in, "height": 135.89 * mm_to_in},
                "rodape_slide": {"x": 15.24 * mm_to_in, "y": 178.31 * mm_to_in,
                                 "width": 307.34 * mm_to_in, "height": 8.38 * mm_to_in},
            },
        }
        spec = importlib.util.spec_from_file_location("pptx_layout_json_render", SCRIPT)
        render_module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(render_module)
        with tempfile.TemporaryDirectory() as temp:
            out_path = Path(temp) / "output.pptx"
            with patch.object(render_module, "load_pptx_layout", return_value=fake_layout):
                render_module.render_pptx(content, tokens, out_path)
            prs = Presentation(str(out_path))
            title_shape = prs.slides[0].shapes[1]
            # 50mm / 25.4 ~= 1.9685 pol -- bem diferente dos 0.5 pol hardcoded;
            # prova que a posição veio do layout.json, não do fallback.
            self.assertAlmostEqual(title_shape.left.inches, 50.0 / 25.4, places=2)
            self.assertAlmostEqual(title_shape.top.inches, 60.0 / 25.4, places=2)


class PdfIdempotencyTests(unittest.TestCase):
    def test_same_json_produces_same_pdf_text(self):
        if not pdf_available():
            self.skipTest("pango indisponível neste ambiente")
        path = ROOT.parent / "_SAMPLES_ENTREGAVEIS/briefing_validation/input.json"
        if not path.exists():
            self.skipTest("Amostra ausente")
        texts = []
        for _ in range(2):
            with tempfile.TemporaryDirectory() as temp:
                out = Path(temp)
                result = subprocess.run([sys.executable, str(SCRIPT), str(path),
                                         "--out-dir", str(out), "--formats", "pdf"],
                                        capture_output=True, text=True)
                self.assertEqual(result.returncode, 0, result.stderr)
                pdf_file = next(out.glob("*.pdf"))
                reader = PdfReader(str(pdf_file))
                texts.append("\n".join(page.extract_text() or "" for page in reader.pages))
        self.assertEqual(texts[0], texts[1], "Mesmo JSON produziu texto diferente entre duas execuções")


class PresentationLayoutJsonTests(unittest.TestCase):
    """Garante que o caminho 'presentation' de render_html usa
    _DESIGN_SYSTEM/layouts/<tipo>.layout.json quando existe, e não muda nada
    quando o arquivo não existe (fallback idêntico ao comportamento anterior).
    Chama render_html direto (sem subprocess, sem WeasyPrint) -- mais rápido e
    não depende de `pango`, ao contrário de testes que convertem para PDF de
    verdade."""

    @classmethod
    def setUpClass(cls):
        cls.render = load_render_module()

    def setUp(self):
        self.content = json.loads((ROOT.parent / "_SAMPLES_ENTREGAVEIS/executive_presentation/input.json").read_text(encoding="utf-8"))
        self.tokens = json.loads((ROOT / "tokens.json").read_text(encoding="utf-8"))
        self.profile = json.loads((ROOT / "profiles.json").read_text(encoding="utf-8"))["types"]["executive_presentation"]

    def test_uses_real_layout_positions(self):
        out = self.render.render_html(self.content, self.tokens, self.profile)
        self.assertIn("left: 16mm; top: 16mm", out)  # slot "logo" (faixa do cabeçalho)
        self.assertIn(".slide { position: relative; }", out)

    def test_page_margin_zeroed_when_layout_present(self):
        out = self.render.render_html(self.content, self.tokens, self.profile)
        # a regra @page { margin: 0.4in } vem do css() compartilhado e continua lá;
        # o que importa é que uma regra @page POSTERIOR zera a margem, porque em CSS
        # a última declaração de uma propriedade em regras @page do mesmo stylesheet
        # vence. Sem isso, as coordenadas absolutas abaixo ficam deslocadas ~10mm
        # da posição real do canvas (confirmado por teste manual com WeasyPrint
        # nesta sessão: @page margin:10mm desloca um elemento position:absolute
        # left:5mm para x0=15mm no PDF final, não 5mm).
        style_start = out.index("<style>")
        style_end = out.index("</style>")
        style = out[style_start:style_end]
        assert style.rindex("@page { margin: 0; }") > style.rindex("margin: 0.4in")

    def test_falls_back_without_layout_file(self):
        # DELIVERABLE_LABELS[content["deliverable_type"]] é resolvido em
        # render_html antes do ramo "presentation", então um deliverable_type
        # inventado (sem entrada em DELIVERABLE_LABELS) gera KeyError antes de
        # chegar a exercitar o fallback -- por isso simulamos "sem arquivo de
        # layout" fazendo load_presentation_layout devolver None, em vez de
        # trocar o deliverable_type.
        with unittest.mock.patch.object(self.render, "load_presentation_layout", return_value=None):
            out = self.render.render_html(self.content, self.tokens, self.profile)
        self.assertNotIn("position: absolute", out)
        self.assertNotIn(".slide { position: relative; }", out)


@unittest.skipUnless(reportlab_available(), "reportlab indisponível")
class ReportLabPdfTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.render_module = load_render_module()

    def render_sample(self, kind: str, folder: Path, formats: str = "pdf") -> subprocess.CompletedProcess[str]:
        sample = ROOT.parent / "_SAMPLES_ENTREGAVEIS" / kind / "input.json"
        return subprocess.run([sys.executable, str(SCRIPT), str(sample), "--out-dir", str(folder),
                               "--formats", formats, "--pdf-engine", "reportlab"],
                              capture_output=True, text=True)

    def test_test_card_sample_one_page_a4_portrait(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            result = self.render_sample("test_card", folder)
            self.assertEqual(result.returncode, 0, result.stderr)
            reader = PdfReader(str(next(folder.glob("*.pdf"))))
            self.assertEqual(len(reader.pages), 1)
            self.assertAlmostEqual(float(reader.pages[0].mediabox.width) * 25.4 / 72, 210, delta=0.5)
            self.assertAlmostEqual(float(reader.pages[0].mediabox.height) * 25.4 / 72, 297, delta=0.5)
            text = " ".join(reader.pages[0].extract_text().split())
            item = json.loads((ROOT.parent / "_SAMPLES_ENTREGAVEIS/test_card/input.json").read_text(encoding="utf-8"))
            self.assertIn(item["title"], text)
            for label in self.render_module.CARD_LABELS["test_card"].values():
                self.assertIn(label, text)
            for block in item["blocks"]:
                self.assertIn(block["data"]["text"], text)
            self.assertIn(self.render_module.identity_line(item), text)

    def test_learning_card_footer_links_test_card(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            result = self.render_sample("learning_card", folder)
            self.assertEqual(result.returncode, 0, result.stderr)
            text = PdfReader(str(next(folder.glob("*.pdf")))).pages[0].extract_text()
            self.assertIn("Test Card sample-test-card v1.0.0", text)

    def test_one_page_sample_landscape(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            result = self.render_sample("briefing_one_page", folder)
            self.assertEqual(result.returncode, 0, result.stderr)
            reader = PdfReader(str(next(folder.glob("*.pdf"))))
            self.assertEqual(len(reader.pages), 1)
            self.assertAlmostEqual(float(reader.pages[0].mediabox.width) * 25.4 / 72, 297, delta=0.5)
            self.assertAlmostEqual(float(reader.pages[0].mediabox.height) * 25.4 / 72, 210, delta=0.5)
            self.assertIn("Como podemos tornar as solicitações mais claras e rastreáveis?", reader.pages[0].extract_text())

    def test_oversized_card_fails_without_partial_files(self):
        item = content("test_card", "TC-big")
        item["title"] = "Título longo para verificar o limite do card. " * 120
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            source = folder / "input.json"
            source.write_text(json.dumps(item), encoding="utf-8")
            result = subprocess.run([sys.executable, str(SCRIPT), str(source), "--out-dir", str(folder),
                                     "--formats", "pdf,md", "--pdf-engine", "reportlab"],
                                    capture_output=True, text=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("excedeu uma página", result.stderr)
            self.assertFalse(list(folder.glob("*.pdf")))
            self.assertFalse(list(folder.glob("*.manifest.json")))

    def test_one_page_overflow_fails(self):
        item = json.loads((ROOT.parent / "_SAMPLES_ENTREGAVEIS/briefing_one_page/input.json").read_text(encoding="utf-8"))
        item["metadata"]["completeness"] = "partial"
        item["blocks"] = [{**block, "id": f'{block["id"]}-{i}'} for i in range(12) for block in item["blocks"]]
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            source = folder / "input.json"
            source.write_text(json.dumps(item), encoding="utf-8")
            result = subprocess.run([sys.executable, str(SCRIPT), str(source), "--out-dir", str(folder),
                                     "--formats", "pdf", "--pdf-engine", "reportlab"],
                                    capture_output=True, text=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("excedeu uma página", result.stderr)
            self.assertFalse(list(folder.glob("*.pdf")))

    def test_works_when_weasyprint_is_broken(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            fake = folder / "fake"
            fake.mkdir()
            (fake / "weasyprint.py").write_text('raise OSError("pango ausente (simulado)")\n', encoding="utf-8")
            sample = ROOT.parent / "_SAMPLES_ENTREGAVEIS/briefing_one_page/input.json"
            result = subprocess.run([sys.executable, str(SCRIPT), str(sample), "--out-dir", str(folder),
                                     "--formats", "pdf", "--pdf-engine", "reportlab"],
                                    env={**os.environ, "PYTHONPATH": str(fake)}, capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual(json.loads(result.stdout)["unavailable"], {})
            self.assertTrue(list(folder.glob("*.pdf")))

    def test_unsupported_layout_rejected_before_writing(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            item = content("executive_presentation", "PRES-1")
            source = folder / "input.json"
            source.write_text(json.dumps(item), encoding="utf-8")
            out = folder / "out"
            result = subprocess.run([sys.executable, str(SCRIPT), str(source), "--out-dir", str(out),
                                     "--formats", "pdf,md", "--pdf-engine", "reportlab"],
                                    capture_output=True, text=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("ainda não suporta", result.stderr)
            self.assertFalse(out.exists())

    def test_deterministic_bytes(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            outputs = []
            for name in ("a", "b"):
                out = folder / name
                result = self.render_sample("test_card", out)
                self.assertEqual(result.returncode, 0, result.stderr)
                outputs.append(next(out.glob("*.pdf")).read_bytes())
            self.assertEqual(outputs[0], outputs[1])

    def test_embeds_static_fonts(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            result = self.render_sample("test_card", folder)
            self.assertEqual(result.returncode, 0, result.stderr)
            page = PdfReader(str(next(folder.glob("*.pdf")))).pages[0]
            names = [str(value.get_object().get("/BaseFont", "")) for value in page["/Resources"]["/Font"].get_object().values()]
            self.assertTrue(any(name.endswith("Montserrat-Bold") for name in names), names)
            self.assertTrue(any(name.endswith("Roboto-Regular") for name in names), names)
            self.assertFalse(any("Thin" in name or "Variable" in name for name in names), names)

    def test_manifest_records_engine(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            result = self.render_sample("test_card", folder)
            self.assertEqual(result.returncode, 0, result.stderr)
            manifest = json.loads(next(folder.glob("*.manifest.json")).read_text(encoding="utf-8"))
            self.assertEqual(manifest["pdf_engine"], "reportlab")
            out = folder / "md"
            result = self.render_sample("test_card", out, formats="md")
            self.assertEqual(result.returncode, 0, result.stderr)
            manifest = json.loads(next(out.glob("*.manifest.json")).read_text(encoding="utf-8"))
            self.assertNotIn("pdf_engine", manifest)


class PdfQualityTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        try:
            import pdfminer  # noqa: F401
            import reportlab  # noqa: F401
        except ImportError as exc:
            raise unittest.SkipTest(f"dependência de PDF ausente: {exc}") from exc
        cls.render_module = load_render_module()

    def test_detects_each_defect(self):
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen.canvas import Canvas

        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "defective.pdf"
            canvas = Canvas(str(path), pagesize=A4)
            canvas.setFont("Helvetica", 12)
            canvas.drawString(580, 400, "Saindo da pagina")
            canvas.setFont("Helvetica", 6)
            canvas.drawString(50, 100, "miudo")
            canvas.setFont("Helvetica", 12)
            canvas.drawString(50, 700, "Texto sobreposto aqui")
            canvas.drawString(55, 702, "Outra linha por cima")
            canvas.save()
            problems = self.render_module.check_pdf_quality(path, json.loads((ROOT / "tokens.json").read_text()))
            self.assertEqual(len(problems), 3, problems)
            self.assertTrue(any("fora da página" in problem for problem in problems))
            self.assertTrue(any("abaixo de 7.5" in problem for problem in problems))
            self.assertTrue(any("sobrepostos" in problem for problem in problems))

    def test_kerning_is_not_overlap(self):
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen.canvas import Canvas

        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "kerning.pdf"
            canvas = Canvas(str(path), pagesize=A4)
            canvas.setFont("Helvetica", 8)
            canvas.drawString(50, 700, "OUT-01 AVA-02 TT-TT")
            canvas.save()
            tokens = json.loads((ROOT / "tokens.json").read_text())
            self.assertEqual(self.render_module.check_pdf_quality(path, tokens), [])

    def test_reportlab_outputs_pass(self):
        for kind in ("test_card", "learning_card", "briefing_one_page"):
            with self.subTest(kind=kind), tempfile.TemporaryDirectory() as temp:
                sample = ROOT.parent / "_SAMPLES_ENTREGAVEIS" / kind / "input.json"
                result = subprocess.run([sys.executable, str(SCRIPT), str(sample), "--out-dir", temp,
                                         "--formats", "pdf", "--pdf-engine", "reportlab"],
                                        capture_output=True, text=True)
                self.assertEqual(result.returncode, 0, result.stderr)

    def test_weasyprint_samples_pass(self):
        if not pdf_available():
            self.skipTest("WeasyPrint/pango indisponível")
        for sample in sorted((ROOT.parent / "_SAMPLES_ENTREGAVEIS").glob("*/input.json")):
            kind = sample.parent.name
            with self.subTest(kind=kind), tempfile.TemporaryDirectory() as temp:
                result = subprocess.run([sys.executable, str(SCRIPT), str(sample), "--out-dir", temp,
                                         "--formats", "pdf"], capture_output=True, text=True)
                self.assertEqual(result.returncode, 0, result.stderr)

    def test_gate_removes_pdf_and_fails(self):
        sample = ROOT.parent / "_SAMPLES_ENTREGAVEIS/test_card/input.json"
        with tempfile.TemporaryDirectory() as temp:
            out = Path(temp)
            with unittest.mock.patch.object(self.render_module, "check_pdf_quality", return_value=["página 1: teste"]):
                with unittest.mock.patch.object(sys, "argv", [str(SCRIPT), str(sample), "--out-dir", str(out),
                                                            "--formats", "pdf,md", "--pdf-engine", "reportlab"]):
                    with self.assertRaisesRegex(ValueError, "reprovado na verificação de qualidade"):
                        self.render_module.main()
            self.assertFalse(list(out.glob("*.pdf")))


@unittest.skipUnless(reportlab_available(), "reportlab indisponível")
class ReportLabReportTests(unittest.TestCase):
    TYPES = {
        "briefing": 1, "briefing_improvement_questionnaire": 1,
        "briefing_validation": 5, "innovation_potential_assessment": 2,
        "experiment_plan": 1, "four_d_workshop_record": 2,
        "four_d_report": 5, "three_e_workshop_record": 1,
        "three_e_report": 2, "innovation_scorecard": 1,
        "experiment_report": 2, "evidence_recommendation": 2,
        "five_c_workshop_record": 2, "five_c_report": 2,
    }

    @classmethod
    def setUpClass(cls):
        cls.renderer = load_render_module()
        spec = importlib.util.spec_from_file_location("pdf_reportlab_tests", ROOT / "scripts/pdf_reportlab.py")
        cls.pdf_module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(cls.pdf_module)

    def _render(self, kind: str, folder: Path) -> PdfReader:
        source = ROOT.parent / "_SAMPLES_ENTREGAVEIS" / kind / "input.json"
        result = subprocess.run([sys.executable, str(SCRIPT), str(source), "--out-dir", str(folder),
                                 "--formats", "pdf", "--pdf-engine", "reportlab"],
                                capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        return PdfReader(str(next(folder.glob("*.pdf"))))

    def test_group_sections_matches_html_rules(self):
        blocks = [
            {"type": "paragraph", "data": {"text": "intro"}},
            {"type": "heading", "data": {"text": "h1"}},
            {"type": "paragraph", "data": {"text": "body"}},
            {"type": "table", "data": {"rows": []}},
            {"type": "heading", "data": {"text": "h2"}},
            {"type": "paragraph", "data": {"text": "body2"}},
            {"type": "page_break", "data": {}},
            {"type": "heading", "data": {"text": "h3"}},
            {"type": "paragraph", "data": {"text": "body3"}},
        ]
        sections = self.pdf_module.group_sections(blocks, self.renderer.block_text)
        self.assertEqual([(item[0], item[1] if item[0] == "section" else "") for item in sections],
                         [("section", "Notas"), ("section", "h1"), ("section", "h2"),
                          ("page_break", ""), ("section", "h3")])
        self.assertFalse(sections[0][3])
        self.assertTrue(sections[1][3])

    def test_all_report_samples_render(self):
        for kind, reference_pages in self.TYPES.items():
            with self.subTest(kind=kind), tempfile.TemporaryDirectory() as temp:
                reader = self._render(kind, Path(temp))
                self.assertLessEqual(len(reader.pages), reference_pages + 1)

    def test_report_text_is_complete(self):
        for kind in self.TYPES:
            with self.subTest(kind=kind), tempfile.TemporaryDirectory() as temp:
                item = json.loads((ROOT.parent / "_SAMPLES_ENTREGAVEIS" / kind / "input.json").read_text())
                reader = self._render(kind, Path(temp))
                extracted = " ".join(" ".join(page.extract_text().split()) for page in reader.pages)
                expected = [item["title"], item.get("summary", "")]
                for block in item["blocks"]:
                    data = block.get("data", {})
                    expected.extend(str(data[key]) for key in ("text", "title", "description")
                                    if isinstance(data.get(key), str) and data[key].strip())
                    expected.extend(str(value) for value in data.get("items", []) if isinstance(value, str))
                    if block["type"] == "table":
                        for _, _, headings, rows in self.renderer.table_views(block):
                            expected.extend(headings)
                            expected.extend(cell for row in rows for cell in row if cell)
                    if block["type"] == "card_grid":
                        for card in data.get("cards", []):
                            expected.extend(str(card[key]) for key in ("title", "value", "text")
                                            if card.get(key) is not None)
                            expected.extend(str(value) for value in card.get("items", []))
                expected.extend(citation["id"] for citation in item["citations"])
                for phrase in expected:
                    if phrase:
                        self.assertIn(" ".join(str(phrase).split()).casefold(), extracted.casefold(),
                                      f"{kind}: {phrase!r}")

    def test_long_table_repeats_header(self):
        item = content("briefing", "BR-long")
        item["blocks"] = [{"id": "large", "type": "table", "data": {
            "headers": ["Código", "Medida", "Valor"],
            "rows": [[f"L-{index:03d}", "Contagem", str(index)] for index in range(120)],
        }}]
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            source = folder / "input.json"
            source.write_text(json.dumps(item, ensure_ascii=False), encoding="utf-8")
            result = subprocess.run([sys.executable, str(SCRIPT), str(source), "--out-dir", str(folder),
                                     "--formats", "pdf", "--pdf-engine", "reportlab"],
                                    capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            pages = PdfReader(str(next(folder.glob("*.pdf")))).pages
            self.assertGreaterEqual(len(pages), 3)
            for page in pages:
                text = page.extract_text()
                if "L-" in text:
                    self.assertIn("Código", text)
                    self.assertIn("Medida", text)
                    self.assertIn("Valor", text)

    def test_footer_page_numbers(self):
        with tempfile.TemporaryDirectory() as temp:
            reader = self._render("four_d_report", Path(temp))
            item = json.loads((ROOT.parent / "_SAMPLES_ENTREGAVEIS/four_d_report/input.json").read_text())
            identity = self.renderer.identity_line(item)
            for index, page in enumerate(reader.pages, 1):
                text = page.extract_text()
                self.assertIn(identity, text)
                self.assertIn(f"página {index} de {len(reader.pages)}", text)

    def test_outline_has_sections(self):
        with tempfile.TemporaryDirectory() as temp:
            reader = self._render("four_d_report", Path(temp))
            item = json.loads((ROOT.parent / "_SAMPLES_ENTREGAVEIS/four_d_report/input.json").read_text())
            headings = [self.renderer.block_text(block) for block in item["blocks"] if block["type"] == "heading"]
            self.assertEqual([entry.title for entry in reader.outline], headings)

    def test_report_deterministic_bytes(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            outputs = []
            for name in ("first", "second"):
                target = folder / name
                self._render("briefing_validation", target)
                outputs.append(next(target.glob("*.pdf")).read_bytes())
            self.assertEqual(outputs[0], outputs[1])

    def test_section_band_never_orphan(self):
        for kind in ("four_d_report", "briefing_validation"):
            with self.subTest(kind=kind), tempfile.TemporaryDirectory() as temp:
                reader = self._render(kind, Path(temp))
                item = json.loads((ROOT.parent / "_SAMPLES_ENTREGAVEIS" / kind / "input.json").read_text())
                headings = {self.renderer.block_text(block).upper() for block in item["blocks"]
                            if block["type"] == "heading"}
                for page in reader.pages:
                    lines = [line.strip() for line in page.extract_text().splitlines() if line.strip()]
                    self.assertNotIn(lines[-2], headings)


if __name__ == "__main__":
    unittest.main()
