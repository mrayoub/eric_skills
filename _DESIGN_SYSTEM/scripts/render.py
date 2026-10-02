#!/usr/bin/env python3
"""Render canonical my_skills deliverables using the ERIC_RC6 visual baseline."""

from __future__ import annotations

import argparse
import base64
import csv
import hashlib
import html
import io
import json
import re
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
RENDERER_VERSION = "1.3.1"
DELIVERABLE_LABELS = {
    "briefing": "Briefing",
    "briefing_one_page": "Resumo do briefing",
    "briefing_improvement_questionnaire": "Perguntas para melhorar o briefing",
    "briefing_validation": "Avaliação do briefing",
    "innovation_potential_assessment": "Avaliação do potencial de inovação",
    "four_d_workshop_record": "Registro do workshop 4D",
    "four_d_report": "Relatório 4D",
    "three_e_workshop_record": "Registro do workshop 3E",
    "three_e_report": "Relatório 3E",
    "innovation_scorecard": "Avaliação das evidências de inovação",
    "experiment_plan": "Plano de experimentação",
    "test_card": "Cartão de teste",
    "learning_card": "Cartão de aprendizado",
    "experiment_report": "Relatório de experimentação",
    "evidence_recommendation": "Evidências e recomendação",
    "executive_presentation": "Apresentação executiva",
    "five_c_workshop_record": "Registro do workshop 5C",
    "five_c_report": "Relatório 5C",
}
COMPLETENESS_LABELS = {
    "complete": "Completo",
    "partial": "Parcial",
    "inconclusive": "Inconclusivo",
}
IQB_TABLE_WIDTHS = {
    "criterios": [14, 15, 10, 41, 10, 10],
    "dimensoes": [14, 7, 10, 12, 9, 10, 10, 28],
    "agrupamento": [35, 65],
}
CARD_LABELS = {
    "test_card": {
        "hypothesis": "1. Acreditamos que — hipótese",
        "test": "2. Para verificar, vamos — teste",
        "measure": "3. Vamos medir — medida",
        "criteria": "4. Estaremos certos se — critério",
    },
    "learning_card": {
        "hypothesis": "1. Acreditávamos que — hipótese testada",
        "observations": "2. Observamos que — observações",
        "learning": "3. Aprendemos que — inferência",
        "action": "4. Portanto, vamos — ação",
    },
}


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def canonical_bytes(value: Any) -> bytes:
    return (json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2) + "\n").encode("utf-8")


def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def safe_name(value: str) -> str:
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]*", value):
        raise ValueError(f"ID ou versão inadequado para nome de arquivo: {value!r}")
    return value


def identity_line(content: dict[str, Any]) -> str:
    meta = content["metadata"]
    state = COMPLETENESS_LABELS[meta.get("completeness") or "complete"]
    label = f'ID {meta["artifact_id"]} · versão {meta["content_version"]} · {state}'
    if isinstance(meta.get("date"), str) and meta["date"].strip():
        label += f' · data {meta["date"]}'
    return label


# Campos dos cards Strategyzer: uma frase objetiva cada.
CARD_FIELD_MAX_CHARS = 280


def validate(content: dict[str, Any], profiles: dict[str, Any]) -> dict[str, Any]:
    try:
        from jsonschema import Draft202012Validator
        from referencing import Registry, Resource
    except ImportError as exc:
        raise RuntimeError("Validação requer jsonschema e referencing: pip install jsonschema") from exc

    kind = content.get("deliverable_type")
    profile = profiles["types"].get(kind)
    if profile is None:
        raise ValueError(f"deliverable_type sem perfil visual: {kind!r}")
    meta = content.get("metadata", {})
    if not isinstance(meta, dict):
        raise ValueError("metadata deve ser um objeto")
    for key in ("artifact_id", "content_version"):
        if not isinstance(meta.get(key), str):
            raise ValueError(f"metadata.{key} deve ser string")
        safe_name(meta[key])
    if meta.get("orientation") not in (None, "portrait" if profile["paper"] == "a4_portrait" else "landscape" if profile["paper"] == "a4_landscape" else "wide"):
        raise ValueError("metadata.orientation contradiz a variante visual do perfil")
    if meta.get("completeness") not in (None, "complete", "partial", "inconclusive"):
        raise ValueError("metadata.completeness deve ser complete, partial ou inconclusive")

    envelope = read_json(ROOT / "contracts" / "deliverable-content.schema.json")
    registry = Registry().with_resource(envelope["$id"], Resource.from_contents(envelope))
    effective_completeness = meta.get("completeness") or "complete"
    schema_file = profile.get("schema") if effective_completeness == "complete" else None
    schema = read_json(ROOT / "contracts" / schema_file) if schema_file else envelope
    errors = sorted(Draft202012Validator(schema, registry=registry).iter_errors(content), key=lambda e: str(e.path))
    if errors:
        raise ValueError("Conteúdo inválido: " + "; ".join(f"{list(e.path)}: {e.message}" for e in errors[:8]))

    if kind in ("test_card", "learning_card"):
        keys = ("hypothesis", "test", "measure", "criteria") if kind == "test_card" else ("hypothesis", "observations", "learning", "action")
        blocks = content["blocks"]
        if len(blocks) != 4:
            raise ValueError(f"{kind}: exige exatamente quatro campos")
        for block, key in zip(blocks, keys):
            data = block["data"]
            text = data.get("text")
            if data.get("section_key") != key:
                raise ValueError(f"{kind}: campo esperado na ordem do método: {key}")
            if block["type"] not in ("paragraph", "heading", "callout", "metric"):
                raise ValueError(f"{kind}.{key}: campo deve usar um bloco textual visível")
            if not isinstance(text, str) or not text.strip():
                raise ValueError(f"{kind}.{key}: exige texto ou indicação explícita de pendência")
            if len(text.strip()) > CARD_FIELD_MAX_CHARS:
                raise ValueError(f"{kind}.{key}: {len(text.strip())} caracteres; o campo do card aceita até "
                                 f"{CARD_FIELD_MAX_CHARS}. Escreva uma frase objetiva e mova detalhes ao plano ou caderno")

    ids = [block["id"] for block in content["blocks"]]
    if len(ids) != len(set(ids)):
        raise ValueError("IDs de blocos devem ser únicos")
    citation_ids = {citation["id"] for citation in content["citations"]}
    for block in content["blocks"]:
        unknown = set(block.get("citation_ids", [])) - citation_ids
        if block["type"] == "card_grid":
            cards = block["data"].get("cards", [])
            if not 1 <= len(cards) <= 12:
                raise ValueError(f"{block['id']}: card_grid exige 1 a 12 cards")
            if block["data"].get("columns", 2) not in (1, 2, 3, 4):
                raise ValueError(f"{block['id']}: card_grid aceita 1 a 4 colunas")
            for card in cards:
                if not isinstance(card, dict) or not card.get("title"):
                    raise ValueError(f"{block['id']}: cada card exige title")
                unknown |= set(card.get("citation_ids", [])) - citation_ids
        if unknown:
            raise ValueError(f"{block['id']}: citações inexistentes: {sorted(unknown)}")

    def nested_citations(value: Any) -> set[str]:
        found: set[str] = set()
        if isinstance(value, dict):
            for key, item in value.items():
                if key == "citation_id" and isinstance(item, str):
                    found.add(item)
                elif key == "citation_ids" and isinstance(item, list):
                    found.update(entry for entry in item if isinstance(entry, str))
                found.update(nested_citations(item))
        elif isinstance(value, list):
            for item in value:
                found.update(nested_citations(item))
        return found

    for block in content["blocks"]:
        unknown = nested_citations(block["data"]) - citation_ids
        if unknown:
            raise ValueError(f"{block['id']}: citações aninhadas inexistentes: {sorted(unknown)}")
    if effective_completeness == "complete":
        import importlib.util

        spec = importlib.util.spec_from_file_location("deliverable_validators", ROOT / "scripts/validators/__init__.py")
        validators = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(validators)
        semantic_errors = validators.run(content)
        if semantic_errors:
            raise ValueError("Inconsistência semântica: " + "; ".join(semantic_errors[:12]))
    return profile


def block_text(block: dict[str, Any]) -> str:
    data = block.get("data", {})
    def value_text(value: Any) -> str:
        if isinstance(value, dict):
            return "; ".join(f"{key}: {value_text(item)}" for key, item in value.items())
        if isinstance(value, list):
            return "; ".join(value_text(item) for item in value)
        return "" if value is None else str(value)

    parts = []
    for key in ("title", "label", "text", "description", "content", "value"):
        if key in data and data[key] not in (None, "", []):
            parts.append(value_text(data[key]))
    for key, value in data.items():
        if key in {"title", "label", "text", "description", "content", "value", "section_key", "layout", "level", "ordered", "headers", "columns", "rows", "tone"}:
            continue
        if value not in (None, "", []):
            parts.append(f"{key.replace('_', ' ').capitalize()}: {value_text(value)}")
    return "\n".join(parts) or json.dumps(data, ensure_ascii=False, sort_keys=True)


def table_rows(block: dict[str, Any]) -> tuple[list[str], list[list[str]]]:
    data = block.get("data", {})
    if isinstance(data.get("dimensions"), dict):
        _, _, headers, rows = table_views(block)[0]
        return headers, rows
    columns = data.get("headers") or data.get("columns") or []
    keys = [str(h.get("key", h.get("label", ""))) if isinstance(h, dict) else str(h) for h in columns]
    headers = [str(h.get("label", h.get("key", ""))) if isinstance(h, dict) else str(h) for h in columns]
    raw_rows = data.get("rows") or []
    if raw_rows and not headers and isinstance(raw_rows[0], dict):
        keys = headers = list(raw_rows[0])
    rows: list[list[str]] = []
    for row in raw_rows:
        if isinstance(row, list):
            rows.append([str(x) if x is not None else "" for x in row])
        elif isinstance(row, dict):
            values = [row.get(k) for k in keys] if keys else row.values()
            rows.append([str(value) if value is not None else "" for value in values])
        else:
            rows.append([str(row)])
    if not headers and not rows:
        rows = [[key.replace("_", " ").capitalize(), "; ".join(map(str, value)) if isinstance(value, list) else str(value)]
                for key, value in data.items() if key not in ("section_key", "layout")]
        headers = ["Campo", "Valor"]
    return headers, rows


def table_views(block: dict[str, Any]) -> list[tuple[str, str, list[str], list[list[str]]]]:
    """Retorna [(slug, legenda, cabeçalhos, linhas)]; três vistas para o IQB."""
    data = block.get("data", {})
    if not isinstance(data.get("dimensions"), dict):
        return [("", data.get("title", ""), *table_rows(block))]
    labels = {"clarity": "Clareza", "evidence": "Evidências", "outcomes": "Resultados", "actionability": "Acionabilidade"}
    criteria: list[list[str]] = []
    dimensions: list[list[str]] = []
    for dimension, detail in data["dimensions"].items():
        label = labels.get(dimension, dimension)
        for item in detail.get("items", []):
            criterion = str(item.get("criterion_id", ""))
            if criterion == detail.get("veto_item_id"):
                criterion += " (veto)"
            criteria.append([label, criterion, "Sim" if item.get("answer") is True else "Não",
                             str(item.get("evidence", "")), str(item.get("evidence_source", "")),
                             str(item.get("citation_id") or "")])
        veto = f"Disparado: {detail.get('veto_detail') or '—'}" if detail.get("veto_triggered") else "Não disparado"
        dimensions.append([label, str(detail.get("yes_count", "")), str(detail.get("raw_score", "")),
                           veto, str(detail.get("veto_max_score", "")), str(detail.get("final_score", "")),
                           str(detail.get("label", "")), str(detail.get("justification", ""))])
    anti = data.get("anti_clustering", {})
    originals = anti.get("original_scores")
    audit = [["Disparada", "Sim" if anti.get("triggered") else "Não"],
             ["Amplitude", str(anti.get("span", ""))],
             ["Notas originais", ", ".join(map(str, originals)) if originals is not None else "—"],
             ["Notas da reavaliação", str(anti.get("re_evaluation_notes") or "—")]]
    return [
        ("criterios", "Critérios avaliados", ["Dimensão", "Critério", "Resposta", "Evidência", "Fonte", "Citação"], criteria),
        ("dimensoes", "Notas por dimensão", ["Dimensão", "Sim (de 9)", "Nota bruta", "Veto", "Teto do veto", "Nota final", "Rótulo", "Justificativa"], dimensions),
        ("agrupamento", "Auditoria de agrupamento", ["Campo", "Valor"], audit),
    ]


def card_text(card: dict[str, Any]) -> str:
    parts = [str(card.get("title", ""))]
    if card.get("value") is not None:
        parts.append(str(card["value"]))
    if card.get("text"):
        parts.append(str(card["text"]))
    parts.extend(str(x) for x in card.get("items", []))
    return " — ".join(x for x in parts if x)


def html_card_grid(block: dict[str, Any]) -> str:
    data = block["data"]
    columns = data.get("columns", 2)
    cards = []
    for card in data["cards"]:
        tone = card.get("tone", "blue")
        if tone not in ("blue", "neutral", "red", "green", "yellow"):
            tone = "neutral"
        title = html.escape(str(card["title"]))
        value = f'<div class="grid-value">{html.escape(str(card["value"]))}</div>' if card.get("value") is not None else ""
        body = f'<p>{html.escape(str(card["text"]))}</p>' if card.get("text") else ""
        items = '<ul>' + ''.join(f'<li>{html.escape(str(x))}</li>' for x in card.get("items", [])) + '</ul>' if card.get("items") else ""
        refs = ''.join(f'<sup>[{html.escape(str(x))}]</sup>' for x in card.get("citation_ids", []))
        cards.append(f'<article class="grid-card tone-{tone}"><h3>{title}</h3>{value}{body}{items}{refs}</article>')
    title = f'<h3 class="grid-title">{html.escape(str(data["title"]))}</h3>' if data.get("title") else ""
    refs = ''.join(f'<sup>[{html.escape(str(x))}]</sup>' for x in block.get("citation_ids", []))
    rows = ['<div class="card-row">' + "".join(cards[i:i + columns]) + "</div>" for i in range(0, len(cards), columns)]
    return f'<div class="block card-grid-block">{title}<div class="card-grid columns-{columns}">{"".join(rows)}</div>{refs}</div>'


def html_report_blocks(blocks: list[dict[str, Any]]) -> str:
    sections: list[tuple[str, bool]] = []
    pending: list[dict[str, Any]] = []
    def flush() -> None:
        if not pending:
            return
        heading = pending[0] if pending[0]["type"] == "heading" else None
        body = pending[1:] if heading else pending[:]
        wide = any(x["type"] in ("table", "card_grid", "canvas", "chart") for x in body)
        title = html.escape(block_text(heading)) if heading else "Notas"
        inner = "".join(html_block(x) for x in body)
        sections.append((f'<section class="section-card {"wide" if wide else ""}"><header class="section-band"><span>{title}</span></header><div class="section-content">{inner}</div></section>', wide))
        pending.clear()
    for block in blocks:
        if block["type"] == "page_break":
            flush()
            sections.append(('<div class="page-break"></div>', True))
        elif block["type"] == "heading":
            flush()
            pending.append(block)
        else:
            pending.append(block)
    flush()
    rendered = []
    i = 0
    while i < len(sections):
        markup, wide = sections[i]
        if not wide and i + 1 < len(sections) and not sections[i + 1][1]:
            rendered.append('<div class="section-pair">' + markup + sections[i + 1][0] + '</div>')
            i += 2
        else:
            rendered.append(markup)
            i += 1
    return '<div class="section-grid">' + ''.join(rendered) + '</div>'


def css(tokens: dict[str, Any], paper: str) -> str:
    color = tokens["color"]
    page = tokens["page"][paper]
    fonts = ROOT / "assets" / "fonts"
    montserrat = fonts / "Montserrat-Variable.ttf"
    roboto = fonts / "Roboto-Variable.ttf"
    if not montserrat.is_file() or not roboto.is_file():
        raise FileNotFoundError("Fontes Montserrat e Roboto incorporadas ausentes")
    montserrat_data = base64.b64encode(montserrat.read_bytes()).decode("ascii")
    roboto_data = base64.b64encode(roboto.read_bytes()).decode("ascii")
    if paper == "presentation_16_9":
        size = "13.333in 7.5in"
        margins = "0.4in"
    else:
        size = f"{page['width']}mm {page['height']}mm"
        margins = f"{page['margin_top']}mm {page['margin_right']}mm {page['margin_bottom']}mm {page['margin_left']}mm"
    iqb_columns = "\n".join(
        f"table.scores-{slug} col:nth-child({index}) {{ width: {width}%; }}"
        for slug, widths in IQB_TABLE_WIDTHS.items() for index, width in enumerate(widths, 1)
    )
    return f"""
@font-face {{ font-family: Montserrat; src: url(data:font/ttf;base64,{montserrat_data}) format('truetype'); font-weight: 100 900; }}
@font-face {{ font-family: Roboto; src: url(data:font/ttf;base64,{roboto_data}) format('truetype'); font-weight: 100 900; }}
@page {{ size: {size}; margin: {margins}; }}
* {{ box-sizing: border-box; }}
html {{ background: {color['surface_muted']}; }}
body {{ max-width: {page.get('width', 338.66)}mm; margin: 0 auto; background: {color['canvas']}; color: {color['text_body']}; font-family: Roboto, Arial, sans-serif; font-size: 9.5pt; line-height: 1.35; }}
.document {{ padding: 0; }}
.masthead {{ display: flex; align-items: center; justify-content: space-between; min-height: 16mm; border-bottom: 1.5pt solid {color['document_primary']}; margin-bottom: 12pt; }}
.masthead img {{ max-width: 38mm; max-height: 11mm; object-fit: contain; }}
.eyebrow {{ font-family: Montserrat, Arial, sans-serif; text-transform: uppercase; font-size: 8pt; font-weight: 700; letter-spacing: .12em; color: {color['document_accent']}; }}
h1, h2, h3 {{ font-family: Montserrat, Arial, sans-serif; color: {color['document_primary']}; break-after: avoid; }}
h1 {{ font-size: 30pt; line-height: 1.1; margin: 6pt 0 16pt; }}
h2 {{ font-size: 13pt; line-height: 1.15; border-bottom: .75pt solid {color['border']}; margin: 16pt 0 8pt; padding-bottom: 4pt; }}
h3 {{ font-size: 11pt; line-height: 1.2; margin: 12pt 0 5pt; }}
p, li {{ orphans: 2; widows: 2; }}
.summary {{ font-size: 10.5pt; color: {color['text_strong']}; margin: 0 0 16pt; }}
.block {{ margin: 8pt 0; break-inside: avoid; }}
.block.long {{ break-inside: auto; }}
.callout, .metric, .panel {{ padding: 12pt; border-radius: 6pt; }}
.callout {{ color: {color['inverse']}; background: {color['document_deep']}; border-left: 3pt solid {color['document_accent']}; }}
.metric {{ color: {color['document_primary']}; background: {color['document_accent_soft']}; font-size: 16pt; font-weight: 700; }}
.panel {{ background: {color['surface_muted']}; border: .75pt solid {color['border']}; }}
.missing {{ border-left: 3pt solid {color['danger']}; background: {color['danger_soft']}; padding: 12pt; }}
table {{ border-collapse: collapse; width: 100%; font-size: 8pt; break-inside: auto; }}
th {{ background: {color['document_primary']}; color: {color['inverse']}; text-align: left; }}
td, th {{ border: .75pt solid {color['border']}; padding: 6pt 8pt; vertical-align: top; overflow-wrap: anywhere; }}
tr {{ break-inside: avoid; }}
table.scores {{ table-layout: fixed; font-size: 8pt; }}
table.scores td, table.scores th {{ padding: 4pt; white-space: normal; overflow-wrap: anywhere; }}
table.scores caption {{ text-align: left; font-weight: 700; padding: 8pt 0 4pt; break-after: avoid; }}
{iqb_columns}
.refs {{ margin-top: 18pt; border-top: .75pt solid {color['border']}; font-size: 8pt; }}
.meta {{ margin-top: 12pt; color: {color['text_muted']}; font-size: 7.5pt; }}
.page-break {{ break-before: page; }}
.card-field {{ border: .75pt solid {color['border']}; border-left: 3pt solid {color['document_accent']}; border-radius: 4pt; padding: 12pt; margin: 10pt 0; break-inside: avoid; }}
.card-field h2 {{ margin: 0 0 7pt; border: 0; padding: 0; }}
.card-field p {{ margin: 0; }}
.section-grid {{ display: block; }}
.section-card {{ width: 100%; margin: 0 0 10pt; background: {color['surface']}; border: .75pt solid {color['border']}; border-radius: 6pt; overflow: hidden; break-inside: avoid; }}
.section-card.wide {{ break-inside: avoid; }}
.section-pair {{ display: table; table-layout: fixed; width: 100%; border-spacing: 5pt 0; margin: 0 -5pt 10pt; break-inside: avoid; }}
.section-pair .section-card {{ display: table-cell; width: 50%; vertical-align: top; }}
.section-band {{ background: {color['document_deep']}; color: {color['inverse']}; padding: 9pt 12pt; font-family: Montserrat, Arial, sans-serif; font-size: 11pt; font-weight: 700; text-transform: uppercase; letter-spacing: .035em; }}
.section-content {{ padding: 10pt 11pt; }}
.section-content .block:first-child, .section-content p:first-child {{ margin-top: 0; }}
.card-grid-block {{ width: 100%; }}
.grid-title {{ margin: 0 0 7pt; border: 0; }}
.card-grid {{ display: block; }}
.card-row {{ display: table; table-layout: fixed; width: 100%; border-spacing: 6pt 0; margin: 0 -6pt 7pt; break-inside: avoid; }}
.grid-card {{ display: table-cell; width: 50%; vertical-align: top; padding: 9pt 10pt; border-radius: 5pt; border-left: 3pt solid {color['document_accent']}; background: {color['document_accent_soft']}; break-inside: avoid; }}
.card-grid.columns-1 .grid-card {{ width: 100%; }}
.card-grid.columns-3 .grid-card {{ width: 33.333%; }}
.card-grid.columns-4 .grid-card {{ width: 25%; }}
.grid-card h3 {{ font-size: 9pt; text-transform: uppercase; letter-spacing: .035em; margin: 0 0 5pt; }}
.grid-card p {{ margin: 0 0 4pt; }}
.grid-card ul {{ margin: 3pt 0 0; padding-left: 13pt; }}
.grid-value {{ font-size: 16pt; font-weight: 700; color: {color['document_primary']}; margin-bottom: 3pt; }}
.grid-card.tone-neutral {{ background: {color['surface_muted']}; border-left-color: {color['border_strong']}; }}
.grid-card.tone-red {{ background: {color['danger_soft']}; border-left-color: {color['danger']}; }}
.grid-card.tone-green {{ background: #E8F4EE; border-left-color: #3F9270; }}
.grid-card.tone-yellow {{ background: #FFF4D9; border-left-color: #BD8A28; }}
.section-content table tbody tr:nth-child(even) {{ background: {color['surface_muted']}; }}
@media screen and (max-width: 740px) {{ .card-row {{ display: block; }} .grid-card {{ display: block; width: 100% !important; margin-bottom: 7pt; }} }}
@media screen {{ body {{ box-shadow: 0 3px 28px #10266126; padding: 16mm; }} }}
@media print {{ html, body {{ background: {color['canvas']}; }} body {{ max-width: none; box-shadow: none; padding: 0; }} }}
"""


def html_block(block: dict[str, Any]) -> str:
    kind = block["type"]
    data = block.get("data", {})
    label = html.escape(block_text(block)).replace("\n", "<br>")
    refs = "".join(f'<sup>[{html.escape(x)}]</sup>' for x in block.get("citation_ids", []))
    if kind == "page_break":
        return '<div class="page-break"></div>'
    if kind == "heading":
        level = int(data.get("level", 2)) if str(data.get("level", 2)).isdigit() else 2
        level = min(3, max(2, level))
        return f'<h{level} id="{html.escape(block["id"])}">{label}</h{level}>'
    if kind == "card_grid":
        return html_card_grid(block)
    if kind == "table":
        iqb = isinstance(data.get("dimensions"), dict)
        tables = []
        for index, (slug, legend, headers, rows) in enumerate(table_views(block)):
            table_class = f' class="scores scores-{slug}"' if iqb else ""
            columns = "<colgroup>" + "<col>" * len(headers) + "</colgroup>" if iqb else ""
            heads = "".join(f"<th>{html.escape(x)}</th>" for x in headers)
            body = "".join("<tr>" + "".join(f"<td>{html.escape(v)}</td>" for v in row) + "</tr>" for row in rows)
            if index == 0 and data.get("title"):
                legend = data["title"]
            caption = f"<caption>{html.escape(str(legend))}</caption>" if legend else ""
            tables.append(f'<table{table_class}>{caption}{columns}<thead><tr>{heads}</tr></thead><tbody>{body}</tbody></table>')
        return f'<div class="block long">{"".join(tables)}{refs}</div>'
    if kind == "list":
        items = data.get("items", [])
        tag = "ol" if data.get("ordered") else "ul"
        return f'<div class="block long"><{tag}>' + "".join(f"<li>{html.escape(str(x))}</li>" for x in items) + f"</{tag}>{refs}</div>"
    if kind in ("callout", "metric"):
        return f'<div class="block {kind}">{label}{refs}</div>'
    if kind in ("image_ref", "chart", "canvas"):
        alt = html.escape(str(data.get("alt") or data.get("description") or "Elemento visual sem descrição"))
        return f'<div class="block panel"><strong>{label}</strong><p>{alt}</p>{refs}</div>'
    status = "missing" if data.get("status") in ("missing", "pending", "inconclusive") else "block long"
    return f'<p class="{status}">{label}{refs}</p>'


def load_presentation_layout(deliverable_type: str) -> dict[str, dict[str, Any]] | None:
    """Lê _DESIGN_SYSTEM/layouts/<tipo>.layout.json e devolve os slots por nome,
    em milímetros (CSS aceita mm nativamente — sem conversão de unidade).
    None se o arquivo não existir."""
    path = ROOT / "layouts" / f"{deliverable_type}.layout.json"
    if not path.is_file():
        return None
    data = json.loads(path.read_text(encoding="utf-8"))
    return {s["name"]: s for s in data.get("slots", []) if s.get("name")}


def render_html(content: dict[str, Any], tokens: dict[str, Any], profile: dict[str, Any]) -> str:
    logo = ROOT / "assets" / "logo_ccee.png"
    if not logo.is_file():
        raise FileNotFoundError("Logo CCEE ausente; não é possível assegurar timbre idempotente")
    uri = "data:image/png;base64," + base64.b64encode(logo.read_bytes()).decode("ascii")
    meta = content["metadata"]
    title = html.escape(content["title"])
    deliverable_label = html.escape(DELIVERABLE_LABELS[content["deliverable_type"]])
    if profile["layout"] == "presentation":
        sections: list[tuple[str, list[dict[str, Any]]]] = []
        current_title = "Conteúdo"
        current_blocks: list[dict[str, Any]] = []
        for block in content["blocks"]:
            if block["type"] in ("heading", "page_break"):
                if current_blocks:
                    sections.append((current_title, current_blocks))
                    current_blocks = []
                if block["type"] == "heading":
                    current_title = block_text(block)
                continue
            current_blocks.append(block)
        if current_blocks:
            sections.append((current_title, current_blocks))
        layout_slots = load_presentation_layout(content["deliverable_type"])

        def slot_css(selector: str, slot_name: str) -> str:
            slot = layout_slots.get(slot_name) if layout_slots else None
            if not slot or not all(k in slot for k in ("x", "y", "width", "height")):
                return ""
            return (f"{selector} {{ position: absolute; left: {slot['x']}mm; top: {slot['y']}mm; "
                     f"width: {slot['width']}mm; height: {slot['height']}mm; }} ")

        def wrap_slide_content(inner_html: str) -> str:
            return f'<div class="slide-content">{inner_html}</div>' if layout_slots else inner_html

        masthead = f'<header class="masthead"><span class="eyebrow">{deliverable_label}</span><img src="{uri}" alt="CCEE"></header>'
        first_summary = f'<p class="summary">{html.escape(content.get("summary", ""))}</p>'
        slides = [f'<section class="slide">{masthead}<h1>{title}</h1>{wrap_slide_content(first_summary)}</section>']
        for section_title, section_blocks in sections:
            body_html = "\n".join(html_block(block) for block in section_blocks)
            slides.append(f'<section class="slide">{masthead}<h1>{html.escape(section_title)}</h1>'
                          + wrap_slide_content(body_html) + '</section>')
        if content["citations"]:
            refs = ''.join(f'<li><strong>{html.escape(c["id"])}</strong> — {html.escape(c["source_id"])}'
                           + (f', {html.escape(c["locator"])}' if c.get("locator") else '') + '</li>'
                           for c in content["citations"])
            slides.append(f'<section class="slide">{masthead}<h1>Fontes</h1>' + wrap_slide_content(f'<ol>{refs}</ol>') + '</section>')
        slide_footer = f'<footer class="meta">{html.escape(identity_line(content))}</footer>'
        slides = [slide.removesuffix("</section>") + slide_footer + "</section>" for slide in slides]
        slide_css = "@media print { .slide { break-after: page; } .slide:last-child { break-after: auto; } } @media screen { .slide { min-height: 170mm; border-bottom: 1px solid #ddd; padding-bottom: 10mm; } }"
        slide_css += (f'.slide .meta {{ margin-top: 10pt; font-family: Roboto; '
                      f'font-size: {tokens["typography"]["styles"]["footer"]["size"]}pt; '
                      f'color: {tokens["color"]["text_muted"]}; }}')
        if layout_slots:
            slide_css += " .slide { position: relative; }"
            slide_css += slot_css(".slide .masthead", "logo")
            slide_css += slot_css(".slide h1", "titulo_slide")
            slide_css += slot_css(".slide .slide-content", "corpo_slide")
            slide_css += slot_css(".slide .meta", "rodape_slide")
        return ('<!doctype html>\n<html lang="pt-BR"><head><meta charset="utf-8">'
                f'<title>{title}</title><style>{css(tokens, profile["paper"])}{slide_css}</style></head>'
                f'<body><article class="document">{"".join(slides)}</article></body></html>\n')
    if content["deliverable_type"] in CARD_LABELS:
        fields = CARD_LABELS[content["deliverable_type"]]
        rendered = []
        for block in content["blocks"]:
            key = block.get("data", {}).get("section_key")
            if key in fields:
                value = html.escape(block_text(block)).replace("\n", "<br>")
                refs = ''.join(f'<sup>[{html.escape(x)}]</sup>' for x in block.get("citation_ids", []))
                rendered.append(f'<section class="card-field"><h2>{html.escape(fields[key])}</h2><p>{value}{refs}</p></section>')
            else:
                rendered.append(html_block(block))
        blocks = "\n".join(rendered)
    else:
        blocks = html_report_blocks(content["blocks"])
    citations = "".join(
        f"<li><strong>{html.escape(c['id'])}</strong> — {html.escape(c['source_id'])}"
        + (f", {html.escape(c['locator'])}" if c.get("locator") else "") + "</li>"
        for c in content["citations"]
    )
    refs = f'<section class="refs"><h2>Fontes</h2><ol>{citations}</ol></section>' if citations else ""
    summary = f'<p class="summary">{html.escape(content["summary"])}</p>' if content.get("summary") else ""
    variant_css = ""
    if profile["layout"] == "executive_one_page":
        variant_css = (".masthead { min-height: 12mm; margin-bottom: 6pt; } "
                       "h1 { font-size: 24pt; margin: 4pt 0 8pt; } "
                       ".summary { margin-bottom: 8pt; } "
                       ".block { margin: 5pt 0; } .callout { padding: 8pt; } "
                       ".refs { margin-top: 8pt; } "
                       ".section-grid { column-count: 2; column-gap: 8pt; } "
                       ".section-card { display: inline-block; width: 100%; margin-bottom: 6pt; } "
                       ".section-pair { display: block; margin: 0; } .section-pair .section-card { display: block; width: 100%; } "
                       ".section-band { padding: 5pt 7pt; font-size: 9pt; } "
                       ".section-content { padding: 5pt 7pt; font-size: 8.5pt; }")
    elif profile["layout"] == "report_landscape":
        variant_css = (".masthead { margin-bottom: 8pt; } "
                       "h1 { font-size: 24pt; margin: 4pt 0 9pt; } "
                       "h2 { margin: 10pt 0 5pt; } "
                       ".summary { margin-bottom: 9pt; } "
                       ".block { margin: 6pt 0; } .refs { margin-top: 9pt; }")
    elif profile["layout"] == "report_portrait":
        variant_css = (".masthead { margin-bottom: 8pt; } "
                       "h1 { font-size: 25pt; margin: 4pt 0 9pt; } "
                       "h2 { margin: 10pt 0 5pt; } "
                       ".summary { margin-bottom: 9pt; } "
                       ".block { margin: 6pt 0; } .refs { margin-top: 9pt; }")
    footer_label = identity_line(content)
    if content["deliverable_type"] == "learning_card" and meta.get("test_card_id"):
        footer_label += f' · Test Card {meta["test_card_id"]} v{meta.get("test_card_version", "")}'
    css_footer_label = json.dumps(footer_label, ensure_ascii=False).replace("<", "\\003c ")
    footer_css = (f'@page {{ @bottom-right {{ content: {css_footer_label}; '
                  f'font-family: Roboto; font-size: 7.5pt; color: {tokens["color"]["text_muted"]}; }} }} '
                  '@media print { footer.meta { display: none; } }')
    return ("<!doctype html>\n<html lang=\"pt-BR\"><head><meta charset=\"utf-8\">"
            f"<title>{title}</title><style>{css(tokens, profile['paper'])}{variant_css}{footer_css}</style></head><body><article class=\"document\">"
            f'<header class="masthead"><span class="eyebrow">{deliverable_label}</span><img src="{uri}" alt="CCEE"></header>'
            f"<h1>{title}</h1>{summary}{blocks}{refs}"
            f'<footer class="meta">{html.escape(footer_label)}</footer>'
            "</article></body></html>\n")


def render_md(content: dict[str, Any]) -> str:
    def grid_cell(text: str) -> str:
        return text.replace("|", "\\|").replace("\r\n", "\n").replace("\r", "\n").replace("\n", "<br>")

    lines = [f"# {content['title']}", "", identity_line(content), ""]
    meta = content.get("metadata", {})
    if content["deliverable_type"] == "learning_card" and meta.get("test_card_id"):
        lines += [f'Test Card: {meta["test_card_id"]} v{meta.get("test_card_version", "")}', ""]
    if content.get("summary"):
        lines += [content["summary"], ""]
    for block in content["blocks"]:
        kind = block["type"]
        card_key = block.get("data", {}).get("section_key")
        card_label = CARD_LABELS.get(content["deliverable_type"], {}).get(card_key)
        if kind == "page_break":
            continue
        if card_label:
            lines += ["## " + card_label, "", block_text(block), ""]
        elif kind == "heading":
            lines += ["## " + block_text(block), ""]
        elif kind == "card_grid":
            if block["data"].get("title"):
                lines += ["### " + str(block["data"]["title"]), ""]
            lines += ["| Card | Conteúdo |", "| --- | --- |"]
            for card in block["data"]["cards"]:
                text = card_text(card)
                if card.get("citation_ids"):
                    text += "\nFontes: " + ", ".join(card["citation_ids"])
                lines.append("| " + grid_cell(str(card["title"])) + " | " + grid_cell(text) + " |")
            lines.append("")
        elif kind == "table":
            for _, legend, headers, rows in table_views(block):
                if legend:
                    lines += ["### " + str(legend), ""]
                if headers:
                    lines += ["| " + " | ".join(value.replace("|", "\\|") for value in headers) + " |",
                              "| " + " | ".join("---" for _ in headers) + " |"]
                lines += ["| " + " | ".join(value.replace("|", "\\|") for value in row) + " |" for row in rows] + [""]
        elif kind == "list":
            lines += [f"- {item}" for item in block.get("data", {}).get("items", [])] + [""]
        else:
            lines += [block_text(block), ""]
        if block.get("citation_ids"):
            lines += ["Fontes: " + ", ".join(block["citation_ids"]), ""]
    if content["citations"]:
        lines += ["## Fontes", ""]
        lines += [f"- {c['id']}: {c['source_id']} {c.get('locator', '')}" for c in content["citations"]]
    return "\n".join(lines).rstrip() + "\n"


def render_docx(content: dict[str, Any], tokens: dict[str, Any], profile: dict[str, Any], path: Path) -> None:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Cm, Pt, RGBColor
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError as exc:
        raise RuntimeError("DOCX requer python-docx: pip install python-docx") from exc
    def shade(cell: Any, fill: str) -> None:
        element = OxmlElement("w:shd")
        element.set(qn("w:fill"), fill.lstrip("#"))
        cell._tc.get_or_add_tcPr().append(element)

    def band(paragraph: Any) -> None:
        element = OxmlElement("w:shd")
        element.set(qn("w:fill"), tokens["color"]["document_deep"].lstrip("#"))
        paragraph._p.get_or_add_pPr().append(element)
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor(255, 255, 255)

    doc = Document()
    section = doc.sections[0]
    page = tokens["page"][profile["paper"]]
    section.page_width, section.page_height = Cm(page["width"] / 10), Cm(page["height"] / 10)
    section.top_margin, section.right_margin = Cm(page["margin_top"] / 10), Cm(page["margin_right"] / 10)
    section.bottom_margin, section.left_margin = Cm(page["margin_bottom"] / 10), Cm(page["margin_left"] / 10)
    style = doc.styles["Normal"]
    style.font.name = tokens["typography"]["font_body"]
    style.font.size = Pt(tokens["typography"]["styles"]["body"]["size"])
    for key in ("Title", "Heading 1", "Heading 2"):
        doc.styles[key].font.name = tokens["typography"]["font_brand"]
        doc.styles[key].font.color.rgb = RGBColor.from_string(tokens["color"]["document_primary"].lstrip("#"))
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.add_run().add_picture(str(ROOT / "assets" / "logo_ccee.png"), width=Cm(3.2))
    doc.add_heading(content["title"], 0)
    identity = doc.add_paragraph(identity_line(content))
    for run in identity.runs:
        run.font.size = Pt(tokens["typography"]["styles"]["footer"]["size"])
        run.font.color.rgb = RGBColor.from_string(tokens["color"]["text_muted"].lstrip("#"))
    meta = content.get("metadata", {})
    if content["deliverable_type"] == "learning_card" and meta.get("test_card_id"):
        doc.add_paragraph(f'Test Card: {meta["test_card_id"]} v{meta.get("test_card_version", "")}')
    if content.get("summary"):
        doc.add_paragraph(content["summary"])
    for block in content["blocks"]:
        kind = block["type"]
        card_key = block.get("data", {}).get("section_key")
        card_label = CARD_LABELS.get(content["deliverable_type"], {}).get(card_key)
        if card_label:
            doc.add_heading(card_label, 1)
        if kind == "heading":
            band(doc.add_heading(block_text(block), 1))
        elif kind == "card_grid":
            cards = block["data"]["cards"]
            columns = block["data"].get("columns", 2)
            for offset in range(0, len(cards), columns):
                row_cards = cards[offset:offset + columns]
                table = doc.add_table(rows=1, cols=len(row_cards))
                table.style = "Table Grid"
                for cell, card in zip(table.rows[0].cells, row_cards):
                    cell.text = card_text(card)
                    tones = {"blue": tokens["color"]["document_accent_soft"], "neutral": tokens["color"]["surface_muted"], "red": tokens["color"]["danger_soft"], "green": "#E8F4EE", "yellow": "#FFF4D9"}
                    shade(cell, tones.get(card.get("tone", "blue"), tones["blue"]))
                    if card.get("citation_ids"):
                        cell.add_paragraph("Fontes: " + ", ".join(card["citation_ids"]))
        elif kind == "table":
            for slug, legend, heads, rows in table_views(block):
                if legend:
                    doc.add_paragraph(str(legend))
                width = max(1, len(heads), *(len(row) for row in rows))
                table = doc.add_table(rows=1 if heads else 0, cols=width)
                table.style = "Table Grid"
                if heads:
                    for cell, value in zip(table.rows[0].cells, heads):
                        cell.text = value
                for row in rows:
                    cells = table.add_row().cells
                    for cell, value in zip(cells, row):
                        cell.text = value
                if isinstance(block["data"].get("dimensions"), dict):
                    table.autofit = False
                    available_width = section.page_width - section.left_margin - section.right_margin
                    for column, percentage in zip(table.columns, IQB_TABLE_WIDTHS[slug]):
                        column.width = int(available_width * percentage / 100)
                    for row in table.rows:
                        for cell, column in zip(row.cells, table.columns):
                            cell.width = column.width
                            for paragraph in cell.paragraphs:
                                for run in paragraph.runs:
                                    run.font.size = Pt(8)
        elif kind == "list":
            for item in block.get("data", {}).get("items", []):
                doc.add_paragraph(str(item), style="List Bullet")
        elif kind == "page_break":
            doc.add_page_break()
        else:
            if content["deliverable_type"] in CARD_LABELS:
                doc.add_paragraph(block_text(block))
            else:
                table = doc.add_table(rows=1, cols=1)
                table.style = "Table Grid"
                cell = table.rows[0].cells[0]
                cell.text = block_text(block)
                shade(cell, tokens["color"]["surface_muted"])
        if block.get("citation_ids"):
            doc.add_paragraph("Fontes: " + ", ".join(block["citation_ids"]))
    if content["citations"]:
        doc.add_heading("Fontes", 1)
        for citation in content["citations"]:
            doc.add_paragraph(f"{citation['id']}: {citation['source_id']} {citation.get('locator', '')}")
    doc.core_properties.title = content["title"]
    doc.core_properties.author = "my_skills"
    doc.save(path)


def render_pptx(content: dict[str, Any], tokens: dict[str, Any], path: Path) -> None:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
    except ImportError as exc:
        raise RuntimeError("PPTX requer python-pptx: pip install python-pptx") from exc
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    primary = RGBColor.from_string(tokens["color"]["document_primary"].lstrip("#"))
    logo = ROOT / "assets" / "logo_ccee.png"

    def slide(title: str, body: str) -> None:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_picture(str(logo), Inches(11.1), Inches(0.35), width=Inches(1.7))
        title_box = s.shapes.add_textbox(Inches(.5), Inches(.5), Inches(10.3), Inches(.8))
        run = title_box.text_frame.paragraphs[0].add_run()
        run.text, run.font.name, run.font.size, run.font.color.rgb = title, tokens["typography"]["font_brand"], Pt(28), primary
        body_box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.6), Inches(1.55), Inches(12.1), Inches(5.35))
        body_box.fill.solid()
        body_box.fill.fore_color.rgb = RGBColor.from_string(tokens["color"]["document_accent_soft"].lstrip("#"))
        body_box.line.color.rgb = RGBColor.from_string(tokens["color"]["border"].lstrip("#"))
        body_box.text_frame.margin_left = Inches(.2)
        body_box.text_frame.margin_top = Inches(.2)
        body_box.text_frame.word_wrap = True
        for line in body.splitlines() or [""]:
            para = body_box.text_frame.paragraphs[0] if not body_box.text_frame.text else body_box.text_frame.add_paragraph()
            para.text = line
            para.font.name, para.font.size = tokens["typography"]["font_body"], Pt(17)
        footer = s.shapes.add_textbox(Inches(.6), Inches(7.02), Inches(12.1), Inches(.33))
        footer.text_frame.word_wrap = True
        paragraph = footer.text_frame.paragraphs[0]
        paragraph.text = identity_line(content)
        paragraph.font.name = tokens["typography"]["font_body"]
        paragraph.font.size = Pt(10)
        paragraph.font.color.rgb = RGBColor.from_string(tokens["color"]["text_muted"].lstrip("#"))

    slide(content["title"], content.get("summary", ""))
    title = "Conteúdo"
    lines: list[str] = []
    for block in content["blocks"]:
        if block["type"] == "heading" or block["type"] == "page_break":
            if lines:
                slide(title, "\n".join(lines))
                lines = []
            if block["type"] == "heading":
                title = block_text(block)
            continue
        line = block_text(block)
        if block["type"] == "table":
            views = []
            for _, legend, heads, rows in table_views(block):
                text = " | ".join(heads) + "\n" + "\n".join(" | ".join(row) for row in rows)
                views.append((str(legend) + "\n" if legend else "") + text)
            line = "\n\n".join(views)
        if block.get("citation_ids"):
            line += " [" + ", ".join(block["citation_ids"]) + "]"
        for chunk in [line[i:i+650] for i in range(0, len(line), 650)] or [""]:
            if sum(len(x) for x in lines) + len(chunk) > 650 or len(lines) >= 8:
                slide(title, "\n".join(lines))
                title += " (continuação)"
                lines = []
            lines.append(chunk)
    if lines:
        slide(title, "\n".join(lines))
    if content["citations"]:
        slide("Fontes", "\n".join(f"{c['id']} — {c['source_id']}" + (f", {c['locator']}" if c.get('locator') else '')
                                     for c in content["citations"]))
    prs.save(path)


def write_csv_tables(content: dict[str, Any], out: Path, stem: str) -> list[Path]:
    written = []
    for block in content["blocks"]:
        export = block.get("export", {})
        if block["type"] != "table" or not export.get("csv"):
            continue
        variant = safe_name(export["variant_key"])
        for index, (slug, _, headers, rows) in enumerate(table_views(block)):
            suffix = f"-{slug}" if index else ""
            path = out / f"{stem}-{variant}{suffix}.csv"
            with path.open("w", encoding="utf-8", newline="") as handle:
                writer = csv.writer(handle)
                if headers:
                    writer.writerow(headers)
                writer.writerows(rows)
            written.append(path)
    if not written:
        raise ValueError("CSV solicitado sem bloco table marcado export.csv=true")
    return written


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("--out-dir", required=True, type=Path)
    parser.add_argument("--formats", help="json,pdf,docx,md,txt; pptx só para apresentação; csv só para tabelas exportáveis")
    args = parser.parse_args()
    content = read_json(args.input)
    tokens = read_json(ROOT / "tokens.json")
    profiles = read_json(ROOT / "profiles.json")
    profile = validate(content, profiles)
    allowed = profile.get("formats", profiles["default_formats"]) + ["html", "csv"]
    requested = args.formats.split(",") if args.formats else profile.get("formats", profiles["default_formats"])
    requested = list(dict.fromkeys(["json", *requested]))
    invalid = set(requested) - set(allowed)
    if invalid:
        raise ValueError(f"Formatos incompatíveis com {content['deliverable_type']}: {sorted(invalid)}")
    out = args.out_dir
    out.mkdir(parents=True, exist_ok=True)
    stem = safe_name(content["metadata"]["artifact_id"]) + "-v" + safe_name(content["metadata"]["content_version"])
    files: list[Path] = []
    canonical = canonical_bytes(content)
    path = out / f"{stem}.json"
    path.write_bytes(canonical)
    files.append(path)
    source_html: str | None = None
    if "html" in requested or "pdf" in requested:
        source_html = render_html(content, tokens, profile)
    if "html" in requested:
        path = out / f"{stem}.html"
        path.write_text(source_html or "", encoding="utf-8")
        files.append(path)
    unavailable: dict[str, str] = {}
    if "pdf" in requested:
        path = out / f"{stem}.pdf"
        try:
            from weasyprint import HTML
            HTML(string=source_html or "", base_url=str(ROOT)).write_pdf(str(path))
        except (ImportError, OSError) as exc:
            import sys

            path.unlink(missing_ok=True)
            unavailable["pdf"] = f"{type(exc).__name__}: {exc}"[:300]
            print(f"AVISO: PDF não gerado ({unavailable['pdf']}). No macOS: brew install pango; "
                  "no Debian/Ubuntu: apt-get install libpango-1.0-0 libpangoft2-1.0-0.", file=sys.stderr)
        else:
            if profile["layout"] in ("executive_one_page", "card_portrait"):
                try:
                    from pypdf import PdfReader
                except ImportError as exc:
                    raise RuntimeError("Verificação de página única requer pypdf: pip install pypdf") from exc
                count = len(PdfReader(str(path)).pages)
                if count != 1:
                    path.unlink(missing_ok=True)
                    raise ValueError(f"{profile['layout']} excedeu uma página ({count}); revise ou divida o conteúdo sem reduzir a fonte")
            files.append(path)
    if "md" in requested or "txt" in requested:
        markdown = render_md(content)
        if "md" in requested:
            path = out / f"{stem}.md"
            path.write_text(markdown, encoding="utf-8")
            files.append(path)
        if "txt" in requested:
            path = out / f"{stem}.txt"
            path.write_text(re.sub(r"^#{1,6} ", "", markdown, flags=re.MULTILINE), encoding="utf-8")
            files.append(path)
    if "docx" in requested:
        path = out / f"{stem}.docx"
        render_docx(content, tokens, profile, path)
        files.append(path)
    if "pptx" in requested:
        path = out / f"{stem}.pptx"
        render_pptx(content, tokens, path)
        files.append(path)
    if "csv" in requested:
        files.extend(write_csv_tables(content, out, stem))
    manifest = {
        "artifact_id": content["metadata"]["artifact_id"],
        "content_version": content["metadata"]["content_version"],
        "deliverable_type": content["deliverable_type"],
        "layout": profile["layout"],
        "paper": profile["paper"],
        "renderer_version": RENDERER_VERSION,
        "profile_version": profiles["profile_version"],
        "branding_version": tokens["profile_version"],
        "content_sha256": sha(canonical),
        "profile_sha256": sha(canonical_bytes(profiles)),
        "tokens_sha256": sha(canonical_bytes(tokens)),
        "files": {path.name: sha(path.read_bytes()) for path in files},
        "unavailable": unavailable,
    }
    (out / f"{stem}.manifest.json").write_bytes(canonical_bytes(manifest))
    print(json.dumps({"manifest": f"{stem}.manifest.json", "files": list(manifest["files"]),
                      "unavailable": unavailable}, ensure_ascii=False))


if __name__ == "__main__":
    main()
